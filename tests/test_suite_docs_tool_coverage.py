"""Lock: the suite's documentation names every first-party tool, and counts them right.

The README's table and ``documentation/docs/index.md`` are generated from the tool
manifest, so a new tool appears there by construction. Two artefacts are hand-written
prose over the same population, and both silently drifted: the CLI reference
documented 8 of the 30 tools' commands, and the new-user presentation named 26
of them - a reader could not learn that Tag Editor, Stem Separator, Spectral
Viewer or ML Denoiser existed at all. The same deck then stated "26 Windows
desktop tools" while the suite shipped 30, and the README's opening line still
said 27.

This lock derives the population from the manifest rather than naming tools
itself, so a tool added later fails here until the prose covers it:

* every first-party tool has its own section heading in
  ``documentation/user_guides/reference/CLI_Commands_Reference.md``;
* the reference's *All Commands* table lists every command the CLI registers,
  sub-commands included;
* every first-party tool's name appears in the generated presentation;
* every document that quotes a tool count agrees with the manifest - the total,
  the per-category counts, the Windows-only subset and the deck's window count
  alike.

Each scan asserts it matched something, so a renamed file or a changed heading
style fails instead of passing vacuously.

Changed in 2.1.5: new file - the coverage lock behind the 30-tool documentation
repair (CLI reference and new-user deck), including the tool-count claims the
shipped prose makes.
Changed in 2.1.5: the count-claim walk reads *tracked* documents instead of
walking the directories, so the generated ``documentation/docs/_build/`` tree -
which stages a copy of the root documents beside the site, the historical
``CHANGELOG.md`` among them - is no longer mistaken for shipped prose.

Changed in 2.1.7: the count-claim walk covers ``documentation/`` as one tree.

"""

from __future__ import annotations

import re
import subprocess
from collections.abc import Iterator
from dataclasses import dataclass
from itertools import chain
from pathlib import Path

import click

from audio_suite_shared.cli_app import cli
from audio_suite_shared.platform_support import windows_only_tools
from audio_suite_shared.tool_plugin import ToolManifest, get_registry

REPO_ROOT = Path(__file__).resolve().parent.parent
REFERENCE = REPO_ROOT / "documentation/user_guides/reference/CLI_Commands_Reference.md"
DECK = (
    REPO_ROOT / "documentation/user_guides/suite_overview/"
    "AudioSuite-Pro_New_User_Presentation.pptx"
)

#: The heading level the reference gives one tool.
TOOL_HEADING = re.compile(r"^### (.+)$", re.MULTILINE)

#: One row of the reference's *All Commands* table: ``| `command` | Tool | ...``.
COMMAND_ROW = re.compile(r"^\| `([^`]+)` \|", re.MULTILINE)


def _tracked_paths() -> set[str]:
    """Return every tracked path, POSIX-style, relative to the repository root."""
    result = subprocess.run(
        ["git", "ls-files"],  # B603/B607: fixed argv, shell=False
        cwd=REPO_ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=True,
    )
    return {line.strip() for line in result.stdout.splitlines() if line.strip()}


def _tracked_markdown_under(directory: str) -> list[Path]:
    """Return every *tracked* Markdown document under *directory*.

    Tracked rather than walked: ``utilities/build_docs_site.py`` stages a copy
    of the repository's root documents under ``documentation/docs/_build/`` - the
    historical ``CHANGELOG.md`` among them, whose counts are frozen on purpose.
    Walking the directory re-admitted those generated copies, so the lock went red for
    anyone who had run the documented build step, naming a file they had not
    touched.

    Args:
        directory: A repo-relative directory name, e.g. ``"documentation"``.

    Returns:
        Sorted paths of the tracked ``.md`` files beneath it.
    """
    return sorted(
        REPO_ROOT / path
        for path in _tracked_paths()
        if path.endswith(".md") and path.startswith(f"{directory}/")
    )


def _count_documents() -> tuple[Path, ...]:
    """Return every document whose prose must quote the current tool count.

    Historical records are deliberately absent: ``CHANGELOG.md``, ``PLAN_*.md``
    and ``AUDIT_*.md`` state the count as of the release they describe, so a
    tool added later must not rewrite them. ``AGENTS.md`` is absent for a
    different reason - its "14 tools" count is the frozen GUI completion-seam
    population, not the suite's tool list.
    """
    return (
        REPO_ROOT / "README.md",
        REPO_ROOT / "CONTRIBUTING.md",
        REPO_ROOT / "SECURITY.md",
        # One tree covers both the site's pages and the shipped guides: every
        # authored document lives under ``documentation/`` (plan D20).
        *_tracked_markdown_under("documentation"),
    )


#: "30 first-party tools", "26 Windows desktop tools", "each of the 30 tools".
_TOTAL_CLAIM = re.compile(r"(\d+)\s+(?:[\w][\w-]*\s+){0,4}(?:tools|modules)\b", re.I)

#: "no need to open 26 separate windows".
_WINDOW_CLAIM = re.compile(r"(\d+)\s+separate windows", re.I)

#: "the other 29 of 30 tools".
_SUBSET_CLAIM = re.compile(r"(\d+)\s+of\s+(\d+)\s+tools", re.I)

#: A category-count table row: ``| ENRICH | `enrich` | 11 | ...``.
_CATEGORY_CLAIM = re.compile(
    r"^\|\s*(?:REFINE|ENRICH|UTIL)\s*\|\s*`([a-z]+)`\s*\|\s*(\d+)\s*\|", re.MULTILINE
)

#: The categories the manifest knows about.
CATEGORIES: tuple[str, ...] = ("enrich", "refinery", "utility")


def _first_party_tools() -> list[ToolManifest]:
    """Return the first-party manifests the user-facing documents must cover."""
    tools = get_registry().first_party
    assert tools, "the tool manifest is empty - the scan would pass vacuously"
    return tools


def _registered_commands() -> set[str]:
    """Return every command the CLI registers, sub-commands as ``parent sub``."""
    found: set[str] = set()
    for name, command in cli.commands.items():
        sub_commands = getattr(command, "commands", None)
        if isinstance(command, click.Group) and sub_commands:
            found.update(f"{name} {sub}" for sub in sub_commands)
        else:
            found.add(name)
    assert found, "the CLI registered no commands - the scan would pass vacuously"
    return found


def _deck_text() -> str:
    """Return every text run in the generated presentation."""
    from pptx import Presentation

    prs = Presentation(str(DECK))
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    chunks.extend(cell.text for cell in row.cells)
    text = "\n".join(chunks)
    assert len(text) > 2000, f"the deck scan read only {len(text)} characters"
    return text


@dataclass(frozen=True)
class _CountClaim:
    """One tool-count claim found in a shipped document.

    Attributes:
        source: The file (and slide, for the deck) the claim was found in.
        line: The 1-indexed line or paragraph the claim sits on.
        text: The claim as written, e.g. ``"30 Windows desktop tools"``.
        kind: ``"total"``, ``"windows"``, ``"subset"`` or ``"category"``.
        numbers: The numbers the claim states, in the order it states them.
        category: The manifest category a ``"category"`` claim counts.
    """

    source: str
    line: int
    text: str
    kind: str
    numbers: tuple[int, ...]
    category: str = ""


def _markdown_units() -> Iterator[tuple[str, int, str]]:
    """Yield ``(source, line, text)`` for every line of every counting document."""
    for path in _count_documents():
        source = str(path.relative_to(REPO_ROOT))
        for number, line in enumerate(
            path.read_text(encoding="utf-8").splitlines(), start=1
        ):
            yield source, number, line


def _deck_units() -> Iterator[tuple[str, int, str]]:
    """Yield ``(source, slide, paragraph)`` for the presentation.

    Paragraph by paragraph, never shape by shape joined together: a footer's
    slide number (``17``) directly above a title (``Model-Based Tools``) is not
    a claim about how many tools exist, and a joined string would look like one.
    """
    from pptx import Presentation

    prs = Presentation(str(DECK))
    for index, slide in enumerate(prs.slides, start=1):
        source = f"{DECK.name} (slide {index})"
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    yield source, index, paragraph.text
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            yield source, index, paragraph.text


def _claims_in(source: str, line: int, text: str) -> list[_CountClaim]:
    """Return every tool-count claim written on one line or paragraph."""
    found: list[_CountClaim] = []
    for match in _TOTAL_CLAIM.finditer(text):
        if " of " in match.group(0).lower():
            # "29 of 30 tools" is a subset claim, checked against the
            # Windows-only remainder rather than against the total.
            continue
        found.append(
            _CountClaim(source, line, match.group(0), "total", (int(match.group(1)),))
        )
    for match in _WINDOW_CLAIM.finditer(text):
        found.append(
            _CountClaim(source, line, match.group(0), "windows", (int(match.group(1)),))
        )
    for match in _SUBSET_CLAIM.finditer(text):
        found.append(
            _CountClaim(
                source,
                line,
                match.group(0),
                "subset",
                (int(match.group(1)), int(match.group(2))),
            )
        )
    for match in _CATEGORY_CLAIM.finditer(text):
        found.append(
            _CountClaim(
                source,
                line,
                match.group(0),
                "category",
                (int(match.group(2)),),
                match.group(1),
            )
        )
    return found


def _count_claims() -> list[_CountClaim]:
    """Return every tool-count claim the shipped prose makes."""
    claims: list[_CountClaim] = []
    for source, number, text in chain(_markdown_units(), _deck_units()):
        claims.extend(_claims_in(source, number, text))
    return claims


class TestCliReferenceCoversEveryTool:
    """``CLI_Commands_Reference.md`` documents every tool and every command."""

    def test_every_tool_has_its_own_section(self) -> None:
        headings = set(TOOL_HEADING.findall(REFERENCE.read_text(encoding="utf-8")))

        assert headings, "the reference has no tool sections - the scan matched nothing"
        missing = sorted(
            tool.name for tool in _first_party_tools() if tool.name not in headings
        )
        assert not missing, (
            "these tools have no section in "
            f"{REFERENCE.relative_to(REPO_ROOT)}: {missing}"
        )

    def test_the_all_commands_table_lists_every_command(self) -> None:
        text = REFERENCE.read_text(encoding="utf-8")
        assert "## All Commands" in text, "the reference has no All Commands section"
        table = text.split("## All Commands", 1)[1]
        rows = set(COMMAND_ROW.findall(table))

        assert rows, "the All Commands table parsed to no rows"
        missing = sorted(_registered_commands() - rows)
        assert not missing, (
            f"these CLI commands are missing from the All Commands table: {missing}"
        )

    def test_every_tool_with_a_command_has_a_command_row(self) -> None:
        """A tool's section is not enough: its command must be in the table too."""
        text = REFERENCE.read_text(encoding="utf-8")
        rows = set(COMMAND_ROW.findall(text.split("## All Commands", 1)[1]))
        with_command = [tool for tool in _first_party_tools() if tool.cli_command]

        assert with_command, "no manifest declares a CLI command - scan is vacuous"
        missing = sorted(
            f"{tool.name} ({tool.cli_command})"
            for tool in with_command
            if tool.cli_command not in rows
        )
        assert not missing, (
            "these tools declare a CLI command that the reference does not list: "
            f"{missing}"
        )


class TestDeckCoversEveryTool:
    """The new-user presentation names every tool the suite ships."""

    def test_every_tool_is_named_somewhere_in_the_deck(self) -> None:
        text = _deck_text()

        missing = sorted(
            tool.name for tool in _first_party_tools() if tool.name not in text
        )
        assert not missing, f"these tools are missing from {DECK.name}: {missing}"


class TestToolCountClaims:
    """Every document that quotes a tool count agrees with the manifest.

    A count is the first claim to drift: the deck said 26 while the suite
    shipped 30, and the README's opening line said 27. Each claim is compared
    with the value the manifest implies for *that* kind of claim - the total,
    the Windows-only remainder, or the category's own size - so a new tool
    fails here until the prose matches the tree.
    """

    def test_generated_build_output_is_not_treated_as_shipped_prose(self) -> None:
        """A staged copy under the build tree is not a document to police.

        ``utilities/build_docs_site.py`` writes the site *and* stages copies of
        the root documents beside it, so the historical ``CHANGELOG.md`` - whose
        "26 separate windows" is correct for the release it describes - landed
        under the documentation tree and was read as a live claim. Ageing the deck
        and the changelog then looked like a regression in a file the author never
        edited. The walk follows tracked documents; this proves it by planting
        an untracked one where the build actually puts them.
        """
        probe = REPO_ROOT / "documentation" / "docs" / "_build" / "_lock_probe.md"
        probe.parent.mkdir(parents=True, exist_ok=True)
        probe.write_text(
            "A suite of 999 Windows desktop tools and 999 tools.\n", encoding="utf-8"
        )
        try:
            scanned = {str(path.relative_to(REPO_ROOT)) for path in _count_documents()}
            assert scanned, "the document walk matched nothing - the scan is vacuous"
            assert str(probe.relative_to(REPO_ROOT)) not in scanned, (
                f"{probe.relative_to(REPO_ROOT)} is untracked build output, but the "
                "count-claim walk read it; the walk must follow tracked documents, "
                "not whatever is on disk"
            )
        finally:
            probe.unlink(missing_ok=True)

    def test_every_count_claim_matches_the_manifest(self) -> None:
        claims = _count_claims()

        assert len(claims) >= 5, (
            f"found only {len(claims)} tool-count claims - the scan is vacuous"
        )

        registry = get_registry()
        total = len(_first_party_tools())
        portable = total - len(windows_only_tools())
        expected = {
            "total": (total,),
            "windows": (total,),
            "subset": (portable, total),
        }
        problems: list[str] = []
        for claim in claims:
            if claim.kind == "category":
                # First-party only: the shipped documents describe the
                # distribution, not the user's enabled plugins.
                want = (
                    len(
                        [
                            tool
                            for tool in registry.by_category(claim.category)
                            if not tool.plugin
                        ]
                    ),
                )
            else:
                want = expected[claim.kind]
            if claim.numbers != want:
                problems.append(
                    f"{claim.source}, line {claim.line}: {claim.text!r} states "
                    f"{claim.numbers} but the manifest gives {want}"
                )

        assert not problems, (
            "stale tool counts in the shipped documentation:\n" + "\n".join(problems)
        )

    def test_category_counts_are_actually_checked(self) -> None:
        """A per-category count is a claim too; the scan must see the rows.

        Not every category has to be quoted as a number - REFINE names its one
        tool instead - so this pins that the rows which do exist are read and
        name a category the manifest knows.
        """
        quoted = {
            claim.category for claim in _count_claims() if claim.kind == "category"
        }

        assert quoted, "the category count table was not scanned"
        assert quoted <= set(CATEGORIES), (
            f"the docs quote counts for unknown categories: "
            f"{sorted(quoted - set(CATEGORIES))}"
        )

    def test_the_scan_reads_the_readme_and_the_deck(self) -> None:
        """Pin the two documents whose counts drifted, so a path fix cannot hide it."""
        sources = {claim.source for claim in _count_claims()}

        assert "README.md" in sources, "the README scan matched nothing"
        assert any(DECK.name in source for source in sources), (
            "the deck scan matched nothing"
        )
