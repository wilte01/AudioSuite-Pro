"""Generate a fresh ReplayGain Analyzer guide from current code."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "replaygain_analyzer\\ReplayGain_Analyzer_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("ReplayGain Analyzer - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current ReplayGain Analyzer implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "ReplayGain Analyzer measures integrated loudness "
    "and true peak for supported audio files, "
    "calculates track and folder-level ReplayGain values, "
    "and can write the results to metadata. "
    "Files are grouped by their containing folder: "
    "each folder receives its own album gain and "
    "album peak based on the files measured there."
)
doc.add_paragraph(
    "The tool can optionally apply the measured track "
    "gain losslessly to MP3 files after writing "
    "the tags. This changes MP3 frame gain without "
    "re-encoding the audio."
)

doc.add_heading("Supported Files", level=1)
doc.add_paragraph(
    "The analyzer uses the shared Audio Suite supported "
    "formats, including MP3, FLAC, OGG, M4A, "
    "MP4, Opus, WAV, and AIFF variants. The selected "
    "directory is scanned recursively and folders "
    "are processed independently."
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open ReplayGain_Analyzer.pyw.",
    "Select a music directory.",
    "Choose a valid target LUFS value.",
    "Leave Dry Run selected for the first pass.",
    "Enable Apply Lossless Gain only when you explicitly want MP3 audio gain adjusted.",
    "Click the analysis start button and review the folder/file results.",
    "Run again in Live Update when the calculated values are acceptable.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the directory picker, target LUFS, "
    "run mode, optional MP3 lossless gain, "
    "presets, and shared run controls. The application "
    "validates the directory and target before "
    "starting. Target LUFS uses the suite's permitted "
    "target values."
)
bullet(
    doc,
    "Target LUFS: ",
    "Desired integrated loudness target. Track gain is "
    "target LUFS minus measured LUFS; album gain uses "
    "the mean LUFS for each folder.",
)
bullet(
    doc,
    "Dry Run / Live Update: ",
    "Dry Run measures and previews values without "
    "tag writes. Live Update writes ReplayGain tags.",
)
bullet(
    doc,
    "Apply Lossless Gain: ",
    "After tagging, bakes track gain into MP3 files "
    "losslessly. It applies only to MP3 files and "
    "is not a re-encode.",
)
bullet(
    doc,
    "Presets: ",
    "Save and restore target LUFS and run mode through shared suite presets.",
)

doc.add_heading("Two-Phase Processing", level=1)
for step in (
    "Recursively discover supported audio files and group them by containing folder.",
    "Phase 1 measures each file's integrated LUFS and "
    "true peak in parallel, with retry handling.",
    "For each folder, calculate album gain as target LUFS minus the mean file LUFS.",
    "Calculate album peak as the maximum true peak among files in that folder.",
    "In Dry Run, report the track and album values without writing.",
    "In Live Update, write track gain, album gain, track peak, and album peak tags.",
    "Optionally apply lossless MP3 gain after successful tagging.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("ReplayGain Values and Tags", level=1)
bullet(
    doc,
    "Track Gain: ",
    "The per-file gain required to reach the target, formatted such as +2.34 dB.",
)
bullet(
    doc,
    "Album Gain: ",
    "The folder-level gain derived from the mean integrated LUFS of that folder.",
)
bullet(doc, "Track Peak: ", "The measured true-peak value for the file.")
bullet(doc, "Album Peak: ", "The highest true peak measured among files in the folder.")
doc.add_paragraph(
    "The shared tag writer stores REPLAYGAIN_TRACK_GAIN, "
    "REPLAYGAIN_ALBUM_GAIN, "
    "REPLAYGAIN_TRACK_PEAK, and REPLAYGAIN_ALBUM_PEAK "
    "using format-appropriate metadata "
    "representations. Live writes participate in "
    "the suite's tag-edit undo batch."
)

doc.add_heading("Dry Run and Live Update", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Measures values and displays folder/file results "
    "but never writes tags or applies MP3 gain.",
)
bullet(
    doc,
    "Live Update: ",
    "Writes all four ReplayGain fields for each "
    "measured file, then optionally applies MP3 "
    "lossless gain.",
)
bullet(
    doc,
    "Measurement failure: ",
    "A file with unavailable LUFS metrics is reported "
    "as an error and excluded from the folder "
    "calculation.",
)
bullet(
    doc,
    "Lossless gain failure: ",
    "Tagging may succeed while the optional post-tag "
    "gain reports an error status; review the "
    "per-file result and log.",
)

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows four worker slots and a progress "
    "count. Folder completion updates the file "
    "tree with file name, measured LUFS, track gain, "
    "album gain, and status. Results reports total "
    "files, successfully processed files, errors, "
    "and elapsed time."
)

doc.add_heading("Logs", level=1)
bullet(doc, "Dry Run log: ", "replaygain_dry_run_log.txt in the selected directory.")
bullet(doc, "Live log: ", "replaygain_sync_log.txt in the selected directory.")
bullet(
    doc,
    "Contents: ",
    "Mode, target LUFS, folder album statistics, "
    "track LUFS/gain/peak values, statuses, errors, "
    "and summary.",
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is replaygain. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app replaygain "
    "<folder> --dry\n"
    "python -m audio_suite_shared.cli_app replaygain "
    "<folder> --target-lufs -16 --live\n"
    "python -m audio_suite_shared.cli_app replaygain "
    "<folder> --apply-gain --live\n"
    "python -m audio_suite_shared.cli_app replaygain "
    "<folder> --target-lufs -14 --json"
)
doc.add_paragraph(
    "The CLI accepts --target-lufs, "
    "--apply-gain/--no-apply-gain, "
    "the shared dry-run/live option, "
    "and --json. Target LUFS is validated by the "
    "shared suite rules. The text summary reports "
    "files and errors; JSON contains the structured "
    "result."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open ENRICH and select "
    "ReplayGain Analyzer. The driver passes the "
    "directory, target LUFS, run mode, start/end "
    "behavior, and inherited orchestration controls."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No audio files found: ",
    "Select a directory containing supported audio "
    "files or check that the files are accessible.",
)
bullet(
    doc,
    "A folder's album gain looks unexpected: ",
    "Album gain is calculated per folder from the "
    "mean LUFS of only the files that measured "
    "successfully.",
)
bullet(
    doc,
    "Tags did not change: ",
    "Confirm Live Update was selected. Dry Run never writes ReplayGain metadata.",
)
bullet(
    doc,
    "MP3 loudness changed unexpectedly: ",
    "Apply Lossless Gain explicitly changes MP3 frame "
    "gain after tagging; disable it to write tags "
    "only.",
)
bullet(
    doc,
    "A file errors: ",
    "Review the per-file status and replaygain log "
    "for FFmpeg measurement, tag-writing, or "
    "gain-application details.",
)
bullet(doc, "Target rejected: ", "Use one of the suite's permitted target LUFS values.")

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Run Dry Run first and inspect track and album "
    "values. Live Update changes metadata in place; "
    "optional MP3 gain also changes the encoded "
    "track's playback gain. Keep backups and use "
    "the suite's Undo Tag Edits workflow for "
    "available metadata snapshots."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
