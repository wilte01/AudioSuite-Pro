"""Generate a fresh File Reorganizer guide from the current implementation."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "file_reorganizer\\File_Reorganizer_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("File Reorganizer - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current File Reorganizer implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "File Reorganizer scans supported audio files recursively, reads their metadata, "
    "evaluates a user-defined path template, and copies or moves each file into the "
    "resolved destination hierarchy. It has separate templates for named album artists "
    "and Various Artists releases, selected from the album artist tag."
)
doc.add_paragraph(
    "The workflow is preview-first: Dry Run is enabled "
    "by default and shows the planned "
    "paths without performing file operations."
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open File_Reorganizer.pyw.",
    "Select the Source Directory containing your audio files.",
    "Select the Destination Directory where the organized files should go.",
    "Choose Copy or Move and leave Dry Run enabled.",
    "Review the Preview list after the scan completes.",
    "Disable Dry Run and click Execute only after checking the planned paths.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup and Actions", level=1)
bullet(
    doc,
    "Source Directory: ",
    "Root folder to scan "
    "recursively. The last suite directory may be restored on startup.",
)
bullet(
    doc,
    "Destination Directory: ",
    "Base directory prepended to each resolved pattern path. "
    "It is required by the CLI and selected in the GUI.",
)
bullet(
    doc,
    "Copy: ",
    "Copies files with shutil.copy2, preserving file metadata "
    "such as timestamps and attributes.",
)
bullet(
    doc,
    "Move: ",
    "Moves files with shutil.move. "
    "Across volumes this may behave as copy followed by delete.",
)
bullet(
    doc,
    "Dry Run: ",
    "Evaluates paths and updates progress without copying or moving anything. "
    "Enabled by default.",
)
bullet(
    doc,
    "Preview: ",
    "Runs the scan and populates the planned destination paths. "
    "A successful preview enables Execute.",
)
bullet(doc, "Stop: ", "Requests that the active scan stop at the next file boundary.")
bullet(
    doc,
    "Presets: ",
    "Save and restore source/destination paths, action, dry-run state, "
    "and both pattern templates through shared suite configuration.",
)

doc.add_heading("File Selection and Tags", level=1)
doc.add_paragraph(
    "The engine processes files whose extensions are in the "
    "shared supported-audio list. "
    "It reads the following lowercase pattern fields:"
)
for item in (
    "%albumartist% - album artist, falling back to artist for some uses",
    "%artist% - track artist",
    "%album% - album title",
    "%title% - track title, with the filename stem "
    "as a fallback when raw tags are absent",
    "%tracknumber% and %discnumber% - track/disc positions",
    "%date%, %originaldate%, and %originalyear% - release/original dates",
    "%genre%, %bpm%, and %key% - common descriptive tags",
):
    doc.add_paragraph(item, style="List Bullet")
doc.add_paragraph(
    "A file with neither a usable album nor title is skipped. "
    "A missing or unreadable tag "
    "is generally represented as an empty string, "
    "allowing the pattern's fallback logic to decide the path."
)

doc.add_heading("Specific Artists and Various Artists", level=1)
doc.add_paragraph(
    "The album artist is compared case-insensitively "
    "with the suite's Various Artists aliases. "
    "Matching files use the Various Artists template; all other "
    "files use the Specific Artists template."
)
doc.add_paragraph("The default Specific Artists structure is:")
doc.add_paragraph(
    "Specific "
    "Artists/<first artist character>/<artist, prefix swapped> "
    "/<year> - <album>/<disc>-<track> <title>"
)
doc.add_paragraph("The default Various Artists structure is:")
doc.add_paragraph("Various Artists/<year> - <album>/<disc>-<track> <artist> - <title>")
doc.add_paragraph(
    "The default templates use original year first, "
    "then original date, then release date, "
    "with 0000 as the final fallback. "
    "They also swap a leading A, An, or The to the end for "
    "sorting, zero-pad track numbers, "
    "and use Unknown Album when the album tag is empty."
)

doc.add_heading("Pattern Language", level=1)
doc.add_paragraph("Patterns use a foobar2000-like syntax:")
bullet(doc, "%tagname%: ", "Expands a tag or previously assigned variable.")
bullet(doc, "$set(name, expression): ", "Assigns an intermediate variable.")
bullet(doc, "$if(condition, then, else): ", "Chooses a branch; empty and 0 are false.")
bullet(doc, "$if2(a, b, ...): ", "Returns the first non-empty argument.")
bullet(doc, "$eq(a, b): ", "Returns 1 when strings are equal, otherwise 0.")
bullet(
    doc,
    "$swapprefix(value, A, An, The): ",
    "Moves a matching leading prefix to the end, such as The Beatles to Beatles, The.",
)
bullet(doc, "$left(value, count): ", "Returns the first count characters.")
bullet(doc, "$num(value, digits): ", "Zero-pads a numeric value, such as 1 to 01.")
bullet(doc, "$replace(value, old, new): ", "Replaces text.")
bullet(doc, "$lower(value) / $upper(value): ", "Changes case.")

doc.add_heading("Path Safety", level=1)
doc.add_paragraph(
    "After evaluation, the engine removes "
    "control characters, replaces path-illegal characters "
    'such as < > : " | ? * with underscores, '
    "normalizes separators, strips leading/trailing "
    "separators, trims components, and drops empty components. "
    "If the resolved path has no "
    "extension, the source file's extension is appended."
)

doc.add_heading("Dry Run, Copy, and Move", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Logs source-to-target plans and increments scan/processed statistics "
    "without touching files.",
)
bullet(
    doc,
    "Copy: ",
    "Creates destination directories as needed "
    "and copies the source file with metadata preservation.",
)
bullet(
    doc,
    "Move: ",
    "Creates destination directories as needed and moves the source file.",
)
bullet(
    doc,
    "Undo: ",
    "Live copy and move operations are recorded "
    "in the suite file-reorganization undo context.",
)
bullet(
    doc,
    "Conflict behavior: ",
    "The engine delegates destination handling "
    "to the underlying copy/move operation; "
    "review the Preview and keep backups before Move runs.",
)

doc.add_heading("Progress and Statistics", level=1)
doc.add_paragraph(
    "The GUI reports the current file and four counters: "
    "Scanned, Processed, Skipped, and Errors. "
    "Skips include files with no usable album/title tags and files whose pattern "
    "resolves to an empty "
    "path. Pattern parsing failures and operating-system/value errors increment Errors."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is reorganize. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app reorganize <source> --output-dir "
    "<destination> --dry\n"
    "python -m audio_suite_shared.cli_app reorganize <source> --output-dir "
    "<destination> --action copy\n"
    "python -m audio_suite_shared.cli_app reorganize <source> --output-dir "
    "<destination> --action move --live\n"
    "python -m audio_suite_shared.cli_app reorganize <source> --output-dir "
    "<destination> --specific-pattern <pattern> --json"
)
doc.add_paragraph(
    "CLI options include --output-dir, --action copy|move, "
    "--dry/--live, --specific-pattern, "
    "--various-pattern, and --json. When action or patterns are omitted, "
    "saved shared configuration "
    "values are used. The command defaults to Dry Run."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open UTIL and select Tag-Based File Reorganizer. "
    "The utility is standalone "
    "and opens its own workflow. "
    "The selected directory is passed through the driver launch context; use "
    "the reorganizer's Destination, Action, Pattern Templates, Preview, "
    "and Execute controls for the operation."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "Files are skipped: ",
    "Confirm that the file has usable album/title metadata and that the pattern "
    "produces a non-empty path.",
)
bullet(
    doc,
    "The path is unexpected: ",
    "Use Preview, inspect the source tags, and check "
    "whether Various Artists classification "
    "or prefix swapping selected a different branch.",
)
bullet(
    doc,
    "Pattern error: ",
    "Check parentheses, commas, function names, tag references, "
    "and numeric arguments. The file is counted as an error and the scan continues.",
)
bullet(
    doc,
    "Destination is empty: ",
    "Select a valid destination directory in the GUI or pass --output-dir on the CLI.",
)
bullet(
    doc,
    "Files did not move: ",
    "Dry Run may still be enabled. Disable it only after reviewing Preview.",
)
bullet(
    doc,
    "A move needs recovery: ",
    "Use the suite's Undo Tag Edits/file-reorganization workflow "
    "when a snapshot is available, and retain backups for important libraries.",
)

doc.add_heading("Configuration and Recovery", level=1)
doc.add_paragraph(
    "Destination, action, and pattern settings are persisted "
    "in the shared audio_suite_config.json. "
    "The tool can migrate an older reorganizer_config.json once, "
    "then remove the legacy file. Live "
    "operations are recorded in the suite's file-reorganization batch context."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
