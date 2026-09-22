"""Generate a fresh Album Art Normalizer user guide from the implementation."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "album_art_normalizer\\Album_Art_Normalizer_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Album Art Normalizer - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Album Art Normalizer implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Album Art Normalizer standardizes embedded artwork in every supported audio "
    "format. It reads the first embedded image, center-crops non-square artwork, "
    "resizes it to a selected square resolution, encodes it as an optimized JPEG, "
    "and writes the updated artwork while preserving the other audio metadata."
)
doc.add_paragraph(
    "The GUI also checks for folder.jpg directly inside the selected directory. "
    "When present, it is independently normalized to exactly 1000 x 1000 pixels."
)

doc.add_heading("Supported Input", level=1)
bullet(
    doc,
    "Embedded artwork: ",
    "Read and written format-aware: ID3 APIC frames for "
    "MP3/WAV/AIFF/AIF, Vorbis Picture blocks for "
    "FLAC/OGG/Opus, and MP4 covr atoms for M4A/MP4.",
)
bullet(
    doc,
    "Audio files: ",
    "All nine suite audio formats are processed by the "
    "embedded-art pipeline: .mp3, .flac, .ogg, .m4a, "
    ".mp4, .opus, .wav, .aiff, and .aif.",
)
bullet(doc, "Images: ", "Embedded data and folder.jpg must be readable by Pillow.")

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Album_Art_Normalizer.pyw from the AudioSuite-Pro folder.",
    "On Setup, select the directory containing the music files.",
    "Choose 600 x 600, 800 x 800, 1000 x 1000, or 1200 x 1200.",
    "Leave Dry Run selected for the first pass.",
    "Click BEGIN ARTWORK NORMALIZATION.",
    "Review Processing and Results, then use Live Update when ready.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("GUI Tabs", level=1)
doc.add_heading("Setup", level=2)
doc.add_paragraph(
    "Setup contains the folder picker, embedded resolution selector, run mode, "
    "and shared controls. The folder must exist and be a directory. The default "
    "embedded target is 800 x 800. The last-used directory can be restored from "
    "the shared suite configuration."
)
bullet(doc, "Dry Run: ", "Use for a preview pass before writing artwork changes.")
bullet(
    doc, "Live Update: ", "Permits resized embedded artwork to be saved to audio files."
)
bullet(
    doc,
    "Auto Start: ",
    "Starts processing automatically when inherited from the Master Driver.",
)
bullet(
    doc,
    "Auto End: ",
    "Closes the application about two seconds after completion when enabled.",
)
bullet(
    doc,
    "Clear Logs and Parallel: ",
    "Mirror Master Driver state and are display-only "
    "for this tool's standalone processing path.",
)

doc.add_heading("Processing", level=2)
doc.add_paragraph(
    "Processing shows a progress bar, live file outcomes, and counts for files "
    "examined, images updated, and errors. Work runs in background workers; the "
    "normalizer uses up to four workers by default."
)

doc.add_heading("Results", level=2)
doc.add_paragraph(
    "Results shows total time, files examined, images updated, skipped files, and "
    "errors. View Log and Open Folder are available. The GUI writes "
    "art_normalize_log.txt in the selected directory."
)

doc.add_heading("Image Processing Rules", level=1)
doc.add_paragraph(
    "Artwork is converted to RGB, center-cropped to the largest possible square, "
    "resized with Lanczos resampling, and encoded as an optimized JPEG using the "
    "suite JPEG quality setting. The crop avoids stretching but removes pixels "
    "from the longer dimension of a non-square image."
)
bullet(doc, "Already at target size: ", "Skipped without rewriting.")
bullet(
    doc,
    "No embedded artwork: ",
    "Skipped because there is no artwork to update.",
)
bullet(
    doc,
    "Unreadable image or tag data: ",
    "Reported as an error while other files continue.",
)

doc.add_heading("folder.jpg", level=1)
doc.add_paragraph(
    "The GUI checks only the selected directory itself for folder.jpg. If found, "
    "it is read as RGB, center-cropped if needed, resized to 1000 x 1000, and "
    "saved in place as an optimized JPEG. A missing or already-correct folder.jpg "
    "is skipped. The CLI does not run this GUI-only step."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph(
    "The registered command is artnormalize. It processes embedded artwork in "
    "every supported audio format. Examples:"
)
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app artnormalize <folder>\n"
    "python -m audio_suite_shared.cli_app artnormalize <folder> --size 1000\n"
    "python -m audio_suite_shared.cli_app artnormalize <folder> --size 1200 --json"
)
doc.add_paragraph(
    "The --size option is an integer square dimension and defaults to 800. The "
    "--json option prints the result as formatted JSON; otherwise the command "
    "prints processed, skipped, error, and elapsed-time values."
)

doc.add_heading("Master Driver", level=1)
doc.add_paragraph(
    "In the Master Driver, open UTIL and select Album Art Normalizer. The card "
    "uses the shared directory and run controls and launches the tool with the "
    "selected path. The standalone GUI exposes all four resolution choices. A "
    "driver launch can inherit control values and auto-start behavior."
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Use Dry Run before a live run and keep backups of important audio files. "
    "Artwork is written back into the original audio files, while folder.jpg is "
    "written back as the directory image. The tool participates in the suite's "
    "batch undo mechanism for recorded tag edits; this does not replace backups."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "The run will not start: ",
    "Choose an existing directory and a valid resolution.",
)
bullet(
    doc,
    "Everything is skipped: ",
    "Files may already match the target or have no embedded artwork.",
)
bullet(
    doc,
    "The image is cropped: ",
    "Center-cropping is intentional for square output without distortion.",
)
bullet(
    doc,
    "A file errors: ",
    "Check whether its tags, artwork bytes, or file "
    "permissions allow reading and writing.",
)
bullet(
    doc,
    "folder.jpg is unchanged: ",
    "It may be missing, outside the selected directory, or already 1000 x 1000.",
)
bullet(
    doc,
    "The CLI rejects the folder: ",
    "Correct the path; CLI argument validation requires an existing directory.",
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
