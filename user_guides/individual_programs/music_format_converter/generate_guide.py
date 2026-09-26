"""Generate a fresh Music Format Converter guide from current code."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "music_format_converter\\Music_Format_Converter_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Music Format Converter - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph(
    "Based on the current Music Format Converter implementation"
)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Music Format Converter converts a directory of audio files to another format with "
    "FFmpeg. It preserves metadata through the conversion path, "
    "supports format-specific "
    "quality controls, can resample audio, "
    "can choose WAV bit depth, and optionally applies "
    "lossless MP3 loudness gain after conversion."
)
doc.add_paragraph(
    "The converter is designed for a single input format per folder. "
    "It scans the input "
    "directory before processing and filters the valid output formats for that input."
)

doc.add_heading("Requirements", level=1)
bullet(
    doc,
    "FFmpeg: ",
    "Required for conversion. The GUI checks availability before starting.",
)
bullet(
    doc,
    "Input directory: ",
    "Must contain supported audio files directly in the selected directory.",
)
bullet(
    doc,
    "Single input format: ",
    "Mixed-format folders are rejected so one conversion rule applies consistently.",
)

doc.add_heading("Supported Formats", level=1)
bullet(
    doc,
    "Lossy: ",
    "MP3, AAC, M4A, OGG, Opus, WMA, and MP4 are recognized as lossy formats.",
)
bullet(
    doc,
    "Lossless/uncompressed: ",
    "FLAC, WAV, ALAC, APE, AIFF, AIF, and WV are recognized as lossless "
    "input categories.",
)
bullet(
    doc,
    "APE output: ",
    "APE can be detected as input but is excluded as an output "
    "because the converter has no supported encoder.",
)
bullet(
    doc,
    "Conversion restrictions: ",
    "Same-format conversion is prohibited, "
    "and lossy input cannot be converted to a lossless output "
    "because quality cannot be recovered.",
)

doc.add_heading("System Requirements", level=1)
doc.add_paragraph(
    "This tool requires ffmpeg and ffprobe to be installed and on PATH. "
    "ffmpeg decodes and encodes the audio formats the suite reads and "
    "writes; ffprobe reads stream metadata (duration, codec, bitrate) "
    "for analysis tools. The suite shows a startup warning naming any "
    "missing binary and how to install it."
)
for _hint in (
    "Windows: winget install Gyan.FFmpeg",
    "macOS:   brew install ffmpeg",
    "Linux:   sudo apt install ffmpeg",
):
    doc.add_paragraph(_hint, style="List Bullet")

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Music_Format_Converter.pyw.",
    "Select an input directory containing one audio format.",
    "Select an output format from the formats allowed for that input.",
    "Choose quality, sample rate, bit depth where applicable, and an output folder.",
    "Leave Delete Originals disabled and review the Dry Run preview first.",
    "Click the conversion start button and monitor Processing.",
    "Review Results and the conversion log before using deletion or a large batch.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the input directory, output format, "
    "quality, sample rate, WAV bit depth, "
    "MP3 normalization, target LUFS, output directory, presets, and run controls. "
    "Changing the "
    "input folder triggers a scan "
    "and updates the detected-format label and output-format choices."
)
bullet(
    doc,
    "Input Directory: ",
    "The source folder to scan. "
    "The converter processes supported files in that directory.",
)
bullet(
    doc,
    "Output Format: ",
    "The selected target extension. The menu excludes the detected input format, "
    "invalid lossy-to-lossless targets, and unsupported APE output.",
)
bullet(doc, "Quality: ", "Format-specific codec quality or bitrate selection.")
bullet(
    doc,
    "Sample Rate: ",
    "Keep source or choose a supported rate such as "
    "44100, 48000, 88200, 96000, 176400, or 192000 Hz.",
)
bullet(
    doc,
    "Bit Depth: ",
    "Relevant to WAV output: Auto, 16, 24, 32, or float, "
    "with source-preserving behavior where supported.",
)
bullet(
    doc,
    "Output Directory: ",
    "Optional separate destination. If blank, output defaults to the input directory.",
)
bullet(
    doc,
    "Delete Originals: ",
    "Deletes the source only after a successful converted output exists, "
    "with confirmation in the GUI.",
)
bullet(
    doc,
    "Normalize MP3: ",
    "Optionally measures the converted MP3 "
    "and applies lossless track gain toward the target LUFS, default -16 LUFS.",
)

doc.add_heading("Quality Options", level=1)
bullet(doc, "MP3: ", "VBR V0 (highest), VBR V2 (high), or CBR 192 kbps (medium).")
bullet(doc, "FLAC: ", "Always compression level 8; no quality tiers apply.")
bullet(doc, "OGG: ", "Vorbis quality 10, 8, or 6.")
bullet(doc, "M4A: ", "AAC at 320, 256, or 192 kbps.")
bullet(doc, "MP4: ", "AAC at 320, 256, or 192 kbps (same tiers as M4A).")
bullet(doc, "Opus: ", "256, 192, or 128 kbps.")
bullet(doc, "AIFF / AIF: ", "24-bit PCM at highest/high, 16-bit PCM at medium.")
bullet(doc, "WMA: ", "320, 256, or 192 kbps.")

doc.add_heading("Conversion Behavior", level=1)
for step in (
    "Discover supported files in the input directory.",
    "Validate that input and output formats are compatible.",
    "Build FFmpeg parameters for the selected format, "
    "quality, sample rate, and bit depth.",
    "Convert each file with retry handling and write "
    "the output filename with the target extension.",
    "Preserve metadata through the conversion path.",
    "Optionally measure and losslessly adjust MP3 output toward the target LUFS.",
    "Optionally delete the original after output success "
    "and record the deletion snapshot.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Dry Run and Live Update", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Simulates each valid conversion, estimates input/output size, "
    "reports deletion and MP3-normalization actions, and writes no audio files.",
)
bullet(
    doc,
    "Live Update: ",
    "Runs FFmpeg conversion and writes outputs to the destination or input directory.",
)
bullet(
    doc,
    "Skipped: ",
    "Same-format conversions, lossy-to-lossless conversions, "
    "unsupported output targets, "
    "and configured recovery skips are reported as skipped.",
)
bullet(
    doc,
    "Errors: ",
    "FFmpeg failures, timeouts, permissions, and unexpected failures "
    "are reported as errors.",
)
bullet(
    doc,
    "Undo: ",
    "Original deletion is recorded in the suite file-delete undo context; "
    "conversion itself does not alter the source unless Delete Originals is enabled.",
)

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows four worker slots, progress, converted/skipped/error counts, "
    "and per-file "
    "status lines. Results reports totals and elapsed "
    "time. Converted lines include source and "
    "output sizes, optional loudness-gain status, and original-deletion status."
)

doc.add_heading("Logs", level=1)
doc.add_paragraph(
    "The converter writes its processing log in the resolved output directory. "
    "The log contains "
    "per-file conversion status, format transition, size information, "
    "skipped reasons, errors, "
    "and summary counts."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is convert. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app convert <folder> --format flac --dry\n"
    "python -m audio_suite_shared.cli_app convert <folder> "
    "--format mp3 --quality highest --output-dir <out>\n"
    "python -m audio_suite_shared.cli_app convert <folder> "
    "--format mp3 --normalize-mp3 --target-lufs -16\n"
    "python -m audio_suite_shared.cli_app convert <folder> "
    "--format m4a --delete-originals --live --json"
)
doc.add_paragraph(
    "CLI options include --format, --quality, --sample-rate, "
    "--bit-depth, --output-dir, "
    "--normalize-mp3, --target-lufs, "
    "--delete-originals/--keep-originals, --dry/--live, "
    "--preset, and --json. Saved presets can supply output and quality settings."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open UTIL and select Music Format Converter. "
    "The card detects input "
    "formats from the shared directory, "
    "passes the shared mode and orchestration state, and launches "
    "the converter workflow. Use the standalone Setup controls "
    "for quality, output, and deletion choices."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "Mixed formats detected: ",
    "Place one input format in the selected directory "
    "or separate the files before scanning.",
)
bullet(
    doc,
    "No output format available: ",
    "The input may have no supported files, may be APE-only, "
    "or every valid output may be prohibited by conversion rules.",
)
bullet(
    doc,
    "FFmpeg required: ",
    "Install FFmpeg and ensure the application can find it before starting.",
)
bullet(
    doc,
    "Conversion skipped: ",
    "Check whether input and output are the same format "
    "or whether a lossy-to-lossless conversion was requested.",
)
bullet(
    doc,
    "Originals were not deleted: ",
    "Deletion occurs only after output success and may fail separately; "
    "review the status and log.",
)
bullet(
    doc,
    "MP3 normalization failed: ",
    "Conversion can still succeed while gain status reports "
    "measurement or gain failure. Review the log and output file.",
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Use Dry Run first and choose a separate output directory while testing. "
    "Keep Delete Originals "
    "disabled until outputs are verified. Live deletion is recorded for "
    "undo support, but independent "
    "backups remain the safest recovery method."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
