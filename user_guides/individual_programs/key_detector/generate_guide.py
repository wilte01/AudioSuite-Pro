"""Generate a fresh Key Detector guide from the current implementation.

Changed in 1.4.0: the Workers description states the machine default
(one per logical CPU, capped at 8) instead of a fixed 4.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "key_detector\\Key_Detector_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Key Detector - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Key Detector implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Key Detector estimates the musical key of audio "
    "files from their signal using "
    "librosa chroma analysis and Krumhansl-Schmuckler "
    "key profiles. It can report the "
    "detected key without changing files or write the "
    "result to the standard musical "
    "key metadata field."
)
doc.add_paragraph(
    "The tool scans supported audio files in a selected "
    "directory, analyzes a configurable "
    "leading duration, and processes files with "
    "configurable worker parallelism."
)

doc.add_heading("Supported Audio", level=1)
doc.add_paragraph(
    "The detector uses the shared Audio Suite "
    "supported-extension list, including MP3, "
    "FLAC, OGG, M4A, Opus, WAV, and AIFF variants. "
    "The selected directory must contain "
    "supported audio files."
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Key_Detector.pyw.",
    "Select the music directory.",
    "Choose an analysis duration and optional parallel processing settings.",
    "Leave Dry Run selected for the first pass.",
    "Click BEGIN KEY DETECTION.",
    "Review detected keys, skipped files, errors, and the log.",
    "Run again in Live Update when you want the key tags written.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the directory picker, run mode, "
    "analysis duration, parallel worker "
    "controls, presets, and the start button. The "
    "last-used directory and parallel settings "
    "are restored through shared suite configuration."
)
bullet(
    doc,
    "Analysis Duration: ",
    "The GUI defaults to 90 seconds. Shorter analysis "
    "is faster but may provide less musical context.",
)
bullet(
    doc,
    "Parallel Processing: ",
    "When enabled, CPU-bound librosa chroma work uses "
    "ProcessPoolExecutor to bypass the Python GIL. "
    "When disabled, the thread-based path is used.",
)
bullet(
    doc,
    "Workers: ",
    "Controls the maximum worker count; the default is one "
    "worker per logical CPU, capped at 8.",
)
bullet(
    doc,
    "Presets: ",
    "Save and restore directory, mode, duration, "
    "parallel mode, and worker-count settings.",
)
bullet(
    doc,
    "Dry Run / Live Update: ",
    "Dry Run detects without writing. Live Update writes the detected key metadata.",
)

doc.add_heading("Detection Algorithm", level=1)
for step in (
    "Load up to the configured duration of audio with librosa.",
    "Convert the signal to mono analysis data and "
    "promote it to float64 for internal processing.",
    "Build a constant-Q chromagram and average it "
    "over time into a 12-element pitch-class profile.",
    "Correlate the profile against all rotations of "
    "the 12-note major and minor Krumhansl-Schmuckler "
    "templates.",
    "Select the strongest major or minor correlation "
    "and convert it to conventional display "
    "notation.",
):
    doc.add_paragraph(step, style="List Number")
doc.add_paragraph(
    "Major keys are displayed without a suffix, while "
    "minor keys end in m. Common flat "
    "spellings are used for Eb, Ab, Bb, and their "
    "supported minor equivalents; C# and F# "
    "remain sharp-spelled in the display conventions."
)

doc.add_heading("Silent and Existing-Key Behavior", level=1)
bullet(
    doc,
    "Silent audio: ",
    "Zero-energy audio cannot produce a meaningful "
    "chromagram correlation and is reported as "
    "an error.",
)
bullet(
    doc,
    "Dry Run with existing key: ",
    "A file that already has a key tag is skipped rather than re-analyzed in Dry Run.",
)
bullet(
    doc,
    "Live Update with existing key: ",
    "The file is analyzed and the tag is rewritten; "
    "the status includes the previous key when "
    "it differs.",
)
bullet(
    doc,
    "Detection failure: ",
    "The file is reported as an error or recoverable skip while other files continue.",
)
bullet(
    doc,
    "Material with no defensible key: ",
    "When the winning correlation does not lead the runner-up by a decisive "
    "margin -- broadband noise, percussive material, atonal music -- the file "
    "is reported as DETECTED (not tagged) and no key is written, even in Live "
    "Update. A file that already has a key tag keeps it. This is deliberate: "
    'Key Detector never writes the literal value "Unknown" into your library, '
    "because that reads as a musical fact it cannot support and would block a "
    "better result on a later run.",
)

doc.add_heading("Dry Run and Live Update", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Reports detected keys and skips files that "
    "already contain a key tag. No metadata "
    "is written.",
)
bullet(
    doc,
    "Live Update: ",
    "Writes the detected key to the standard key "
    "metadata field, including format-appropriate "
    "representations through the shared tag helper. Files whose key could not "
    "be determined are reported instead of written.",
)
bullet(doc, "Undo: ", "Live tag changes participate in the suite's batch undo context.")

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows four worker-slot statuses, a "
    "progress bar, processed/skipped/error "
    "counts, and per-file messages. Results reports "
    "total files, analyzed files, skipped "
    "files, errors, and elapsed time. The completion "
    "label distinguishes Analysis Complete "
    "from Key Tagging Complete."
)

doc.add_heading("Logs", level=1)
bullet(doc, "Dry Run log: ", "key_dry_run_log.txt in the selected directory.")
bullet(doc, "Live log: ", "key_sync_log.txt in the selected directory.")
bullet(
    doc,
    "Contents: ",
    "Mode, analysis duration, parallel settings, "
    "per-file detected/existing keys, skips, errors, "
    "counts, and elapsed time.",
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is key. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app key "
    "<folder> --dry\n"
    "python -m audio_suite_shared.cli_app key "
    "<folder> --duration 120\n"
    "python -m audio_suite_shared.cli_app key "
    "<folder> --parallel --workers 4\n"
    "python -m audio_suite_shared.cli_app key "
    "<folder> --live --json"
)
doc.add_paragraph(
    "The CLI defaults to a 30-second analysis duration, "
    "unlike the GUI's 90-second default. "
    "It accepts --duration, "
    "--parallel/--no-parallel, --workers, "
    "the shared dry-run/live option, and --json. "
    "The duration is validated from 10 through "
    "600 seconds."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open ENRICH and select "
    "Key Detector. The driver passes the directory, "
    "run mode, duration, worker settings, "
    "start/end behavior, and orchestration flags "
    "through the launch handshake. A driven launch "
    "can auto-start and displays inherited controls."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No audio files found: ",
    "Select a directory containing supported audio files.",
)
bullet(
    doc,
    "Keys are skipped: ",
    "In Dry Run, files that already have a key tag "
    "are intentionally skipped. Use Live Update "
    "to refresh them.",
)
bullet(
    doc,
    "Results are uncertain: ",
    "Increase analysis duration to provide more "
    "musical context, especially for tracks with "
    "long intros or varied sections.",
)
bullet(
    doc,
    "Silent file error: ",
    "Pure silence has no usable pitch-class profile and cannot yield a meaningful key.",
)
bullet(
    doc,
    "Parallel processing fails: ",
    "Disable parallel mode or reduce worker count if "
    "process startup or system resources are "
    "constrained.",
)
bullet(
    doc,
    "Tags did not change: ",
    "Confirm Live Update was selected; Dry Run never writes metadata.",
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Run Dry Run first and review the detected keys. "
    "Live Update modifies metadata in place, "
    "so keep backups of important files. Use the "
    "suite's Undo Tag Edits workflow for available "
    "tag snapshots."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
