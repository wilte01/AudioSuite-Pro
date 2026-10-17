"""Generate a fresh Audio Fingerprinter guide from the current implementation.

Changed in 1.2.0: documents the MusicBrainz enrichment, ISRC tag-write and
ISRC submission
options (Sprint 2 Item 2.1).
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "audio_fingerprinter\\Audio_Fingerprinter_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Audio Fingerprinter - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Audio Fingerprinter implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Audio Fingerprinter generates Chromaprint "
    "fingerprints with the external fpcalc utility "
    "and can identify tracks through AcoustID. It "
    "can run as a fingerprint-only analysis, identify "
    "likely recordings, or write AcoustID metadata "
    "to matched files. The selected directory is "
    "scanned recursively for supported audio files "
    "using four worker slots."
)

doc.add_heading("Prerequisites", level=1)
bullet(
    doc,
    "fpcalc: ",
    "Install Chromaprint fpcalc, add it to PATH, "
    "set FPCALC, or choose its executable in Setup.",
)
bullet(
    doc,
    "AcoustID application key: ",
    "Required when lookup is enabled. It can be "
    "stored encrypted in suite configuration or "
    "supplied through ACOUSTID_API_KEY or a one-run "
    "CLI option.",
)
bullet(
    doc,
    "Network access: ",
    "Lookup mode contacts AcoustID with a default "
    "30-second timeout and a one-request-per-second "
    "rate limit.",
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Audio_Fingerprinter.pyw.",
    "Select the directory containing audio files.",
    "Confirm that fpcalc is found and configure an "
    "AcoustID application key if lookup is enabled.",
    "Leave Dry Run selected and use the default minimum score of 0.5.",
    "Keep Write AcoustID tags disabled for the first pass.",
    "Click START FINGERPRINTING and review the Results tab and log.",
    "Use Live Update with tag writing enabled only after reviewing matches.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("GUI Setup", level=1)
doc.add_paragraph(
    "Setup contains the directory picker, fpcalc path, "
    "API key controls, lookup and tag-writing "
    "switches, score threshold, analysis length, and "
    "shared run controls. The last-used directory is "
    "restored from suite configuration."
)
bullet(
    doc,
    "fpcalc Binary: ",
    "Leave blank for FPCALC/PATH resolution or browse "
    "to an executable. The status line reports the "
    "resolved path.",
)
bullet(
    doc,
    "AcoustID API Key: ",
    "Enter a key and Save Key to encrypt it into "
    "suite configuration, or Clear to remove the "
    "stored key.",
)
bullet(
    doc,
    "Look up tracks: ",
    "Queries AcoustID after fingerprint generation; "
    "enabled by default and requires a key.",
)
bullet(
    doc,
    "Write AcoustID tags: ",
    "Writes match metadata only for a match in Live Update mode; disabled by default.",
)
bullet(
    doc,
    "Write MusicBrainz identifiers: ",
    "Adds MUSICBRAINZ_TRACKID, MUSICBRAINZ_ARTISTID "
    "and (when one release is unambiguous) "
    "MUSICBRAINZ_ALBUMID to the AcoustID tags. Needs "
    "the option above and Live Update; disabled by "
    "default. The album id is never guessed - a "
    "recording is typically released many times, and "
    "the id is written only when exactly one release "
    "survives the selection criteria below.",
)
bullet(
    doc,
    "MusicBrainz enrichment: ",
    "Queries MusicBrainz directly for the matched "
    "recording: its ISRCs and its release's date, "
    "label and country. Off by default, because "
    "MusicBrainz allows one request per second and "
    "an album therefore costs one request per unique "
    "recording plus one per unique release. A failed "
    "lookup never fails a file: the file keeps its "
    "AcoustID result and the message names the "
    "failure.",
)
bullet(
    doc,
    "Write ISRCs: ",
    "Writes the retrieved ISRCs in each format's "
    "native convention - the ID3 TSRC frame for "
    "MP3/WAV/AIFF/DSF, the "
    "----:com.apple.iTunes:ISRC freeform atom for "
    "M4A/MP4, and an ISRC comment/item for "
    "FLAC/OGG/Opus/WavPack/APE/Musepack. Needs the "
    "enrichment option above, Write AcoustID tags, "
    "and Live Update; disabled by default. A lookup "
    "that returns no ISRCs never blanks an ISRC the "
    "file already had.",
)
bullet(
    doc,
    "Submit ISRCs to MusicBrainz: ",
    "Offers the ISRCs each file's own tags carry to "
    "MusicBrainz for its matched recording. Needs "
    "MusicBrainz enrichment, because that lookup is "
    "what tells the run which ISRCs are already "
    "registered, so only new codes are sent. Disabled "
    "by default. In Dry Run the exact XML payload is "
    "written to the log and nothing is posted; a Live "
    "Update needs a token.",
)
bullet(
    doc,
    "MusicBrainz Token: ",
    "An OAuth access token with the submit_isrc "
    "scope, stored encrypted in the suite config. The "
    "MUSICBRAINZ_TOKEN environment variable overrides "
    "the stored value. The token is never accepted as "
    "a command-line argument, because argv is visible "
    "in process listings and shell history. The "
    "legacy application-token alternative (HTTP Digest) "
    "is CLI-only: --mb-auth app-token --mb-user <name>.",
)
bullet(
    doc,
    "MBID Region: ",
    "Optional two-letter country code (ISO 3166-1 "
    "alpha-2, e.g. GB) stored in suite configuration. "
    "It only breaks ties between otherwise identical "
    "releases (same track and disc counts, same "
    "release group); it can never select a release "
    "on its own.",
)
bullet(
    doc,
    "Album Match: ",
    "Clusters each folder's lookups by MusicBrainz release "
    "group and tags only when the album is decided. A "
    "release group forms a cluster when at least three "
    "files (the --album-min default) carry candidates in "
    "it; the cluster is decided when its score (files "
    "voting times average confidence) leads the runner-up "
    "by 25 percent. Well-known albums often tie exactly "
    "(one test album produced eight clusters at the same "
    "score, because a recording appears on every release "
    "group containing it); when the files share one album "
    "tag and it matches exactly one tied cluster, that "
    "cluster is decided - the user's own tag is the "
    "evidence, and a match against several clusters still "
    "refuses to guess. Ambiguous albums are left untagged "
    "with the whole tied set named in the log. Files no "
    "decided cluster claims fall back to per-file matching. "
    "Disabled by default.",
)
bullet(
    doc,
    "Min Score: ",
    "Accepts candidates from 0.0 to 1.0; default 0.5. Higher values are stricter.",
)
bullet(
    doc,
    "Max Length: ",
    "Maximum seconds sent to fpcalc, from 1 to 600; "
    "default 120. fpcalc still reports full file "
    "duration for lookup checks.",
)
bullet(
    doc,
    "Dry Run / Live Update: ",
    "Dry Run prevents tag writes. Live Update is required for tag writing.",
)

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows total and processed counts, "
    "errors, per-file status lines, and four "
    "worker-slot labels. Each file is fingerprinted, "
    "optionally looked up, and optionally tagged. "
    "Results reports identified matches, no matches, "
    "tags written, lookup skips, errors, and "
    "elapsed time."
)
doc.add_paragraph(
    "The GUI writes fingerprint_dry_run_log.txt or "
    "fingerprint_live_log.txt in the input directory. "
    "Results are sorted by file path when displayed."
)
doc.add_paragraph(
    "In Album Match mode the log starts with one "
    "[ALBUM] line per folder summarising the cluster "
    "decision, files decided by an album carry an "
    "'album consensus' note in their message, and the "
    "summary adds an Album Ambiguous count when any "
    "folder tied."
)

doc.add_heading("Fingerprinting Pipeline", level=1)
for step in (
    "Discover supported audio files recursively below the selected directory.",
    "Resolve fpcalc using the explicit path, FPCALC, or PATH.",
    "Run fpcalc with the selected maximum length and parse DURATION and FINGERPRINT.",
    "Skip lookup when disabled or when no key is available.",
    "Send the fingerprint and duration to AcoustID when lookup is enabled.",
    "Filter candidates by score and duration "
    "consistency, then choose the best recording.",
    "Optionally write Acoustid Fingerprint and Acoustid Id in Live Update mode.",
):
    doc.add_paragraph(step, style="List Number")
doc.add_paragraph(
    "Album Match mode runs the same pipeline in two "
    "phases: every file is fingerprinted first (the live "
    "log shows ANALYSED lines), then all lookups run "
    "through the rate-limited batch client, each folder "
    "is clustered by MusicBrainz release group, and only "
    "decided albums are tagged."
)

doc.add_heading("AcoustID Matching", level=1)
bullet(
    doc,
    "Score threshold: ",
    "Candidates below the selected score are discarded; scores range from 0.0 to 1.0.",
)
bullet(
    doc,
    "Duration guard: ",
    "Candidates must have a duration consistent with "
    "the file using the service's seven-second "
    "threshold plus a five-percent allowance for "
    "long tracks.",
)
bullet(
    doc,
    "Candidate limit: ",
    "Up to ten distinct recording candidates are retained per lookup.",
)
bullet(
    doc,
    "Near ties: ",
    "When top scores are within 0.10, existing artist "
    "and title tags resolve the tie when they agree. "
    "With MusicBrainz enrichment on, the remaining "
    "ambiguity is offered to MusicBrainz: a recording "
    "whose title and artist match the file's tags "
    "wins, then one that appears on a release matching "
    "the file's album tag, then one that has ISRCs "
    "registered. MusicBrainz never overrides a tie it "
    "cannot separate - rank order is kept and the "
    "ambiguity and runner-up are reported.",
)
bullet(
    doc,
    "No match: ",
    "A file receives NO MATCH when no candidate "
    "remains above the score and duration guardrails.",
)
bullet(
    doc,
    "Album clustering: ",
    "In Album Match mode each folder's candidates are "
    "grouped by MusicBrainz release group. A group with "
    "at least the album-minimum of voting files forms a "
    "cluster scored as (files voting) times (average "
    "confidence). A cluster is decided only when it "
    "leads the runner-up by 25 percent; otherwise the "
    "album is reported ambiguous and its files are left "
    "untagged with both clusters named. A decided "
    "cluster pulls each voting file to its best "
    "candidate inside the winning release group, which "
    "resolves remix-versus-original near ties that "
    "per-file scores cannot. Files outside the decided "
    "cluster keep the per-file matching described "
    "above, including the tag-aware near-tie break.",
)

doc.add_heading("Tags and Key Storage", level=1)
doc.add_paragraph(
    "When a match is found and tag writing is enabled "
    "in Live Update mode, the shared writer records "
    "the Chromaprint as Acoustid Fingerprint and the "
    "selected service identifier as Acoustid Id. Tag "
    "changes participate in the suite batch undo "
    "system."
)
doc.add_paragraph(
    "When Write MusicBrainz identifiers is also "
    "enabled, the run writes the matched recording as "
    "MUSICBRAINZ_TRACKID, the credited artist as "
    "MUSICBRAINZ_ARTISTID, and - only when exactly one "
    "release survives - the release as "
    "MUSICBRAINZ_ALBUMID. The identifiers are stored "
    "in each format's native convention: TXXX frames "
    "for MP3/WAV/AIFF, Vorbis comments for "
    "FLAC/OGG/Opus, and iTunes freeform atoms for "
    "M4A/MP4."
)
doc.add_paragraph(
    "The album id uses the releases already reported "
    "by AcoustID - no separate MusicBrainz lookup is "
    "made. Candidates are narrowed in order by an "
    "existing MUSICBRAINZ_ALBUMID on the file, the "
    "album title, the folder's track count and disc "
    "count (Album Match folders provide these), the "
    "release-group type, and finally the MBID Region "
    "preference. A criterion that would eliminate "
    "every candidate is ignored, and if more than one "
    "release still survives the field is simply not "
    "written. In per-file runs the track and disc "
    "counts are usually unknown, so the album id "
    "resolves less often than the track and artist "
    "ids."
)
doc.add_paragraph(
    "The effective API key is resolved from "
    "ACOUSTID_API_KEY first, then the encrypted "
    "acoustid_api_key value in shared suite "
    "configuration. A CLI --api-key value is used "
    "for that run. Do not place keys in scripts or "
    "logs."
)

doc.add_heading("MusicBrainz Enrichment and ISRCs", level=1)
doc.add_paragraph(
    "Enrichment is an addition to the AcoustID result, "
    "never a replacement for it. Each run creates one "
    "rate-limited MusicBrainz client with a per-entity "
    "cache, so an album of twelve tracks on one release "
    "costs twelve recording requests and one release "
    "request - not twenty-four. Requests are spaced to "
    "MusicBrainz's documented one per second limit, so "
    "an album run takes noticeably longer with "
    "enrichment enabled."
)
doc.add_paragraph(
    "ISRCs are validated and normalised before the "
    "write: values are upper-cased (ISO 3901 ISRCs are "
    "case-insensitive) and duplicates dropped. The "
    "write is additive to the AcoustID tags: it happens "
    "only after those tags land successfully, only in "
    "Live Update, and only when Write AcoustID tags and "
    "Write ISRCs are both enabled. Every write is "
    "recorded in the batch undo log with the other tag "
    "edits."
)
doc.add_paragraph(
    "Release-level facts (date, label, country) are "
    "reported in the run log and the JSON output; they "
    "are not written as tags in this release. The facts "
    "and any MUSICBRAINZ_ALBUMID always describe the "
    "same release, because both use the same selection "
    "rules."
)

doc.add_heading("Submitting ISRCs to MusicBrainz", level=1)
doc.add_paragraph(
    "Submission is the reverse direction of enrichment: "
    "it uploads ISRCs the files already carry, which "
    "usually means the codes a CD rip wrote, to the "
    "recording AcoustID matched. Codes MusicBrainz "
    "already has for that recording, and codes another "
    "file in the same batch already contributed, are "
    "dropped, so the payload contains only work the "
    "service can act on. A file whose tags carry an "
    "ISRC MusicBrainz considers invalid is reported "
    "and skipped; a bad code never blocks the rest of "
    "the run."
)
doc.add_paragraph(
    "A submission is an edit in your account, not a "
    "write: MusicBrainz keeps new edits open for review "
    "for seven days, and the account's email address "
    "must be verified. Always run Dry Run first. The "
    "payload shown there is generated by the same "
    "function that builds the live request, and the "
    "codes and recordings are sorted, so the document "
    "you review is byte-for-byte what a Live Update "
    "posts."
)
doc.add_paragraph(
    "The request goes to the recording endpoint with "
    "the mandatory client parameter, as "
    "application/xml; charset=UTF-8. Authentication is "
    "an OAuth bearer token with the submit_isrc scope "
    "or, for accounts that generated one, a legacy "
    "application token over HTTP Digest. Digest is "
    "deprecated by MusicBrainz and is scheduled for "
    "removal; prefer OAuth."
)
doc.add_paragraph(
    "A refusal is reported, never retried: HTTP 401/403 "
    "means the token lacks the scope, expired, or the "
    "account's email is unverified; HTTP 400 means the "
    "service rejected the payload as invalid; a network "
    "or throttling failure leaves nothing submitted. "
    "No submission outcome changes a file's result "
    "status, because the files are already tagged."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is fingerprint. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --no-lookup\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --threshold 0.7 --max-length 120\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --write-tags --live --api-key <key>\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --album-mode --album-min 3 --write-tags --live\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --write-tags --write-mbids --mbid-region GB --live\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --musicbrainz --write-tags --live\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --musicbrainz --write-tags --write-isrc --live\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --musicbrainz --submit-isrc --dry\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --musicbrainz --submit-isrc --live\n"
    "python -m audio_suite_shared.cli_app fingerprint "
    "<folder> --fpcalc <path> --json"
)
doc.add_paragraph(
    "Options include --lookup/--no-lookup, "
    "--write-tags/--no-write-tags, "
    "--write-mbids/--no-write-mbids (needs "
    "--write-tags), --mbid-region (two-letter release "
    "country tie-break), --musicbrainz/--no-musicbrainz "
    "(enrich matches from the MusicBrainz database), "
    "--write-isrc/--no-write-isrc (needs --musicbrainz "
    "and --write-tags), --submit-isrc/--no-submit-isrc "
    "(needs --musicbrainz; previews the payload in Dry "
    "Run, posts in Live Update), --mb-auth "
    "(oauth or app-token) and --mb-user (needed only by "
    "app-token), --threshold, "
    "--max-length, --fpcalc, --api-key, "
    "--album-mode/--no-album-mode (cluster each folder "
    "by release group; tag only decided albums), "
    "--album-min (minimum voting files per release "
    "group, default 3, at least 1), --dry, "
    "and --json. Lookup requires an API key. Text "
    "output reports files, identified matches, tagged "
    "files, errors, and elapsed time."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open ENRICH and select "
    "Audio Fingerprinter. The driver passes the "
    "directory, mode, start/end behavior, and "
    "fingerprint options through the launch handshake. "
    "A driven launch can auto-start and displays "
    "inherited controls."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "fpcalc required: ",
    "Install Chromaprint, set FPCALC, add fpcalc to PATH, or browse to the executable.",
)
bullet(
    doc,
    "API key required: ",
    "Configure encrypted storage, set "
    "ACOUSTID_API_KEY, or provide --api-key when "
    "lookup is enabled.",
)
bullet(
    doc,
    "No matches: ",
    "Check the threshold, increase analysis length "
    "when appropriate, and confirm the recording "
    "is represented by AcoustID.",
)
bullet(
    doc,
    "Wrong match: ",
    "Review score and existing artist/title tags. "
    "Near ties are tag-resolved when possible and "
    "otherwise marked ambiguous.",
)
bullet(
    doc,
    "Tags not written: ",
    "Confirm Live Update and Write AcoustID tags are "
    "enabled, and that files are writable. Dry Run "
    "never writes tags.",
)
bullet(
    doc,
    "MusicBrainz album id missing: ",
    "The matched recording spans several "
    "indistinguishable releases, so the run omits "
    "MUSICBRAINZ_ALBUMID rather than guess. Set an "
    "MBID Region, use Album Match so the folder's "
    "track and disc counts can narrow the candidates, "
    "or accept that only the track and artist ids "
    "were written.",
)
bullet(
    doc,
    "Album files left untagged (ALBUM AMBIGUOUS): ",
    "The folder's release-group clusters tied inside "
    "the decision margin, so album mode refuses to "
    "guess; the log names every tied cluster and its "
    "scores, and the files' shared album tag already "
    "decides the tie when it matches exactly one of "
    "them. Otherwise lower --album-min, fix the "
    "album tag, or resolve the ambiguity manually; "
    "a run without Album Match still applies per-file "
    "matching.",
)
bullet(
    doc,
    "MusicBrainz token required: ",
    "A Live Update with Submit ISRCs needs a token: "
    "store one on the Setup tab, or set "
    "MUSICBRAINZ_TOKEN. A Dry Run needs no token - it "
    "only shows the payload. The token is never read "
    "from argv.",
)
bullet(
    doc,
    "ISRC submission NOT AUTHORISED: ",
    "MusicBrainz refused the edit. Check that the "
    "token carries the submit_isrc scope, has not "
    "expired, and that the account's email address is "
    "verified. Nothing was submitted and nothing is "
    "retried automatically.",
)
bullet(
    doc,
    "ISRC submission REJECTED: ",
    "The service rejected the payload as invalid. "
    "Review the preview: codes must be 12 characters, "
    "CC-XXX-YY-NNNNN, and the recording ids must be "
    "current (a merged recording needs its new id).",
)
bullet(
    doc,
    "ISRC submission skipped: ",
    "The log names each file that contributed "
    "nothing - because its codes were already "
    "registered, its tags carried no ISRC, or the "
    "code was not a valid ISRC. Nothing is guessed "
    "and nothing is dropped silently.",
)
bullet(
    doc,
    "A file errors: ",
    "Check codec support, fpcalc output, permissions, "
    "network access, and the per-file log message.",
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Start with Dry Run and tag writing disabled. "
    "Fingerprinting does not alter audio files, but "
    "Live Update modifies metadata in place, and ISRC "
    "submission creates a public edit in your "
    "MusicBrainz account. Keep backups and use the "
    "suite's Undo Tag Edits workflow for available tag "
    "snapshots; a pending MusicBrainz edit is withdrawn "
    "on the MusicBrainz site, not from here."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
