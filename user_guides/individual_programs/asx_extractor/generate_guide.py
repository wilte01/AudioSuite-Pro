"""Generate a fresh ASX Playlist Extractor guide from the current code."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "asx_extractor\\ASX_Playlist_Extractor_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("ASX Playlist Extractor - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph(
    "Based on the current ASX Playlist Extractor implementation"
)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "ASX Playlist Extractor processes Windows Media "
    "ASX playlist files in a selected directory. "
    "It reads each playlist's first usable ref link, "
    "resolves the linked audio file, verifies that "
    "the target exists, and optionally copies that "
    "audio file into the selected directory. In "
    "permanent mode it then removes the ASX playlist, "
    "leaving the audio file as the extracted result."
)
doc.add_paragraph(
    "The tool is deliberately staged: it does not "
    "remove an ASX playlist when its audio copy "
    "fails. This keeps the playlist available when "
    "an extraction cannot be completed."
)

doc.add_heading("Input and Output", level=1)
bullet(
    doc,
    "Input: ",
    "Top-level files whose extension is .asx, matched case-insensitively.",
)
bullet(
    doc,
    "Playlist reference: ",
    "The first <ref> XML element with a non-empty href attribute.",
)
bullet(
    doc,
    "Audio result: ",
    "The referenced audio file is copied into the "
    "selected directory with metadata preserved "
    "by copy2.",
)
bullet(
    doc,
    "Playlist cleanup: ",
    "In permanent mode, the ASX file is sent to the "
    "system Trash when send2trash is available; "
    "otherwise it is deleted permanently.",
)
bullet(
    doc,
    "Log: ",
    "Extraction_Log.txt is written in the selected "
    "directory with failures, successful extractions, "
    "mode, date, and summary counts.",
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open ASX_Playlist_Extractor.pyw from the AudioSuite-Pro folder.",
    "On Setup, select the directory containing the ASX playlists.",
    "Choose Dry Run (Simulated) first to inspect playlist parsing and targets.",
    "Click START EXTRACTION.",
    "Review the Processing steps and then open the extraction log from Results.",
    "Use Execute Extraction (Permanent) only after confirming the dry-run results.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("GUI Tabs", level=1)
doc.add_heading("Setup", level=2)
doc.add_paragraph(
    "Setup contains the target directory, action mode, "
    "shared run controls, presets, and the start "
    "button. The last-used directory can be restored "
    "from shared suite configuration. A Master Driver "
    "launch can provide the directory and inherited "
    "control values through the launcher handshake."
)
bullet(
    doc,
    "Dry Run (Simulated): ",
    "Parses and verifies playlists and writes the "
    "report, but does not copy audio or remove "
    "ASX files.",
)
bullet(
    doc,
    "Execute Extraction (Permanent): ",
    "Copies successful audio targets into the "
    "selected directory and removes only playlists "
    "whose copies and cleanup succeed.",
)
bullet(
    doc,
    "Presets: ",
    "Save, load, and delete the current action mode "
    "through the shared suite preset storage.",
)
bullet(
    doc,
    "Auto Start: ",
    "Begins extraction automatically after a driven launch when enabled.",
)
bullet(
    doc,
    "Auto End: ",
    "Closes the window about two seconds after the run completes when enabled.",
)
bullet(
    doc,
    "Clear Logs / Parallel: ",
    "Mirror Master Driver orchestration state and "
    "are display-only for this extractor's processing "
    "path.",
)

doc.add_heading("Processing", level=2)
doc.add_paragraph(
    "The Processing tab presents six stages: scanning "
    "for ASX files, parsing ASX contents, verifying "
    "source audio and paths, copying audio, removing "
    "or trashing original ASX files, and finalizing "
    "the priority error log. Active stages show an "
    "arrow; completed stages show a check."
)
doc.add_paragraph(
    "Each playlist is classified as a success or "
    "failure during parsing and verification. "
    "Permanent extraction may reclassify a parsed "
    "success as a failure if copying or ASX removal "
    "fails."
)

doc.add_heading("Results", level=2)
doc.add_paragraph(
    "Results provides View Log and Open Folder. View "
    "Log becomes available after Extraction_Log.txt "
    "is written. Open Folder opens the selected "
    "source/output directory because extracted audio "
    "is copied there."
)

doc.add_heading("ASX Parsing Rules", level=1)
doc.add_paragraph(
    "The parser reads playlist text as XML using "
    "ElementTree. An XML declaration is removed "
    "before parsing when present, which accommodates "
    "declarations whose stated encoding does not "
    "match the actual text. The first <ref> with an "
    "href is used. Playlists that declare a DTD or "
    "entity (DOCTYPE/ENTITY) are rejected outright to "
    "block XML entity-expansion attacks."
)
bullet(
    doc,
    "Relative href: ",
    "Resolved relative to the directory containing the ASX file.",
)
bullet(doc, "Absolute href: ", "Resolved as an absolute filesystem path.")
bullet(
    doc,
    "Alternate-extension fallback: ",
    "If the referenced audio file is missing, the parser "
    "tries the first existing same-stem sibling file in "
    "any of the nine suite supported audio extensions "
    "(.mp3, .flac, .ogg, .m4a, .mp4, .opus, .wav, .aiff, "
    ".aif), preferring the shared "
    "SUPPORTED_EXTENSIONS order.",
)
bullet(doc, "Malformed XML: ", "Reported as INVALID XML IN ASX.")
bullet(
    doc,
    "DTD/entity payload: ",
    "Reported as DTD-BEARING ASX REJECTED before any parsing occurs.",
)
bullet(doc, "No reference: ", "Reported as NO REF LINK IN ASX.")
bullet(
    doc,
    "Missing target: ",
    "Reported as FILE NOT FOUND or DIRECTORY NOT FOUND depending on the resolved path.",
)

doc.add_heading("Dry Run and Permanent Extraction", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Runs scan, parse, and verification, records what "
    "was found, writes the report, and leaves both "
    "source audio and ASX playlists unchanged.",
)
bullet(
    doc,
    "Permanent mode: ",
    "Copies each verified audio file with metadata "
    "preserved, then removes the corresponding ASX "
    "only after the copy succeeds.",
)
bullet(
    doc,
    "Copy failure: ",
    "The ASX is retained and the item is recorded as COPY FAILED.",
)
bullet(
    doc,
    "Removal failure: ",
    "The audio copy remains, but the item is recorded "
    "as AUDIO COPIED BUT ASX REMOVAL FAILED.",
)
bullet(
    doc,
    "Undo support: ",
    "Successful copies are recorded in the suite "
    "file-reorganization undo context. Keep "
    "independent backups when permanently changing "
    "a library.",
)

doc.add_heading("Extraction Log", level=1)
doc.add_paragraph(
    "Extraction_Log.txt contains the processed "
    "directory, timestamp, selected mode, a failures "
    "section, a successful extractions section, and "
    "a final count summary. Failure records include "
    "the ASX filename, reason, and attempted audio "
    "path. Successful records include the ASX "
    "filename and resolved audio path."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is asx-extract. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app asx-extract "
    "<folder> --dry\n"
    "python -m audio_suite_shared.cli_app asx-extract "
    "<folder>\n"
    "python -m audio_suite_shared.cli_app asx-extract "
    "<folder> --json"
)
doc.add_paragraph(
    "The CLI accepts --dry and --json. Without --dry, "
    "extraction is permanent. The JSON option prints "
    "the structured result; normal text output "
    "reports total ASX files, extracted successes, "
    "failures, and errors. The directory must exist "
    "and be a directory."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open UTIL and select ASX "
    "Playlist Extractor. The driver passes the shared "
    "directory, action mode, Auto Start, Auto End, "
    "Clear Logs, and Parallel handshake values. The "
    "child window displays inherited controls and "
    "can start automatically when requested."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No ASX files found: ",
    "Confirm that the selected directory contains "
    "top-level .asx files; subdirectories are "
    "not scanned.",
)
bullet(
    doc,
    "Invalid XML: ",
    "Open the ASX file and check that it is "
    "well-formed XML with a usable "
    '<ref href="..."> element.',
)
bullet(
    doc,
    "File not found: ",
    "Check the href path relative to the playlist. "
    "The parser also tries the first same-stem sibling "
    "in a supported audio extension when applicable.",
)
bullet(
    doc,
    "Copy failed: ",
    "Check destination permissions, file locks, "
    "available space, and whether the source audio "
    "is readable.",
)
bullet(
    doc,
    "ASX remains after extraction: ",
    "The audio copy or playlist removal failed; "
    "review the failure entry in "
    "Extraction_Log.txt.",
)
bullet(
    doc,
    "The run did not change files: ",
    "Confirm that Dry Run (Simulated) was not selected.",
)

doc.add_heading("Safety Recommendations", level=1)
doc.add_paragraph(
    "Always run Dry Run first and inspect the log. "
    "Permanent mode changes the directory by copying "
    "audio and removing ASX playlists. Although "
    "send2trash normally provides a recovery path for "
    "removed playlists, do not rely on Trash or undo "
    "records as a substitute for a backup."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
