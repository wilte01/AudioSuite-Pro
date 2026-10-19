"""Generate a fresh Mood Analyzer guide from current code.

Changed in 1.4.0: the Workers description states the machine default
(one per logical CPU, capped at 8) instead of a fixed 4.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "mood_analyzer\\Mood_Analyzer_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Mood Analyzer - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Mood Analyzer implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Mood Analyzer estimates musical mood directly "
    "from the audio signal. It analyzes "
    "a selected directory, derives arousal and valence "
    "features, maps them to three mood "
    "words, and can write mood, confidence, and energy "
    "metadata. It does not rely on "
    "pre-existing mood tags."
)
doc.add_paragraph(
    "The standalone tool processes one directory. For "
    "recursive whole-library processing, "
    "use Mood Analysis Batch Tagger. Both tools use "
    "the same underlying mood analysis engine."
)

doc.add_heading("Supported Files", level=1)
doc.add_paragraph(
    "The analyzer uses the shared Audio Suite supported "
    "audio extensions, including MP3, "
    "FLAC, OGG, M4A, MP4, Opus, WAV, and AIFF "
    "variants. Subfolders are not descended into "
    "by this standalone tool."
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Mood_Analyzer.pyw.",
    "Select the directory containing the audio files.",
    "Choose an analysis duration and optional parallel settings.",
    "Leave Dry Run selected for the first pass.",
    "Click the analysis start button.",
    "Review the per-file results and mode-specific log.",
    "Run again in Live Update when you want mood tags written.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the directory picker, run mode, "
    "analysis duration, parallel worker "
    "controls, presets, and shared run controls. The "
    "last-used directory and parallel settings "
    "are stored in shared suite configuration."
)
bullet(
    doc,
    "Analysis Duration: ",
    "The GUI defaults to 90 seconds. The CLI accepts "
    "10 through 600 seconds. Shorter values are "
    "faster but sample less of each track.",
)
bullet(
    doc,
    "Parallel Processing: ",
    "When enabled, CPU-bound mood analysis uses "
    "ProcessPoolExecutor; otherwise the thread-based "
    "path is used.",
)
bullet(
    doc,
    "Workers: ",
    "Sets the maximum worker count; the default is one "
    "worker per logical CPU, capped at 8.",
)
bullet(
    doc,
    "Presets: ",
    "Save and restore directory, mode, duration, parallel mode, and worker count.",
)
bullet(
    doc,
    "Dry Run / Live Update: ",
    "Dry Run reports without writing. Live Update writes the derived values to files.",
)

doc.add_heading("Mood Model", level=1)
doc.add_paragraph(
    "The signal mood engine scores two axes. Arousal "
    "represents energy versus calm and uses "
    "tempo, spectral flux, zero-crossing rate, and "
    "RMS energy. Valence represents positive "
    "versus dark character and uses musical mode, "
    "key-detection strength, spectral brightness, "
    "and low-band energy. The resulting point is "
    "mapped to three descriptive words."
)
bullet(
    doc,
    "Mood words: ",
    "Stored as a semicolon-separated string such as Energetic;Bright;Confident.",
)
bullet(
    doc,
    "Confidence: ",
    "A decimal from 0 to 1 describing how strongly the analysis matched.",
)
bullet(doc, "Energy: ", "An integer band from 1, calm, through 6, very energetic.")

doc.add_heading("Processing Behavior", level=1)
for step in (
    "Discover supported audio files directly inside the selected directory.",
    "Load each file for the configured analysis duration.",
    "Analyze the signal and derive mood, confidence, and energy.",
    "In Dry Run, report the values without writing metadata.",
    "In Live Update, write the mood, confidence, and energy tags.",
    "Write the mode-specific report and show the final summary.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Dry Run and Live Update", level=1)
bullet(doc, "Dry Run: ", "Analyzes and reports each file without changing it.")
bullet(
    doc,
    "Live Update: ",
    "Writes the three mood-related values through the shared tag helpers.",
)
bullet(
    doc,
    "Undecodable or silent files: ",
    "They are skipped and not written; analysis "
    "failures are reflected in the summary and log.",
)
bullet(
    doc,
    "Errors: ",
    "Unexpected failures are counted separately from ordinary skipped analysis files.",
)
bullet(
    doc,
    "Undo: ",
    "Live tag writes run inside the suite tag-edit "
    "batch and can be reverted when snapshots "
    "are available.",
)

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows four worker slots, progress, "
    "processed/skipped/error counts, and per-file "
    "status messages. Results reports total files, "
    "analyzed/tagged files, skipped files, errors, "
    "and elapsed time. The tool can open the source "
    "folder and the generated log."
)

doc.add_heading("Logs", level=1)
bullet(doc, "Dry Run log: ", "mood_dry_run_log.txt in the selected directory.")
bullet(doc, "Live log: ", "mood_sync_log.txt in the selected directory.")
bullet(
    doc,
    "Contents: ",
    "Mode, analysis duration, parallel settings, "
    "per-file mood results, skips, errors, counts, "
    "and elapsed time.",
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is mood-analyzer. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app "
    "mood-analyzer <folder> --dry\n"
    "python -m audio_suite_shared.cli_app "
    "mood-analyzer <folder> --duration 60\n"
    "python -m audio_suite_shared.cli_app "
    "mood-analyzer <folder> --parallel --workers 4\n"
    "python -m audio_suite_shared.cli_app "
    "mood-analyzer <folder> --live --json"
)
doc.add_paragraph(
    "The CLI accepts --duration, "
    "--parallel/--no-parallel, --workers, "
    "the shared dry-run/live option, and --json. "
    "Its text summary reports processed, skipped, "
    "errors, and elapsed time."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open ENRICH and select "
    "Mood Analyzer. The driver passes the directory, "
    "mode, duration, worker settings, start/end "
    "behavior, and orchestration flags through the "
    "launch handshake. A driven launch can auto-start "
    "and displays inherited controls."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No supported files found: ",
    "Select a directory containing supported audio files directly in that folder.",
)
bullet(
    doc,
    "A file is skipped: ",
    "It may be undecodable, unsupported by the active decoder, or effectively silent.",
)
bullet(
    doc,
    "The mood seems subjective: ",
    "The result is signal-derived from measurable "
    "audio features, not lyrics, genre tags, or "
    "personal judgment.",
)
bullet(
    doc,
    "Results differ after changing duration: ",
    "The analyzer samples a different portion of the "
    "track; longer durations can provide more context "
    "but take longer.",
)
bullet(
    doc,
    "Parallel processing fails: ",
    "Disable parallel mode or reduce worker count if "
    "process startup or resources are constrained.",
)
bullet(
    doc,
    "Tags did not change: ",
    "Confirm Live Update was selected. Dry Run never writes metadata.",
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Run Dry Run first and review the log. Live Update "
    "modifies metadata in place, so keep "
    "backups of important files and use the suite's "
    "Undo Tag Edits workflow for available snapshots."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
