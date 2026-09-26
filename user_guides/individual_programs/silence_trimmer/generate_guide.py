"""Generate a fresh Silence Trimmer guide from current code."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "silence_trimmer\\Silence_Trimmer_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Silence Trimmer - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Silence Trimmer implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Silence Trimmer detects and removes leading and trailing "
    "silence from supported audio "
    "files. It can optionally apply fades after trimming, "
    "preserves existing tags, and "
    "processes files through temporary worker copies before "
    "replacing the source in Live "
    "Update."
)
doc.add_paragraph(
    "The tool is read-only in Dry Run and reports the detected "
    "trim amount without exporting "
    "or modifying files."
)

doc.add_heading("Supported Files", level=1)
doc.add_paragraph(
    "The processor uses the shared Audio Suite supported "
    "extensions, including MP3, FLAC, OGG"
    ", M4A, MP4, Opus, WAV, and AIFF variants. The selected "
    "directory must contain supported "
    "files."
)

doc.add_heading("System Requirements", level=1)
doc.add_paragraph(
    "This tool requires ffmpeg and ffprobe to be installed and on PATH. "
    "ffmpeg decodes and encodes the audio formats the suite reads and "
    "writes; ffprobe reads stream metadata (duration, codec, bitrate) "
    "for analysis tools. The suite shows a startup warning naming any "
    "missing binary and how to install it."
)
for _hint in (
    "Windows: winget install Gyan.FFmpeg",
    "macOS:   brew install ffmpeg",
    "Linux:   sudo apt install ffmpeg",
):
    doc.add_paragraph(_hint, style="List Bullet")

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Silence_Trimmer.pyw.",
    "Select the music directory.",
    "Set the silence threshold and decide whether to apply fades.",
    "Leave Dry Run selected for the first pass.",
    "Click the trimming start button.",
    "Review per-file trim amounts, skips, errors, and audition "
    "the result when available.",
    "Run again in Live Update when the trim plan is acceptable.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the directory picker, execution mode, "
    "silence threshold, fade settings, "
    "parallel controls, presets, and the start button. The "
    "last-used directory and worker "
    "settings are restored through shared suite configuration."
)
bullet(
    doc,
    "Silence Threshold: ",
    "GUI range -80 to -20 dB. Audio below this threshold is "
    "treated as silence; the default "
    "is -50 dB.",
)
bullet(
    doc,
    "Apply Fades: ",
    "Adds fade-in and fade-out after trimming. Disabled by default.",
)
bullet(doc, "Fade-in: ", "GUI range 100 to 3000 ms; default 500 ms.")
bullet(doc, "Fade-out: ", "GUI range 100 to 5000 ms; default 1000 ms.")
bullet(
    doc,
    "Parallel Processing: ",
    "The tool is already I/O-oriented and uses a thread pool; "
    "the parallel flag is accepted "
    "for consistency while worker count controls concurrency.",
)
bullet(
    doc,
    "Presets: ",
    "Save and restore threshold, fade settings, parallel "
    "settings, worker count, and run "
    "mode.",
)

doc.add_heading("Silence Detection", level=1)
doc.add_paragraph(
    "The core uses pydub's leading-silence detector twice: once "
    "on the original audio for "
    "the beginning and once on a reversed copy for the ending. "
    "The detected start and end "
    "durations are reported separately and summed into the "
    "total trim."
)
bullet(doc, "No silence: ", "The file is marked SKIP and is not rewritten.")
bullet(
    doc,
    "Dry Run: ",
    "The file is marked DRY_RUN with the total milliseconds that would be removed.",
)
bullet(
    doc,
    "Live Update: ",
    "The detected ranges are removed and the optional fades are "
    "applied to the remaining "
    "audio.",
)
bullet(
    doc,
    "Fade clamping: ",
    "Fade durations are clamped to the trimmed segment length "
    "so short audio is not extended "
    "or corrupted by an overlong fade.",
)

doc.add_heading("Processing Workflow", level=1)
for step in (
    "Discover supported files in the selected directory.",
    "Copy each source into a worker temporary directory and read its original tags.",
    "Load the temporary audio and detect leading and trailing silence.",
    "Skip clean files or report the Dry Run trim plan.",
    "In Live Update, trim, apply optional fades, export with "
    "format-appropriate settings, "
    "and restore tags.",
    "Record a tag-edit snapshot and replace the original file after successful export.",
    "Remove temporary worker files and report the per-file result.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Dry Run and Live Update", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Calculates trim amounts without exporting or changing "
    "files. Result audio is unavailable"
    " for playback because nothing was written.",
)
bullet(
    doc,
    "Live Update: ",
    "Exports trimmed audio, restores original tags, and "
    "replaces the source after successful "
    "processing.",
)
bullet(
    doc,
    "Tag preservation: ",
    "Existing tags are read before processing and restored to the exported file.",
)
bullet(
    doc,
    "Undo: ",
    "Live edits participate in the suite tag-edit batch for available recovery.",
)

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing displays four worker slots, progress, "
    "processed/succeeded/skipped/error "
    "counts, and per-file actions such as trim durations or No "
    "silence detected. Results "
    "reports total, reduced/succeeded files, skips, errors, and "
    "elapsed time. The Results tab "
    "can open the log and source folder and audition original "
    "versus result for successful "
    "Live Update output."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is silence-trim. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app silence-trim <folder>"
    " --dry\n"
    "python -m audio_suite_shared.cli_app silence-trim <folder>"
    " --threshold -50\n"
    "python -m audio_suite_shared.cli_app silence-trim <folder>"
    " --fade --fade-in 500 --fade-out"
    " 1000 --live\n"
    "python -m audio_suite_shared.cli_app silence-trim <folder>"
    " --parallel --workers 4 --json"
)
doc.add_paragraph(
    "Options include --threshold, --fade/--no-fade, --fade-in, "
    "--fade-out, "
    "--parallel/--no-parallel, "
    "--workers, the shared dry-run/live option, and --json. CLI "
    "defaults are threshold -50 dBFS "
    "and 25 ms fades, differing from the GUI defaults of "
    "500/1000 ms when fades are enabled."
)

doc.add_heading("Audio Refinery and Master Driver", level=1)
doc.add_paragraph(
    "Silence Trimmer is available as the trim stage in Audio "
    "Refinery, where it can run "
    "before noise reduction, enhancement, and normalization. "
    "From the Master Driver, use the "
    "Refinery workflow or the registered utility card as provided"
    " by the suite."
)

doc.add_heading("Troubleshooting", level=1)
bullet(doc, "No files found: ", "Select a directory containing supported audio files.")
bullet(
    doc,
    "Silence was not removed: ",
    "Raise the threshold toward -20 dB if the quiet material "
    "is above the current threshold;"
    " review the resulting trim amount in Dry Run first.",
)
bullet(
    doc,
    "Too much audio was removed: ",
    "Lower the threshold toward -80 dB and inspect the file "
    "before running Live Update.",
)
bullet(
    doc,
    "Fades sound unexpected: ",
    "Use shorter fade values and remember that fades are "
    "applied after trimming and clamped "
    "to the remaining segment length.",
)
bullet(
    doc,
    "A file errors: ",
    "Check codec support, permissions, temporary disk space, "
    "and the per-file status/log.",
)
bullet(
    doc,
    "Tags changed unexpectedly: ",
    "Confirm Live Update was selected and use the suite's undo "
    "snapshots or backups for "
    "recovery.",
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Run Dry Run first and audition representative results. Live"
    " Update replaces source files "
    "only after successful temporary export, but keep backups of "
    "important audio. Use the "
    "suite's Undo Tag Edits workflow for available metadata "
    "snapshots."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
