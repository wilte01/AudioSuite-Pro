"""Generate a fresh Genre Sync guide from the current implementation."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "genre_sync\\Genre_Sync_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Genre Sync - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Genre Sync implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Genre Sync derives the standard Genre field from "
    "custom genre metadata already "
    "stored in each audio file. It reads the file's "
    "TITLEGENRE, ARTISTGENRE, "
    "ALBUMGENRE, ALBUMARTISTGENRE, and Album Artist "
    "values, applies a deterministic "
    "priority tree, and can write the selected genre "
    "plus a derived-source marker."
)
doc.add_paragraph(
    "The tool processes supported audio files "
    "recursively with four worker slots. Dry "
    "Run reports the decisions without changing tags; "
    "Live Update writes changes in place."
)

doc.add_heading("Supported Input", level=1)
doc.add_paragraph(
    "Genre Sync uses the shared Audio Suite "
    "supported-extension list, including MP3, "
    "FLAC, OGG, M4A, Opus, WAV, and AIFF variants. "
    "The selected directory must contain "
    "at least one supported audio file."
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Genre_Sync.pyw.",
    "Select the music directory.",
    "Leave Dry Run selected for the first pass.",
    "Click BEGIN GENRE UPDATE.",
    "Review the selected source and skipped/error entries in Processing and Results.",
    "Run again in Live Update when the derived Genre values are acceptable.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the directory picker, mode, shared "
    "run controls, presets, a summary "
    "of the resolution rules, and the start button. "
    "The last-used directory is restored "
    "from shared suite configuration."
)
bullet(
    doc, "Dry Run: ", "Analyze and report the selected genre without writing metadata."
)
bullet(
    doc,
    "Live Update: ",
    "Write the standard Genre field and the derived "
    "marker when a usable source is resolved.",
)
bullet(
    doc,
    "Presets: ",
    "Save, load, and delete the run mode through shared suite preset storage.",
)
bullet(
    doc,
    "Auto Start / Auto End: ",
    "A Master Driver launch can start automatically "
    "and close the window after completion when "
    "configured.",
)
bullet(
    doc,
    "Clear Logs / Parallel: ",
    "These values mirror Master Driver orchestration "
    "state; the tool's own worker processing "
    "remains managed by its processor.",
)

doc.add_heading("Genre Resolution Rules", level=1)
doc.add_paragraph("For each file, the resolver applies these rules in order:")
for step in (
    "TITLEGENRE always wins when it contains a non-blank value.",
    "For an Album Artist equal to Various Artists, "
    "use ARTISTGENRE first, then ALBUMGENRE.",
    "For all other albums, use ALBUMGENRE first, then ALBUMARTISTGENRE.",
    "If no usable field exists, report No genre fields "
    "available and do not update the file.",
):
    doc.add_paragraph(step, style="List Number")
doc.add_paragraph(
    "The resolver treats ordinary whitespace and "
    "several Unicode blank characters as empty. "
    "A field that exists but contains only visual "
    "whitespace cannot become the selected genre."
)

doc.add_heading("Tags Written", level=1)
doc.add_paragraph(
    "On a Live Update when a source is resolved, Genre "
    "Sync writes the selected value to the "
    "format-appropriate standard Genre field, such as "
    "TCON for MP3, and writes the derived "
    "source marker through the shared tag helper. "
    "The marker identifies the winning source:"
)
bullet(doc, "Title selection: ", "Derived value Title.")
bullet(doc, "Artist selection for Various Artists: ", "Derived value Artist.")
bullet(doc, "Album selection: ", "Derived value Album.")
bullet(doc, "Album Artist selection: ", "Derived value Album Artist.")
bullet(doc, "Undo: ", "Live tag writes are recorded in the suite's batch undo context.")

doc.add_heading("Status Categories", level=1)
bullet(
    doc,
    "UPDATED: ",
    "Live Update changed the genre fields, or the "
    "corresponding dry-run entry says Would set.",
)
bullet(
    doc,
    "DRY RUN: ",
    "A usable source was selected and the file would be updated in Live Update.",
)
bullet(
    doc, "SKIPPED: ", "The current standard Genre already equals the resolved genre."
)
bullet(doc, "NO GENRE: ", "No usable custom genre field matched the rules.")
bullet(
    doc, "ERROR: ", "Tag reading, resolution, or writing raised an error for the file."
)

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows four worker slots, a progress "
    "bar, processed count, updated count, "
    "already-correct skips, no-genre count, errors, "
    "and per-file decision text. Work is "
    "retried through the shared processor recovery "
    "path. Results reports the same counts "
    "and elapsed time, then opens the Results tab "
    "when complete."
)

doc.add_heading("Logs", level=1)
bullet(doc, "Dry Run log: ", "genre_dry_run_log.txt in the selected directory.")
bullet(doc, "Live log: ", "genre_sync_log.txt in the selected directory.")
bullet(
    doc,
    "Sections: ",
    "Mode, updated or would-update files, "
    "already-correct skips, files with no genre "
    "fields, errors, and a summary with counts and "
    "elapsed time.",
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is genre-sync. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app genre-sync "
    "<folder> --dry\n"
    "python -m audio_suite_shared.cli_app genre-sync "
    "<folder> --live\n"
    "python -m audio_suite_shared.cli_app genre-sync "
    "<folder> --json"
)
doc.add_paragraph(
    "The CLI accepts the shared directory, "
    "dry-run/live, and JSON options. It validates "
    "that the supplied path exists and is a directory. "
    "Text output reports files, updated, "
    "skipped, no-genre, errors, and elapsed time."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open ENRICH and select "
    "Genre Sync. The driver passes the directory, "
    "Dry Run or Live Update mode, Auto Start, "
    "Auto End, and display-only orchestration state. "
    "A driven launch can auto-start and locks "
    "inherited controls in the child window."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No supported audio files found: ",
    "Select a directory containing supported audio files.",
)
bullet(
    doc,
    "No genre fields: ",
    "Check TITLEGENRE, ARTISTGENRE, ALBUMGENRE, and "
    "ALBUMARTISTGENRE values and confirm they are "
    "not blank or whitespace-only.",
)
bullet(
    doc,
    "A file is already correct: ",
    "SKIPPED is expected when the standard Genre already equals the resolved value.",
)
bullet(
    doc,
    "Unexpected source selected: ",
    "Check the Album Artist value: Various Artists "
    "changes precedence from Album Genre to "
    "Artist Genre.",
)
bullet(
    doc,
    "Tags did not change: ",
    "Confirm Live Update was selected. Dry Run never writes metadata.",
)
bullet(
    doc,
    "A file errors: ",
    "Review the per-file message and "
    "genre_sync_log.txt for tag-reader or writer "
    "details.",
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Run Dry Run first and review which field wins for "
    "each file. Live Update modifies tags "
    "in place. Keep backups of important files and "
    "use the suite's Undo Tag Edits workflow "
    "for available snapshots."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
