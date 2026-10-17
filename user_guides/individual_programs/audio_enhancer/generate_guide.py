"""Generate a fresh Audio Enhancer guide from the current implementation.

Changed in 1.1.0: documented the five DSP effects (reverse, karaoke, phaser,
flanger, chorus), the widened input-format list, and the formats the enhancer
skips because it cannot re-encode them.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "audio_enhancer\\Audio_Enhancer_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Audio Enhancer - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Audio Enhancer implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Audio Enhancer applies a configurable DSP chain to audio "
    "files and writes "
    "enhanced copies to an output folder. It measures the "
    "source and processed RMS "
    "levels, reports the gain change, writes the audio in its "
    "original format family, "
    "and copies source metadata tags to the result."
)
doc.add_paragraph(
    "The application accepts either one audio file or a folder "
    "of supported audio "
    "files. Leave the output folder blank to replace files in "
    "place, or select a "
    "different output folder to preserve the originals."
)

doc.add_heading("Supported Audio", level=1)
doc.add_paragraph(
    "The processor uses the shared Audio Suite format list. "
    "Supported extensions are "
    ".mp3, .flac, .ogg, .m4a, .mp4, .opus, .wav, .aiff, and "
    ".aif. A single selected "
    "file is processed directly; a selected directory is scanned"
    " for supported files "
    "in that directory."
)
bullet(
    doc,
    "MP3 output: ",
    "Preserves the source CBR/VBR quality selection when writing the enhanced file.",
)
bullet(
    doc,
    "FFmpeg-backed formats: ",
    "Formats such as M4A, MP4, Opus, and other configured "
    "formats use the tool's FFmpeg I/O "
    "path when required.",
)
bullet(
    doc,
    "Formats skipped: ",
    "The shared Audio Suite format list also includes .wv, .spx, "
    ".ape, .mpc, and .dsf. The enhancer cannot re-encode these "
    "(WavPack and Speex have no encoder on this path, and Monkey's "
    "Audio, Musepack, and DSD are read-only), so such files are "
    "reported as skipped rather than processed.",
)
bullet(
    doc,
    "Metadata: ",
    "Tags are copied after the audio output is written, "
    "including ID3, Vorbis-style, MP4, "
    "and available WAV metadata through the shared tag "
    "implementation.",
)

doc.add_heading("System Requirements", level=1)
doc.add_paragraph(
    "This tool requires ffmpeg and ffprobe to be installed and on PATH. "
    "ffmpeg decodes and encodes the audio formats the suite reads and "
    "writes; ffprobe reads stream metadata (duration, codec, bitrate) "
    "for analysis tools. The suite shows a startup warning naming any "
    "missing binary and how to install it."
)
for _hint in (
    "Windows: winget install Gyan.FFmpeg",
    "macOS:   brew install ffmpeg",
    "Linux:   sudo apt install ffmpeg",
):
    doc.add_paragraph(_hint, style="List Bullet")

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Audio_Enhancer.pyw from the AudioSuite-Pro folder.",
    "Choose a source file or folder.",
    "Leave Output folder blank for in-place replacement, or "
    "select a separate output folder.",
    "Select Mild, Moderate, Aggressive, or Custom.",
    "For a first pass, use a separate output folder and the Moderate preset.",
    "Click BEGIN ENHANCEMENT.",
    "Review the Processing and Results tabs, then audition the "
    "original and result when "
    "available.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("GUI Tabs", level=1)
doc.add_heading("Setup", level=2)
doc.add_paragraph(
    "Setup contains the source path, output folder, the Preset "
    "selector, and the Custom "
    "DSP summary row. Browse for a folder or a single audio file. "
    "If the output folder "
    "is blank, the application uses the source directory and "
    "overwrites the source "
    "filename with the enhanced output. The last source location"
    " is remembered in "
    "shared suite configuration."
)
bullet(
    doc,
    "Preset: ",
    "One selector for both kinds of preset. The built-in presets (Mild, "
    "Moderate, Aggressive, Custom) change DSP values only. Saved presets "
    "appear below a divider and also restore the source and output paths "
    "stored with them. The caption under the selector states which kind is "
    "active and, for a saved preset, names the built-in preset it carries. "
    "Save stores the current settings under a new name; Delete removes the "
    "selected saved preset and is greyed out while a built-in preset is "
    "selected.",
)
bullet(
    doc,
    "Custom: ",
    "Selecting it shows a Configure Custom DSP button and opens the parameter "
    "editor in its own non-modal window, which stays open while you start a "
    "run and applies changes immediately. The Setup tab keeps only the summary "
    "row, so the BEGIN ENHANCEMENT button is always visible.",
)
bullet(
    doc,
    "BEGIN ENHANCEMENT: ",
    "Validates the source path, resolves the output directory, "
    "then starts background "
    "processing.",
)

doc.add_heading("Processing", level=2)
doc.add_paragraph(
    "Processing shows an overall progress bar, "
    "processed/enhanced/error counts, and "
    "per-file status lines. Four worker slots are used by "
    "default. Each worker may "
    "retry a processing operation up to three times with backoff"
    " before reporting an "
    "error. A file that disappears during processing is reported"
    " as skipped; permission "
    "and other processing failures are reported as errors."
)

doc.add_heading("Results", level=2)
doc.add_paragraph(
    "Results shows total time, total files, enhanced files, and "
    "errors. The log is "
    "written as enhancement_log.txt in the output directory. "
    "When a successful result "
    "is available, the Results tab exposes original and enhanced "
    "audition controls. "
    "For an in-place run, the app preserves an audition copy of "
    "the original before "
    "replacement; for a separate output run, it compares the "
    "source file with the "
    "new output file."
)

doc.add_heading("DSP Presets", level=1)
doc.add_paragraph(
    "Preset values are defined in code and configure the "
    "complete enhancement pass. "
    "The presets are starting points, not loudness guarantees; "
    "always listen to a "
    "representative result before processing a large library."
)
bullet(
    doc,
    "Mild: ",
    "Noise gate -70 dB, compression threshold -18 dB at 2.5:1,"
    " gentle low/presence/high-"
    "shelf boosts, +2 dB output gain, -1 dB limiter.",
)
bullet(
    doc,
    "Moderate: ",
    "Noise gate -65 dB, compression threshold -22 dB at 3.5:1,"
    " stronger EQ shaping, "
    "+4 dB output gain, -0.5 dB limiter. This is the default.",
)
bullet(
    doc,
    "Aggressive: ",
    "Noise gate -60 dB, compression threshold -26 dB at 5:1, "
    "stronger EQ and +6 dB output "
    "gain, -0.3 dB limiter.",
)
bullet(
    doc,
    "Custom: ",
    "Opens the Custom DSP editor; its current values are used for the "
    "enhancement pass.",
)

doc.add_heading("Custom DSP Controls", level=1)
bullet(
    doc,
    "Noise Gate Threshold (dB): ",
    "Sets the threshold for quiet-signal gating; range -90 to -40 dB.",
)
bullet(
    doc,
    "Compressor Threshold (dB): ",
    "Sets the level where compression begins; range -40 to -6 dB.",
)
bullet(doc, "Compressor Ratio: ", "Sets compression strength; range 1:1 to 10:1.")
bullet(
    doc,
    "Compressor Attack / Release: ",
    "Controls compressor response timing; attack range 1 to "
    "50 ms and release range 20 to "
    "500 ms.",
)
bullet(
    doc,
    "Low Shelf Gain / Frequency: ",
    "Shapes low frequencies with gain from -6 to +12 dB and "
    "frequency from 40 to 400 Hz.",
)
bullet(
    doc,
    "Presence Gain / Frequency: ",
    "Shapes the presence band with gain from -6 to +12 dB and "
    "frequency from 500 to "
    "8000 Hz.",
)
bullet(
    doc,
    "High Shelf Gain / Frequency: ",
    "Shapes high frequencies with gain from -6 to +12 dB and "
    "frequency from 2000 to "
    "16000 Hz.",
)
bullet(doc, "Output Gain: ", "Applies final gain from -6 to +18 dB before limiting.")
bullet(
    doc,
    "Limiter Threshold: ",
    "Sets the final ceiling; the value must be at or below 0 dB"
    " and the UI range is -12 to "
    "-0.1 dB.",
)

doc.add_heading("Effects", level=1)
doc.add_paragraph(
    "The Custom DSP editor also exposes five optional effects, and the command "
    "line can enable the same effects through flags. All five default to "
    "off, so a preset run or an untouched Custom run produces exactly the "
    "same output as before the effects existed. Applying a preset resets "
    "only the numeric DSP values; it never changes the effect toggles. "
    "Reverse and karaoke are applied before the main chain, while the "
    "phaser, flanger, and chorus run after the EQ shaping. Dry Run disables "
    "all five effects so it still measures the base chain only."
)
bullet(
    doc,
    "Reverse: ",
    "Reverses the audio playback direction.",
)
bullet(
    doc,
    "Karaoke: ",
    "Removes centre-panned content (typically lead vocals) by subtracting "
    "one channel from the other. Karaoke Strength, 0 to 1, blends between "
    "the original and the fully removed signal. Mono input is returned "
    "unchanged.",
)
bullet(
    doc,
    "Phaser: ",
    "A sweeping phase effect applied after the EQ shaping.",
)
bullet(
    doc,
    "Flanger: ",
    "A short modulated delay mixed with the dry signal.",
)
bullet(
    doc,
    "Chorus: ",
    "Multiple slightly detuned copies added after the flanger.",
)

doc.add_heading("Processing Pipeline", level=1)
for step in (
    "Inspect the source for format, sample rate, channels, "
    "duration, bitrate, VBR state, "
    "and tag count.",
    "Read audio through pedalboard I/O or the FFmpeg-backed path"
    " for formats that need it.",
    "Apply any enabled effects, then the enhancement chain: reverse "
    "and karaoke first, then gate/compression, EQ shaping, phaser, "
    "flanger, and chorus, followed by output gain and limiting.",
    "Measure RMS before and after processing and run the post-chain clip check.",
    "Write the enhanced audio to the output path, preserving "
    "the source extension and "
    "relevant quality settings.",
    "Copy tags from source to output and report the number of "
    "copied tag frames or fields.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("In-Place versus Separate Output", level=1)
bullet(
    doc,
    "Separate output: ",
    "Recommended for review. The source remains available and "
    "the result is written to the "
    "chosen output directory with the same filename.",
)
bullet(
    doc,
    "In-place output: ",
    "The source directory is used as output, so the original "
    "file is replaced. The app "
    "prepares an audition copy for the first result when "
    "possible.",
)
bullet(
    doc,
    "Output creation: ",
    "The output directory is created when needed. A log is "
    "always associated with the "
    "resolved output directory.",
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph(
    "The registered command is enhance. It requires an output "
    "directory and uses a "
    "named preset. Examples:"
)
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app enhance <folder> "
    "--output-dir <output>\n"
    "python -m audio_suite_shared.cli_app enhance <folder> "
    "--output-dir <output> --preset mild\n"
    "python -m audio_suite_shared.cli_app enhance <folder> "
    "--output-dir <output> --preset "
    "aggressive --json"
)
doc.add_paragraph(
    "The CLI accepts Mild, Moderate, Aggressive, or Custom presets "
    "case-insensitively. Custom uses the neutral default DSP values, "
    "since the CLI has no per-slider options. Use --json "
    "for structured output; normal output reports files, "
    "enhanced count, errors, and "
    "elapsed time."
)
doc.add_paragraph(
    "The effects are available as flags: --reverse, --phaser, "
    "--flanger, --chorus, and --karaoke (or --no-karaoke to leave it "
    "off), plus --karaoke-strength with a value from 0.0 to 1.0. "
    "For example:\n"
    "python -m audio_suite_shared.cli_app enhance <folder> "
    "--output-dir <output> --karaoke --karaoke-strength 0.8"
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "Audio Enhancer is reachable three ways: its own window, the Audio "
    "Enhancer stage of the Master Driver's REFINE tab, and the Audio "
    "Enhancer card on the Master Driver's UTIL tab. The UTIL card "
    "requires an output folder and pre-fills the preset with Mild, "
    "Moderate, Aggressive, or Custom; Custom starts the tool with the DSP "
    "values it currently has loaded, because the card has no sliders. "
    "From the REFINE tab the driver supplies the source, stage "
    "configuration, output, and inherited orchestration settings. Direct "
    "launch is useful when you need the full Audio Enhancer Setup "
    "controls and audition view."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No supported audio files found: ",
    "Select a supported file or a folder containing supported "
    "extensions directly inside "
    "it.",
)
bullet(
    doc,
    "The result sounds too strong: ",
    "Use Mild or reduce Custom output gain, EQ gains, "
    "compression ratio, or compression "
    "threshold strength.",
)
bullet(
    doc,
    "Tags are missing: ",
    "Review the result log. Tag copying depends on readable "
    "source tags and supported "
    "destination metadata.",
)
bullet(
    doc,
    "A file reports an error: ",
    "Check codec support, file permissions, locks, available "
    "disk space, and the detailed "
    "enhancement_log.txt entry.",
)
bullet(
    doc,
    "The output overwrote the source: ",
    "The Output folder was blank or resolved to the source "
    "directory. Use a separate output"
    " folder for review runs.",
)
bullet(
    doc,
    "Audition is unavailable: ",
    "No successful output or no usable original/result pair was"
    " available for the selected "
    "run.",
)

doc.add_heading("Logs and Recovery", level=1)
doc.add_paragraph(
    "The enhancement_log.txt report records each completed, "
    "skipped, or failed file, "
    "the measured RMS gain for successful files, copied tag "
    "counts, and the final "
    "summary. In-place enhancement changes source audio files, "
    "so use a separate output "
    "directory or maintain an independent backup before "
    "processing important material."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
