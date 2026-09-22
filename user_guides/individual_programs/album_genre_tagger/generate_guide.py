"""Generate a fresh Album Genre Tagger guide from the current implementation."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "album_genre_tagger\\AlbumGenre_Tagger_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Album Genre Tagger - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Album Genre Tagger implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Album Genre Tagger consolidates per-track "
    "TITLEGENRE metadata into one album-level genre "
    "value. It reads every supported audio file in a selected "
    "directory, checks how many files contain TITLEGENRE, "
    "builds a sorted list of unique genre values, and "
    "writes that result to ALBUMGENRE on every file "
    "when the coverage requirement is met. It also "
    "marks the result with ALBUMDERIVED = Y."
)
doc.add_paragraph(
    "The tool is intentionally guarded by a coverage "
    "check. If too few files have TITLEGENRE, it "
    "stops before Phase 2 so an incomplete track-level "
    "sample does not produce an album-level tag."
)

doc.add_heading("Supported Input and Output", level=1)
bullet(
    doc,
    "Input files: ",
    "Any of the nine suite audio formats in the selected "
    "directory are scanned: .mp3, .flac, .ogg, .m4a, .mp4, "
    ".opus, .wav, .aiff, and .aif.",
)
bullet(
    doc,
    "Input tag: ",
    "TITLEGENRE is read from each file through the shared, "
    "format-aware tag handler (TXXX for ID3 files, Vorbis "
    "comments for FLAC/OGG/Opus, and MP4 freeform tags for "
    "M4A/MP4).",
)
bullet(
    doc,
    "Output tag: ",
    "ALBUMGENRE receives the sorted unique genres "
    "joined with semicolons and spaces, stored through "
    "the same format-aware handler as the input.",
)
bullet(
    doc,
    "Derived flag: ",
    "ALBUMDERIVED is written as Y when the consolidation write runs.",
)
bullet(
    doc,
    "Other files: ",
    "Files outside the nine supported audio extensions "
    "and subdirectory contents are not part of this operation.",
)

doc.add_heading("How Consolidation Works", level=1)
for step in (
    "The selected directory is scanned for supported audio files.",
    "Phase 1 reads TITLEGENRE from each supported audio file "
    "and records readable values.",
    "Coverage is calculated as files with TITLEGENRE divided by total audio files.",
    "If coverage meets the configured minimum, unique "
    "genre values are sorted alphabetically.",
    "The sorted values are joined into one string, for example Rock; Soul; Vocal.",
    "Phase 2 writes ALBUMGENRE and ALBUMDERIVED to every supported audio file.",
):
    doc.add_paragraph(step, style="List Number")
doc.add_paragraph(
    "A file without tags, without TITLEGENRE, or "
    "with an unreadable tag is counted as a read "
    "error and does not contribute to coverage."
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open AlbumGenre_Tagger.pyw from the AudioSuite-Pro folder.",
    "On Setup, select the directory containing the album's audio files.",
    "Set Minimum Coverage to the percentage of files that must have TITLEGENRE.",
    "Leave Dry Run selected for the first pass.",
    "Click BEGIN GENRE CONSOLIDATION.",
    "Review the Phase 1 coverage result and the proposed consolidated genre.",
    "Run again with Live Update when the result is correct.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("GUI Tabs", level=1)
doc.add_heading("Setup", level=2)
doc.add_paragraph(
    "Setup contains the directory picker, run mode, "
    "Minimum Coverage field, shared run controls, "
    "presets, and the start button. The selected path "
    "must exist and be a directory. The last-used "
    "directory is restored from shared suite "
    "configuration when no path is supplied."
)
bullet(
    doc,
    "Minimum Coverage (%): ",
    "An integer from 0 through 100. The default is 50,"
    " meaning at least half of the audio files must "
    "have TITLEGENRE.",
)
bullet(
    doc,
    "Dry Run: ",
    "Reads and consolidates the data and reports what "
    "would be written without changing tags.",
)
bullet(
    doc,
    "Live Update: ",
    "Writes ALBUMGENRE and ALBUMDERIVED to the supported "
    "audio files after the coverage gate passes.",
)
bullet(
    doc,
    "Presets: ",
    "Save and reload the current run mode and minimum "
    "coverage through the shared suite configuration,"
    " or delete a saved preset.",
)
bullet(
    doc,
    "Auto Start / Auto End: ",
    "Auto Start can begin a Master Driver launch "
    "automatically. Auto End closes the window after "
    "completion or a coverage failure when enabled.",
)
bullet(
    doc,
    "Clear Logs / Parallel: ",
    "These values mirror Master Driver orchestration "
    "state; the Album Genre Tagger analysis itself "
    "is single-threaded.",
)

doc.add_heading("Processing", level=2)
doc.add_paragraph(
    "The Processing tab shows Phase 1 progress with "
    "processed count, files that contain TITLEGENRE,"
    " and read errors. It then lists Phase 2 dry-run "
    "or update statuses for each supported audio file. The progress "
    "bar reserves an additional step for the write "
    "phase."
)

doc.add_heading("Results", level=2)
doc.add_paragraph(
    "Results shows elapsed time, total files, files "
    "with TITLEGENRE and their coverage percentage,"
    " unique genres found, files updated or would be "
    "updated, and read or write errors. It also "
    "displays the consolidated album genre string. "
    "View Log opens AlbumGenre_Tagger_Log.txt, and "
    "Open Folder opens the source directory because "
    "tags are changed in place."
)

doc.add_heading("Coverage Gate", level=1)
doc.add_paragraph(
    "The gate compares the actual coverage percentage "
    "with Minimum Coverage. The comparison requires "
    "coverage to be at least the configured value. "
    "For example, with 10 audio files and 5 readable "
    "TITLEGENRE values, coverage is 50.0 percent: "
    "a 50 percent requirement passes, while a 60 "
    "percent requirement fails."
)
doc.add_paragraph(
    "When the gate fails, the application writes a "
    "failure log and shows the total, files with "
    "TITLEGENRE, measured coverage, required "
    "coverage, and an insufficient coverage status. "
    "No ALBUMGENRE or ALBUMDERIVED writes are "
    "attempted."
)

doc.add_heading("Dry Run and Live Update", level=1)
bullet(
    doc,
    "Dry Run: ",
    "The writer validates each target operation and "
    "reports [DRY RUN], but does not modify any "
    "tags.",
)
bullet(
    doc,
    "Live Update: ",
    "The writer saves the consolidated value and "
    "ALBUMDERIVED flag to each supported audio file "
    "after Phase 1 passes.",
)
bullet(
    doc,
    "In-place operation: ",
    "There is no output directory; source files are "
    "updated in their selected directory.",
)
bullet(
    doc,
    "Undo support: ",
    "The run is placed inside the suite batch undo "
    "context for tag edits. Keep backups for "
    "dependable recovery.",
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph(
    "The registered command is album-genre. It "
    "supports unattended analysis and uses the same "
    "coverage gate as the GUI. Examples:"
)
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app album-genre "
    "<folder> --dry\n"
    "python -m audio_suite_shared.cli_app album-genre "
    "<folder> --coverage 75\n"
    "python -m audio_suite_shared.cli_app album-genre "
    "<folder> --coverage 50 --json"
)
doc.add_paragraph(
    "The --coverage option accepts 0 through 100 and "
    "defaults to 50. Use the shared dry-run option "
    "to preview changes and --json to print the "
    "structured result. The command reports coverage "
    "failures separately from ordinary write errors."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open ENRICH and select "
    "Album Genre Tagger. The driver passes the shared "
    "directory, Dry Run or Live Update mode, start "
    "and end flags, and the configured coverage "
    "value. A driven launch can auto-start, while "
    "the child window displays inherited orchestration"
    " controls."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No supported audio files found: ",
    "Select a directory containing files in one of the "
    "nine supported audio formats. Other file types and "
    "subdirectories are not scanned by this tool.",
)
bullet(
    doc,
    "Coverage check failed: ",
    "Add or repair TITLEGENRE tags, choose a directory"
    " with the intended album files, or lower the "
    "minimum only when the available metadata is "
    "sufficient for your workflow.",
)
bullet(
    doc,
    "Genres are missing: ",
    "The source file may have no tags, no TITLEGENRE field, or an "
    "unreadable tag. Review Phase 1 entries and the log.",
)
bullet(
    doc,
    "The result contains several genres: ",
    "That is expected when multiple unique TITLEGENRE"
    " values exist; the tool sorts them and joins "
    "them in one ALBUMGENRE string.",
)
bullet(
    doc,
    "Write errors occurred: ",
    "Check file permissions, locks, and the per-file "
    "messages in AlbumGenre_Tagger_Log.txt.",
)
bullet(
    doc,
    "Nothing changed: ",
    "Confirm that the run was Live Update rather than "
    "Dry Run and that the coverage gate passed.",
)

doc.add_heading("Logs and Recovery", level=1)
doc.add_paragraph(
    "The GUI writes AlbumGenre_Tagger_Log.txt in the "
    "selected directory. The log contains the mode, "
    "directory, coverage requirement, Phase 1 results,"
    " coverage decision, unique genre counts, "
    "consolidated tags, Phase 2 statuses, and summary."
    " Use the suite's Undo Tag Edits workflow for "
    "available snapshots, and retain an independent "
    "backup before live metadata changes."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
