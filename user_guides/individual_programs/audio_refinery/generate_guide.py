"""Generate a fresh Audio Refinery guide from the current implementation.

Changed in 1.1.0: documented the DSP effect controls available under the
Custom enhancement preset (reverse, karaoke, phaser, flanger, chorus) and the
matching CLI effect flags.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "audio_refinery\\Audio_Refinery_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Audio Refinery - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Audio Refinery implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Audio Refinery is a single-pass audio processing pipeline. It loads each source "
    "file once, runs the selected stages in memory, exports the result, restores the "
    "original tags, and optionally applies loudness normalization. The final file is "
    "placed in a separate output folder or atomically replaces the source when running "
    "in place."
)
doc.add_paragraph(
    "The pipeline is designed to combine several suite tools without "
    "repeatedly decoding "
    "and encoding the same file. "
    "It validates file integrity before expensive DSP work, "
    "uses four worker slots by default, retries transient failures, and records "
    "completed "
    "files in a hash-based registry."
)

doc.add_heading("Supported Audio", level=1)
doc.add_paragraph(
    "The refinery discovers the shared Audio Suite audio extensions, including MP3, "
    "FLAC, OGG, M4A, Opus, and WAV. The pipeline uses pydub/FFmpeg-backed decoding and "
    "format-aware export. Files must pass existence, readability, "
    "size, container-header, "
    "decode, sample-rate, duration, and channel checks before processing."
)

doc.add_heading("System Requirements", level=1)
doc.add_paragraph(
    "This tool requires ffmpeg and ffprobe to be installed and on PATH. "
    "ffmpeg decodes and encodes the audio formats the suite reads and "
    "writes; ffprobe reads stream metadata (duration, codec, bitrate) "
    "for analysis tools. The suite shows a startup warning naming any "
    "missing binary and how to install it."
)
for _hint in (
    "Windows: winget install Gyan.FFmpeg",
    "macOS:   brew install ffmpeg",
    "Linux:   sudo apt install ffmpeg",
):
    doc.add_paragraph(_hint, style="List Bullet")

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Audio_Refinery.pyw.",
    "Select a source folder containing supported audio files.",
    "Choose an output folder, or leave it blank for in-place replacement.",
    "Select the stages to run and expand their settings as needed.",
    "Leave Dry Run selected for the first pass.",
    "Click BEGIN REFINERY.",
    "Review the Processing statuses, refinery_log.txt, and Results audition controls.",
    "Run in Live Update only after confirming the intended stages and output.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains source and output paths, presets, the shared run controls, five "
    "stage panels, Reset Progress, and BEGIN REFINERY. The default active stages are "
    "declick, trim, enhance, and normalise; noise reduction is optional. At least one "
    "stage must be enabled."
)
bullet(
    doc,
    "Source folder: ",
    (
        "Required directory containing supported audio files. "
        "The application validates "
        "the directory and checks for audio before starting."
    ),
)
bullet(
    doc,
    "Output folder: ",
    (
        "Optional. Blank means in-place replacement; "
        "a selected existing directory receives results with the source filenames."
    ),
)
bullet(
    doc,
    "Dry Run: ",
    (
        "Validates and decodes files, reports RMS and duration information, "
        "and skips all modification and export stages."
    ),
)
bullet(
    doc,
    "Live Update: ",
    (
        "Runs the selected stages and writes results. The embedded Refinery "
        "always uses live stage processing after its driver-level dry-run decision."
    ),
)
bullet(
    doc,
    "Presets: ",
    (
        "Save, load, and delete the complete stage configuration, "
        "including stage toggles and parameters."
    ),
)
bullet(
    doc,
    "Reset Progress: ",
    (
        "Clears the source folder's completion registry "
        "so every file can be processed again."
    ),
)

doc.add_heading("Pipeline Stages", level=1)
doc.add_paragraph("When enabled, stages run in this implementation order:")
for step in (
    "Integrity validation: checks existence, readability, file size, "
    "magic bytes, decoding, sample rate, duration, and channels.",
    "Load: reads original tags and decodes the source audio.",
    "Pop & Click Remover: detects and repairs clicks using "
    "sensitivity and maximum defect width.",
    "Silence Trimming / Fading: removes silence using a dBFS threshold "
    "and can apply fade-in and fade-out.",
    "Noise Reduction: optionally removes mains hum and/or broadband noise using "
    "its noise profile and gate settings.",
    "Audio Enhancement: applies the Mild, Moderate, Aggressive, or Custom DSP "
    "preset; Custom adds per-effect toggles (reverse, karaoke, phaser, flanger, "
    "chorus).",
    "Export: writes the processed audio once using "
    "format-appropriate encoding settings.",
    "Normalise: optionally measures loudness and applies "
    "the configured LUFS target and tolerance.",
    "Finish: restores tags, records an undo snapshot, and atomically "
    "replaces or creates the output file.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Stage Controls", level=1)
bullet(
    doc,
    "Declick sensitivity: ",
    "Default 0.31; controls click detection. Lower values are more "
    "aggressive according to the stage model.",
)
bullet(doc, "Maximum defect width: ", "Default 2 ms; limits the click width repaired.")
bullet(doc, "Trim threshold: ", "Default -50 dBFS; defines silence for trimming.")
bullet(
    doc,
    "Fades: ",
    "Optional fade-in default 500 ms and fade-out default 1000 ms after trimming.",
)
bullet(
    doc,
    "Noise reduction: ",
    "Optional hum removal and broadband reduction. "
    "Defaults include 60 Hz hum, 8 harmonics, Q 30, a 0-500 ms noise profile, "
    "and a 3 dB gate threshold.",
)
bullet(
    doc,
    "Enhancement preset: ",
    "Mild, Moderate, Aggressive, or Custom. When "
    "Normalise is active, the refinery suppresses enhancement output gain "
    "so final loudness is controlled by the normalizer.",
)
bullet(
    doc,
    "Enhancement effects: ",
    "Relevant to the Custom preset: Reverse audio, Karaoke (centre-vocal "
    "removal with a 0.0-1.0 strength), Phaser, Flanger, and Chorus. Each is "
    "off by default, and the built-in Mild, Moderate, and Aggressive presets "
    "leave them off.",
)
bullet(
    doc,
    "Target LUFS: ",
    "Permitted targets are -12, -14, -16, -18, and -23 LUFS; default -16.0.",
)
bullet(
    doc,
    "Tolerance: ",
    "Default 0.5 LUFS; used by the normalization stage to avoid unnecessary adjustment "
    "when loudness is already close to target.",
)

doc.add_heading("Dry Run and Output Safety", level=1)
doc.add_paragraph(
    "In Dry Run, each valid file is loaded and measured, then "
    "the pipeline returns without "
    "running modification stages or writing an output. Results "
    "include a DryRun stage with "
    "RMS, duration, and channel information."
)
doc.add_paragraph(
    "In live processing, each file is written "
    "to a temporary file in the output directory. "
    "Tags are restored to that temporary file, and only "
    "then is it atomically moved into "
    "the final output path. If a stage or export fails, "
    "the temporary file is removed and "
    "the original remains intact."
)

doc.add_heading("Completion Registry", level=1)
doc.add_paragraph(
    "The refinery stores successful completion records "
    "in _refinery_completed.log in the "
    "source folder. "
    "Each record contains a SHA-256 hash and filename. A file is skipped "
    "only when both its name and pre-processing hash match a recorded completion, so a "
    "renamed or modified file is eligible for processing again. Use Reset Progress to "
    "clear all records for the selected source folder."
)

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows four worker slots, progress, and per-file DONE, SKIP, or ERROR "
    "entries. Each completed file's log includes stage messages and metrics. Results "
    "reports processed, succeeded, skipped, errors, and elapsed time. The Results tab "
    "can compare the original and result audio "
    "when a suitable pair exists; Dry Run has "
    "no result audio to play."
)

doc.add_heading("Logs and Recovery", level=1)
doc.add_paragraph(
    "The pipeline writes "
    "refinery_log.txt in the resolved output directory. It includes "
    "per-file stage output and a final "
    "total, processed, succeeded, skipped, error, and "
    "time summary. Tag edits are recorded in the suite batch undo system. "
    "Keep independent "
    "backups before in-place processing "
    "even though failures are protected by temporary "
    "output and atomic replacement."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is refine. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app refine <folder> --dry\n"
    "python -m audio_suite_shared.cli_app refine <folder> --output <output> --stages "
    "declick,trim,enhance,normalise\n"
    "python -m audio_suite_shared.cli_app refine <folder> "
    "--stages declick,enhance --preset Mild\n"
    "python -m audio_suite_shared.cli_app refine <folder> --reduce-hum "
    "--hum-frequency 60 --json"
)
doc.add_paragraph(
    "CLI options cover output directory, dry run, stages, "
    "declick sensitivity and width, "
    "trim threshold and fades, hum and broadband reduction, enhancement preset, "
    "the enhancement effect flags (--reverse, --karaoke/--no-karaoke, "
    "--karaoke-strength, --phaser, --flanger, --chorus), target "
    "LUFS, tolerance, and JSON output. "
    "If --output is omitted, results are written in place."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "Audio Refinery is the Master Driver's REFINE "
    "pipeline. The driver embeds the Refinery "
    "Setup controls, passes source and output folders, "
    "and supplies inherited Auto Start, "
    "Auto End, Clear Logs, and mode values. The driver uses its "
    "own batch orchestration; "
    "the Refinery's Parallel control is not an independent pipeline switch."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No audio files found: ",
    "Select a source folder containing "
    "supported files directly discoverable by the processor.",
)
bullet(
    doc,
    "Integrity check failed: ",
    "The file may be empty, unreadable, malformed, "
    "undecodable, outside size/duration/sample-rate/channel limits, "
    "or have invalid container headers.",
)
bullet(
    doc,
    "Files are skipped on a later run: ",
    "They match the completion registry by filename and SHA-256. "
    "Use Reset Progress when a full rerun is intended.",
)
bullet(
    doc,
    "Output directory error: ",
    "The selected output directory must already exist "
    "and be a directory; leave it blank for in-place mode.",
)
bullet(
    doc,
    "Normalization behaves unexpectedly: ",
    "Remember that the normalizer controls final level when active, and "
    "enhancement output gain is suppressed in that case.",
)
bullet(
    doc,
    "A stage fails: ",
    "Review refinery_log.txt for the stage-specific error. "
    "The original source is protected if the final replacement is not reached.",
)

doc.add_heading("Safety Recommendations", level=1)
doc.add_paragraph(
    "Run Dry Run first, prefer a separate output folder while "
    "tuning stages, and listen "
    "to the original/result audition pair. Use in-place mode only with backups. "
    "Reset the "
    "completion registry when configuration changes require "
    "already completed files to be "
    "processed again."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
