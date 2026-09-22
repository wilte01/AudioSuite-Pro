"""Generate a fresh Mood Analysis Batch Tagger guide from current code."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "mood_analysis_batch_tagger\\Mood_Analysis_Batch_Tagger_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Mood Analysis Batch Tagger - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph(
    "Based on the current Mood Analysis Batch Tagger implementation"
)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Mood Analysis Batch Tagger recursively scans a "
    "root music directory and all of its "
    "subfolders, analyzes mood from the audio signal, "
    "and can write mood, confidence, "
    "and energy tags. It uses the same signal mood "
    "engine as the standalone Mood Analyzer, "
    "so the same track produces the same values in "
    "both tools."
)

doc.add_heading("Supported Files", level=1)
doc.add_paragraph(
    "The batch engine uses the shared Audio Suite "
    "supported extensions, including MP3, "
    "FLAC, OGG, M4A, MP4, Opus, WAV, and AIFF "
    "variants. Unsupported files are ignored."
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Mood_Analysis_Batch_Tagger.pyw.",
    "Select the root directory containing your music library.",
    "Choose an analysis duration and optional parallel processing settings.",
    "Leave Dry Run selected for the first pass.",
    "Click the batch analysis start button.",
    "Review folders, per-file results, counts, and the log.",
    "Run again in Live Update when you want mood tags written.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the root folder, mode, duration, "
    "worker settings, presets, and shared "
    "run controls. The root directory is scanned "
    "recursively. Duration and parallel settings "
    "are saved in shared suite configuration."
)
bullet(
    doc,
    "Analysis Duration: ",
    "The default is 90 seconds. The CLI and parser "
    "allow 10 through 600 seconds. Shorter windows "
    "are faster but sample less of each track.",
)
bullet(
    doc,
    "Parallel Processing: ",
    "When enabled, CPU-bound mood analysis uses "
    "ProcessPoolExecutor to bypass the GIL. When "
    "disabled, a thread pool is used.",
)
bullet(doc, "Workers: ", "Sets the maximum worker count; the default is 4.")
bullet(
    doc,
    "Presets: ",
    "Save and restore root directory, mode, duration, parallel mode, and worker count.",
)
bullet(
    doc,
    "Dry Run / Live Update: ",
    "Dry Run analyzes without writing. Live Update writes the mood tags.",
)

doc.add_heading("Mood Analysis", level=1)
doc.add_paragraph(
    "The shared signal analyzer derives two axes from "
    "the audio: arousal, representing "
    "energy or calmness, and valence, representing "
    "positive or dark character. Features "
    "include tempo, spectral flux, zero-crossing rate, "
    "RMS energy, musical mode, spectral "
    "brightness, and low-frequency energy. The axes "
    "are mapped to three mood words, a "
    "confidence score from 0 to 1, and an energy band "
    "from 1 to 6."
)
bullet(
    doc,
    "Mood value: ",
    "Stored as three semicolon-separated words such as Energetic;Bright;Confident.",
)
bullet(
    doc,
    "Confidence: ",
    "Stored as a decimal such as 0.87 indicating how strongly the analysis matched.",
)
bullet(
    doc,
    "Energy: ",
    "Stored as an integer band from 1, calm, through 6, very energetic.",
)

doc.add_heading("Batch Workflow", level=1)
for step in (
    "Walk the root directory recursively and sort folders and filenames.",
    "Enter each folder in the Processing display and count its supported files.",
    "Load and analyze each file through the shared Mood Analyzer engine.",
    "In Dry Run, report the values without writing tags.",
    "In Live Update, record a tag snapshot and write "
    "mood, confidence, and energy fields.",
    "Write mood_analysis_batch_log.txt in the root directory and report final counts.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Dry Run and Live Update", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Reports the mood that would be written and leaves source files unchanged.",
)
bullet(
    doc,
    "Live Update: ",
    "Writes the three mood-related metadata values to each successfully analyzed file.",
)
bullet(
    doc,
    "Skipped files: ",
    "Files that cannot be decoded or analyzed, "
    "including unsupported/corrupt/silent audio, "
    "are counted as skipped when the failure is a "
    "known analysis or codec issue.",
)
bullet(
    doc,
    "Errors: ",
    "Unexpected failures are counted separately as errors and recorded in the log.",
)
bullet(
    doc,
    "Undo: ",
    "Live writes use the suite tag-edit batch and "
    "can be reverted through Undo Tag Edits when "
    "snapshots are available.",
)

doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows folders as they are entered, "
    "four active worker slots, file progress, "
    "and tagged/skipped/error counts. A successful "
    "file appears as OK in Live Update or a "
    "dry-run entry describing the values that would "
    "be written."
)
doc.add_paragraph(
    "Results reports folders scanned, files analyzed, "
    "files tagged, skipped files, and errors. "
    "Open Folder opens the root directory and "
    "View Log opens the generated batch report."
)

doc.add_heading("Log File", level=1)
doc.add_paragraph(
    "Each completed run writes "
    "mood_analysis_batch_log.txt in the selected root "
    "directory. The log records mode, timestamp, root, "
    "file count, duration, parallel settings, folder "
    "sections, per-file mood/confidence/energy "
    "information, skips, errors, and final totals."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is mood-analysis-batch. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app "
    "mood-analysis-batch <folder> --dry\n"
    "python -m audio_suite_shared.cli_app "
    "mood-analysis-batch <folder> --duration 60\n"
    "python -m audio_suite_shared.cli_app "
    "mood-analysis-batch <folder> --parallel "
    "--workers 4\n"
    "python -m audio_suite_shared.cli_app "
    "mood-analysis-batch <folder> --json"
)
doc.add_paragraph(
    "Options include --duration, "
    "--parallel/--no-parallel, --workers, "
    "the shared dry-run/live option, and --json. "
    "The CLI returns directory, duration, mode, "
    "worker settings, tagged, skipped, error, "
    "folder, and log-path values."
)

doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open UTIL and select "
    "Mood Analysis Batch Tagger. The driver passes "
    "the root directory, mode, duration, worker "
    "settings, start/end behavior, and orchestration "
    "flags through the launch handshake. A driven "
    "launch can auto-start and displays inherited "
    "controls."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "A folder was not scanned: ",
    "Confirm it is below the selected root and contains supported audio extensions.",
)
bullet(
    doc,
    "A file is skipped: ",
    "The file may be corrupt, unsupported by its "
    "decoder, or effectively silent. The log "
    "contains the reason when available.",
)
bullet(
    doc,
    "The mood seems subjective: ",
    "The result is signal-derived from measurable "
    "audio features, not lyrics, genre labels, or "
    "personal listening judgment.",
)
bullet(
    doc,
    "Parallel processing fails: ",
    "Disable parallel mode or reduce workers if "
    "process startup or system resources are "
    "constrained.",
)
bullet(
    doc,
    "Tags did not change: ",
    "Confirm Live Update was selected. Dry Run never writes files.",
)
bullet(
    doc,
    "No log is available: ",
    "A run must reach completion and be able to write to the selected root directory.",
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Run Dry Run first and review the batch log. Live "
    "Update modifies metadata throughout the "
    "library, so maintain backups of important files "
    "and use Undo Tag Edits for available tag-edit "
    "snapshots."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
