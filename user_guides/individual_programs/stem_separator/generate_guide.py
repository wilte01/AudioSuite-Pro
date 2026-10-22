"""Generate a fresh Stem Separator guide from the current implementation.

Changed in 2.0.0: initial version (Sprint 6, Item 6.1).  The version subtitle
renders ``declared_version()`` and the generation month dynamically, and the
output path is resolved from this file rather than an absolute checkout path.
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from audio_suite_shared.project_version import declared_version

OUTPUT = Path(__file__).with_name("Stem_Separator_User_Guide.docx")


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
h = doc.add_heading("Stem Separator \u2014 User Guide", level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(
    f"Version {declared_version()} \u2014 {date.today().strftime('%B %Y')}"
).alignment = WD_ALIGN_PARAGRAPH.CENTER

# ── Overview ──
doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Stem Separator splits a mixed track into its component parts with Demucs, "
    "an open-source source-separation model: vocals, drums, bass and everything "
    "else. It writes each part as its own file in an output folder and can also "
    "write a mix of the parts you selected."
)
doc.add_paragraph(
    "The source files are never modified. Every run produces new files, so there "
    "is nothing to undo: keep the originals where they are and delete the output "
    "folder if a separation is not useful."
)
doc.add_paragraph(
    "Separation is computationally expensive. A five-minute stereo track takes "
    "roughly one to two minutes on a modern CPU, and considerably less on a CUDA "
    "GPU. Demucs downloads its model weights on first use (about 80 MB, cached "
    "afterwards), so the first run of a session needs a network connection."
)

# ── Supported files ──
doc.add_heading("Supported Files", level=1)
doc.add_paragraph(
    "The processor reads the shared Audio Suite supported formats, including "
    "MP3, FLAC, OGG, M4A, MP4, Opus, WAV and AIFF variants. Files are resampled "
    "to 44.1 kHz stereo before separation, because that is the shape Demucs' "
    "model accepts; a mono source is duplicated to two channels."
)

# ── System requirements ──
doc.add_heading("System Requirements", level=1)
doc.add_paragraph(
    "This tool requires ffmpeg and ffprobe to be installed and on PATH. ffmpeg "
    "decodes and encodes the audio formats the suite reads and writes; ffprobe "
    "reads stream metadata (duration, codec, bitrate) for analysis tools. The "
    "suite shows a startup warning naming any missing binary and how to install "
    "it."
)
for _hint in (
    "Windows: winget install Gyan.FFmpeg",
    "macOS:   brew install ffmpeg",
    "Linux:   sudo apt install ffmpeg",
):
    doc.add_paragraph(_hint, style="List Bullet")
doc.add_paragraph(
    "Demucs and torch are an optional dependency, installed with the suite's "
    "'ml' extra (uv sync --extra ml). Without it the tool opens, explains what "
    "is missing and refuses a live run; a dry run still works."
)
doc.add_paragraph(
    "The PyPI torch wheel for Windows is a CPU build. A machine with a NVIDIA "
    "card therefore still reports CPU unless a CUDA build of torch is installed; "
    "the Device line on the Setup tab states exactly what torch measured."
)

# ── Quick start ──
doc.add_heading("Quick Start", level=1)
for step in (
    "Open Stem_Separator.pyw.",
    "Select the music directory that holds the files to split.",
    "Choose an output folder for the stems.",
    "Leave all four stems ticked, or clear the ones you do not need.",
    "Leave Dry Run selected and start: the report lists what would be written.",
    "Switch to Live Update and start again to write the stems.",
    "Open the Results tab to read the counts and the log.",
):
    doc.add_paragraph(step, style="List Number")

# ── Setup ──
doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the source and output directory pickers, the stem selection, "
    "the model and device controls, the run mode and the shared run controls."
)
bullet(
    doc,
    "Stem output folder: ",
    "Required. Each track gets its own subfolder inside it. The folder is "
    "remembered between runs.",
)
bullet(
    doc,
    "Stems: ",
    "Vocals, drums, bass and other. At least one must be ticked; only ticked "
    "stems are separated, which shortens a run that needs one part.",
)
bullet(
    doc,
    "Also write mix.wav: ",
    "Writes the sum of the ticked stems next to them. A mix of all four "
    "reconstructs the original track as closely as the model can.",
)
bullet(
    doc,
    "Demucs model: ",
    "htdemucs is the default. htdemucs_ft is the fine-tuned variant: better "
    "stems, roughly four times slower because it runs one model per stem.",
)
bullet(
    doc,
    "Device: ",
    "auto uses CUDA when torch reports it usable and CPU otherwise. cpu and "
    "cuda force one; forcing cuda when it is not usable stops the run with a "
    "message that says why.",
)
bullet(
    doc,
    "Dry Run / Live Update: ",
    "Dry Run reports the plan and writes nothing. Live Update separates and "
    "writes the stems.",
)

# ── How separation works ──
doc.add_heading("How Separation Works", level=1)
for step in (
    "Discover the supported audio files in the selected directory.",
    "Decode the next file once, resampling it to 44.1 kHz stereo.",
    "Run the Demucs model on it to produce the four stems.",
    "Keep the stems you selected and write each one as a 16-bit WAV.",
    "Write mix.wav when a mix was requested.",
    "Report the file and move on to the next one.",
):
    doc.add_paragraph(step, style="List Number")
doc.add_paragraph(
    "Several files are separated in parallel by the shared worker pool. The "
    "model is loaded once per run and reused, so the loading cost is paid once "
    "rather than per file."
)

# ── Output layout ──
doc.add_heading("Output Layout", level=1)
doc.add_paragraph("Each source file writes into its own folder:")
doc.add_paragraph(
    "    <output>/<track>/vocals.wav\n"
    "    <output>/<track>/drums.wav\n"
    "    <output>/<track>/bass.wav\n"
    "    <output>/<track>/other.wav\n"
    "    <output>/<track>/mix.wav        (only when a mix is requested)"
)
doc.add_paragraph(
    "Two sources with the same name but different extensions (song.flac and "
    "song.wav) would target the same folder. The first one is separated and the "
    "second is reported as skipped, so one file never overwrites another's stems."
)

# ── Dry run and live update ──
doc.add_heading("Dry Run and Live Update", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Lists each file and the stems that would be written. It does not load the "
    "model, download weights or create the per-track folder, and it works "
    "without the 'ml' extra.",
)
bullet(
    doc,
    "Live Update: ",
    "Separates each file and writes the stems. The source file is not touched.",
)
bullet(
    doc,
    "Disk space: ",
    "Each stem is a full-length file, so a four-stem separation of a lossless "
    "album needs several times the album's size. WAV stems are uncompressed.",
)

# ── Processing and results ──
doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows the worker slots, progress, processed/separated/skipped/"
    "error counts and a line per file, such as [SEPARATED], [SKIPPED] or "
    "[ERROR]. Results reports the files examined, how many were separated, how "
    "many stem files were written, the device used, the elapsed time and the "
    "output folder, and can open the log and the output folder."
)

# ── Command line ──
doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is stems. Examples:")
doc.add_paragraph(
    "audiosuite stems <folder> --output-dir <stems> --dry\n"
    "audiosuite stems <folder> --output-dir <stems> --stems vocals,bass\n"
    "audiosuite stems <folder> --output-dir <stems> --model htdemucs_ft --mix\n"
    "audiosuite stems <folder> --output-dir <stems> --device cpu --live --json"
)
doc.add_paragraph(
    "Options include --output-dir (required), --model htdemucs|htdemucs_ft, "
    "--stems (comma-separated), --mix/--no-mix, --device auto|cpu|cuda, the "
    "shared --dry/--live option and --json. A live run without the 'ml' extra "
    "exits non-zero with a message naming the missing dependency."
)

# ── Master Driver ──
doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "From the Master Driver, open UTIL and select Stem Separator. The card "
    "takes the output folder, model, device and stem selection. The driver "
    "passes the directory, mode, start/end behaviour and those settings through "
    "the launch handshake. With no output folder the card shows a required-field "
    "hint instead of launching a run that could not write anything."
)

# ── Troubleshooting ──
doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "Demucs is not installed: ",
    "Install the 'ml' extra (uv sync --extra ml). A dry run works without it.",
)
bullet(
    doc,
    "Device says CPU on a machine with a GPU: ",
    "The installed torch is a CPU build. Install a CUDA build of torch; the "
    "reason line names which case applies.",
)
bullet(
    doc,
    "No files found: ",
    "Select a directory containing supported audio files; subfolders are not scanned.",
)
bullet(
    doc,
    "A file errors: ",
    "Check codec support, file permissions, free temporary disk space and the "
    "per-file log line.",
)
bullet(
    doc,
    "The first run pauses: ",
    "Demucs is downloading its model weights. Later runs use the cache.",
)
bullet(
    doc,
    "A track was skipped as a duplicate: ",
    "Another source with the same name already produced that output folder in "
    "the same run. Separate them in separate runs.",
)

# ── Safety ──
doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Use Dry Run first to confirm the file list and the destination, then "
    "audition the stems before deleting anything. Separated stems are never a "
    "perfect reconstruction; keep the original files, and keep a backup of "
    "anything irreplaceable."
)

doc.save(str(OUTPUT))
print(f"Saved to {OUTPUT}")
