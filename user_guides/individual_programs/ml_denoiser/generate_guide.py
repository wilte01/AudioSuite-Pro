"""Generate a fresh ML Denoiser guide from the current implementation.

Changed in 2.0.0: initial version (Sprint 6, Item 6.2).  The version subtitle is
rendered from ``declared_version()`` and the generation month, and the output
path is resolved from this file rather than an absolute checkout path, because
``tests/test_user_guide_hygiene.py`` locks both shapes.
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from audio_suite_shared.project_version import declared_version

OUTPUT = Path(__file__).with_name("ML_Denoiser_User_Guide.docx")


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
h = doc.add_heading("ML Denoiser \u2014 User Guide", level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(
    f"Version {declared_version()} \u2014 {date.today().strftime('%B %Y')}"
).alignment = WD_ALIGN_PARAGRAPH.CENTER

# ── Overview ──
doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "ML Denoiser removes broadband noise with RNNoise, a small recurrent neural "
    "network trained on speech. Unlike the signal-based Noise Reduction tool, it "
    "needs no noise-profile selection: the model estimates, frame by frame, how "
    "much of the signal is speech and reconstructs each frame with the noise "
    "band suppressed."
)
doc.add_paragraph(
    "The model is fast (measured at roughly ten times realtime on a CPU) and it "
    "is speech-oriented. That last part matters: it attenuates tonal and musical "
    "content heavily, so the tool reports the level it measured for every file "
    "and offers an amount blend instead of promising a transparent denoise."
)
doc.add_paragraph(
    "Files are processed in place in Live Update, with existing tags preserved, "
    "and each file is written only after its processed copy has been produced "
    "in a temporary directory."
)

# ── Requirements ──
doc.add_heading("System Requirements", level=1)
doc.add_paragraph(
    "This tool requires ffmpeg and ffprobe to be installed and on PATH. "
    "ffmpeg decodes and encodes the audio formats the suite reads and "
    "writes; ffprobe reads stream metadata (duration, codec, bitrate) "
    "for analysis tools. The suite shows a startup warning naming any "
    "missing binary and how to install it."
)
for _hint in (
    "Windows: winget install Gyan.FFmpeg",
    "macOS:   brew install ffmpeg",
    "Linux:   sudo apt install ffmpeg",
):
    doc.add_paragraph(_hint, style="List Bullet")

doc.add_paragraph(
    "The neural model itself is an optional dependency. It ships in the `ml` "
    "extra (approximately 13 MB, no GPU runtime):"
)
doc.add_paragraph("uv sync --extra ml", style="List Bullet")
doc.add_paragraph(
    "Without that extra the tool opens, validates a run and explains what to "
    "install; it does not start a batch that would fail on every file. A dry "
    "run also needs the extra, because a dry run measures with the model - that "
    "measurement is the point of it."
)

# ── Supported files ──
doc.add_heading("Supported Files", level=1)
doc.add_paragraph(
    "The processor uses the shared Audio Suite supported formats, including MP3, "
    "FLAC, OGG, M4A, MP4, Opus, WAV and AIFF variants. The selected directory "
    "must contain at least one supported file. Read-only formats are reported as "
    "skipped, because the tool rewrites the source file."
)

# ── Quick start ──
doc.add_heading("Quick Start", level=1)
for step in (
    "Open ML_Denoiser.pyw (or launch it from the Master Driver's UTIL tab).",
    "Select the music directory.",
    "Leave Dry Run selected and start; the run reports the attenuation each "
    "file would receive and writes nothing.",
    "Adjust the Denoise Amount if the measured attenuation is too aggressive "
    "for the material.",
    "Switch to Live Update and run again.",
    "Use the Results audition controls to compare an original and a denoised "
    "file when a live result is available.",
):
    doc.add_paragraph(step, style="List Number")

# ── Setup ──
doc.add_heading("Setup", level=1)
doc.add_paragraph(
    "Setup contains the directory picker, the run mode, one denoise control, "
    "presets and the shared run controls."
)
bullet(
    doc,
    "Denoise Amount (%): ",
    "How much of the model's output to keep. 100 % is the model's own result "
    "and 1 % is almost the untouched source. Lower values are a blend, not a "
    "weaker model: the output is mixed with the original signal.",
)
bullet(
    doc,
    "Dry Run / Live Update: ",
    "Dry Run measures and reports without writing; Live Update exports the "
    "processed audio and replaces the source file.",
)
bullet(
    doc,
    "Presets: ",
    "Save and restore the amount and the run mode through the suite's shared presets.",
)

# ── What it does ──
doc.add_heading("How the Model Is Applied", level=1)
doc.add_paragraph(
    "RNNoise works on 10 ms frames of 48 kHz audio. The tool feeds the decoded "
    "file to the model at its own sample rate, lets the library resample "
    "internally, and then forces the result back to the exact source length: a "
    "final frame that the resampler shortened is restored from the source, and "
    "one that overshot is trimmed. The output keeps the source's sample rate, "
    "channel count and duration."
)
doc.add_paragraph(
    "The model's output is 16-bit, so a denoised file carries 16-bit resolution "
    "whatever the source's bit depth was. A file whose measured level does not "
    "move by at least 0.1 dB is reported as clean and is not rewritten, which "
    "makes a second pass over an already-denoised folder a no-op instead of a "
    "re-encode."
)

# ── Modes ──
doc.add_heading("Dry Run and Live Update", level=1)
bullet(
    doc,
    "Dry Run: ",
    "Runs the model, reports each file's measured attenuation in dB and the "
    "before/after level, and writes no audio, no tags and no temporary file. "
    "Result playback is unavailable because nothing was written.",
)
bullet(
    doc,
    "Live Update: ",
    "Exports the processed audio, restores the original tags and replaces the "
    "source file only after that export succeeded.",
)
bullet(
    doc,
    "Clean files: ",
    "Files the model left alone are marked skipped and are not rewritten.",
)
bullet(
    doc,
    "Undo: ",
    "Every live run records a full copy of each file before it is rewritten, so "
    "the suite's Undo Tag Edits workflow restores the original audio and tags "
    "byte for byte.",
)

# ── Results ──
doc.add_heading("Processing and Results", level=1)
doc.add_paragraph(
    "Processing shows the worker slots, a progress bar and per-file statuses: "
    "FIXED (rewritten), FOUND (a dry run measured a change), SKIP (nothing to "
    "do) and ERROR. The status line carries the measured attenuation, for "
    "example 'Attenuated 6.0 dB (-18.4 → -24.4 dBFS)'. Results reports total, "
    "denoised, skipped, errors and elapsed time, and can open the log and the "
    "source folder."
)

# ── CLI ──
doc.add_heading("Command Line", level=1)
doc.add_paragraph(
    "The registered command is denoise. The amount is a percentage here too, "
    "the same unit the GUI slider shows. Examples:"
)
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app denoise <folder> --dry\n"
    "python -m audio_suite_shared.cli_app denoise <folder> --amount 50 --dry\n"
    "python -m audio_suite_shared.cli_app denoise <folder> --amount 100 --live\n"
    "python -m audio_suite_shared.cli_app denoise <folder> --live --json"
)
doc.add_paragraph(
    "Options: --amount 1-100, the shared --dry/--live pair and --json. A run "
    "without the `ml` extra exits non-zero and names what to install."
)

# ── Master Driver ──
doc.add_heading("Master Driver Use", level=1)
doc.add_paragraph(
    "In the Master Driver's UTIL tab, select ML Denoiser and set its Amount. "
    "The driver passes the directory, the run mode, the start/end behaviour and "
    "the amount through the launch handshake. The module toggle decides whether "
    "the tool joins a batch; its dependency-graph entry is a leaf, so it imposes "
    "no ordering on the other tools."
)

# ── Troubleshooting ──
doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "A dialog about the ml extra: ",
    "Install it with 'uv sync --extra ml' (or 'pip install pyrnnoise'). The "
    "message names the exact command.",
)
bullet(
    doc,
    "Music sounds dull or thin: ",
    "Lower the Denoise Amount. RNNoise is trained on speech and attenuates tonal "
    "content, so any single instrument or sustained note may lose energy.",
)
bullet(
    doc,
    "No files found: ",
    "Select a directory containing supported audio files.",
)
bullet(
    doc,
    "Every file is skipped: ",
    "The model found no measurable attenuation, which is what a dry run over an "
    "already-processed folder looks like. Check the per-file line for the "
    "measured level.",
)
bullet(
    doc,
    "A file errors: ",
    "Check codec support, file permissions, free space in the temporary "
    "directory and the per-file status in the log.",
)

# ── Safety ──
doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Run Dry Run first and, for music, audition representative results: the "
    "model is not a transparent denoiser. Live Update replaces source files only "
    "after a successful temporary export, keeps the tags and records an undo "
    "snapshot. Retain backups of irreplaceable audio; the suite's Undo Tag Edits "
    "workflow restores a whole run, but no tool can recover a file deleted "
    "outside it."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
