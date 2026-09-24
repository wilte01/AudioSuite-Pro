"""Generate a fresh CD Ripper guide from the current implementation."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "cd_ripper\\CD_Ripper_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("CD Ripper - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current CD Ripper implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "CD Ripper reads an audio CD, looks up release metadata, extracts selected "
    "tracks as raw CDDA sectors, and encodes them with FFmpeg into FLAC, MP3, "
    "OGG, M4A, WAV, or Opus files. The GUI separates disc setup and metadata "
    "review from the actual rip progress and final results."
)
doc.add_paragraph(
    "On Windows, extraction uses Win32 SCSI Pass-Through with the MMC READ CD command. "
    "Administrator rights are required for raw CD audio extraction."
)

doc.add_heading("Requirements and Platform", level=1)
bullet(
    doc,
    "Operating system: ",
    "CD extraction is implemented for Windows. The app reports that "
    "ripping is unsupported on other operating systems.",
)
bullet(doc, "Drive: ", "A detected optical CD drive with an inserted audio CD.")
bullet(
    doc,
    "FFmpeg: ",
    "Required for probing and encoding. The app prefers a bundled core "
    "ffmpeg executable when present, otherwise it uses ffmpeg from PATH.",
)
bullet(
    doc,
    "Disc metadata: ",
    "libdiscid and musicbrainzngs are optional but provide the most accurate disc ID "
    "and MusicBrainz lookup. Without libdiscid, ffprobe supplies a TOC fallback.",
)
bullet(
    doc,
    "Windows library: ",
    "The launcher attempts to preload discid.dll when it exists beside the launcher.",
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Insert an audio CD and open CD_Ripper.pyw.",
    "Refresh drives and select the optical drive.",
    "Click Read Disc to read the table of contents and query MusicBrainz.",
    "Review or edit album and track metadata.",
    "Choose an output directory, format, quality, and optional CUE sheet.",
    "Click PROCEED TO RIP, then START RIPPING.",
    "Monitor each track and review the Results tab when complete.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Setup & Lookup", level=1)
doc.add_paragraph(
    "The first tab is disabled from proceeding until a disc has been read. It contains "
    "drive selection, disc reading, C2 diagnostics, "
    "MusicBrainz release candidates, album "
    "metadata, editable track rows, output configuration, and format quality settings."
)
bullet(
    doc,
    "CD Drive: ",
    "Refresh the drive list, select the optical drive, "
    "and insert the audio CD before reading.",
)
bullet(
    doc,
    "Read Disc: ",
    "Reads the TOC in a worker thread, populates track numbers and durations, "
    "and attempts MusicBrainz lookup.",
)
bullet(
    doc,
    "MusicBrainz Matches: ",
    "Choose among returned releases and click Apply. "
    "If there is no match, enter metadata manually.",
)
bullet(
    doc,
    "Album Metadata: ",
    "Edit Album, Album Artist, Year, and Genre, then "
    "Apply to all tracks to propagate album-level fields.",
)
bullet(
    doc,
    "Track metadata: ",
    "Double-click a track Title or Artist to edit it inline. "
    "Track number and duration are not inline-editable.",
)
bullet(
    doc,
    "Output Directory: ",
    "Required before proceeding; the last configured location "
    "is retained by the tool preferences.",
)
bullet(
    doc,
    "Verify Rips: ",
    "Choose the per-track verification level before ripping. "
    "Twice (the default) re-reads each extracted track and compares the two "
    "independent reads. AccurateRip additionally looks up the disc's "
    "AccurateRip database records and cross-checks the rip against them over "
    "the network. Off disables verification entirely and only encodes the "
    "single extraction.",
)

doc.add_heading("MusicBrainz Lookup", level=1)
doc.add_paragraph(
    "When libdiscid is available, the reader calculates a disc ID "
    "and queries MusicBrainz "
    "by that identity, including artists, recordings, and release groups. "
    "If no disc ID is "
    "available, the fallback searches by track count and returns "
    "up to five candidates. "
    "A "
    "successful release applies album, year, album artist, "
    "track titles, and track artists."
)
doc.add_paragraph(
    "Lookup failure is non-fatal: the track list can remain available, "
    "and you can fill in "
    "metadata manually."
)

doc.add_heading("C2 Reliability Diagnostic", level=1)
doc.add_paragraph(
    "Test C2 Reliability reads a small region twice and compares "
    "the results. The diagnostic "
    "can report weak positive evidence that the drive populates C2 error "
    "flags, an all-zero "
    "C2 field that is inconclusive, or differing audio data "
    "that makes the test location "
    "invalid or indicates a possible read problem. "
    "The full detail is written to the CD "
    "ripper log. Treat C2 results as diagnostic evidence, "
    "not a guarantee of drive behavior."
)

doc.add_heading("Output Formats and Quality", level=1)
doc.add_paragraph(
    "The format panel supports FLAC, MP3, OGG, M4A, WAV, and Opus. "
    "Quality choices are format-specific:"
)
bullet(doc, "FLAC: ", "Level 8 max compression, Level 5 balanced, or Level 0 fastest.")
bullet(doc, "MP3: ", "VBR V0, VBR V2, or CBR 320, 256, or 192 kbps.")
bullet(doc, "OGG: ", "Quality 10, 8, 6, 5, or 3.")
bullet(doc, "M4A: ", "AAC at 320, 256, 192, or 128 kbps.")
bullet(doc, "Opus: ", "320, 192, 128, 96, or 64 kbps.")
bullet(doc, "WAV: ", "Uncompressed 16-bit PCM; no quality tiers apply.")
bullet(
    doc,
    "CUE sheet: ",
    "Writes a CUE sheet alongside successfully ripped tracks when enabled.",
)

doc.add_heading("Output Layout and Metadata", level=1)
doc.add_paragraph(
    "Each track is written under the selected output directory using this structure:"
)
doc.add_paragraph(
    "<output>\\<album artist>\\[<year>] <album>\\<track number> - <title>.<extension>"
)
doc.add_paragraph(
    "Invalid filename characters are replaced with underscores. "
    "The encoder writes title, "
    "artist, album artist, album, date, genre, and track metadata "
    "when values are available."
)
doc.add_paragraph(
    "When CUE output is enabled, the generated sheet is named from "
    "the sanitized album title "
    "and includes the successfully encoded tracks and their titles and performers."
)

doc.add_heading("Rip Tab", level=1)
doc.add_paragraph(
    "The Rip tab displays every selected track "
    "with an individual progress bar and status. "
    "START RIPPING begins the session; "
    "Cancel requests cancellation of the current extraction "
    "or encoding task. "
    "The track grid reports extraction and FFmpeg encoding progress. Split "
    "CUE Image opens a CUE sheet and uses the selected output directory "
    "to split a single-file "
    "image into tracks."
)
doc.add_paragraph(
    "For each track, the session extracts to a temporary WAV, "
    "verifies that the extraction is "
    "non-empty, encodes to the selected format, removes the temporary WAV, "
    "and reports success "
    "or the specific failure. The drive is spun down after the session "
    "on a best-effort basis."
)

doc.add_heading("Rip Verification", level=1)
doc.add_paragraph(
    "When verification is enabled, each track is checked after extraction "
    "and before encoding. "
    "A track that fails verification is NOT written to disk; its status reports "
    "MISMATCH and the session marks the task failed."
)
bullet(
    doc,
    "Off: ",
    "No re-read or network cross-check. The single extraction is encoded as before.",
)
bullet(
    doc,
    "Twice: ",
    "The track's sectors are read a second time and the two reads "
    "are compared. Identical reads produce a VERIFIED status. "
    "Twice requires only the local drive and works offline.",
)
bullet(
    doc,
    "AccurateRip: ",
    "The disc's AccurateRip disc ID is computed from the TOC and the "
    "database is queried over the network. The rip is first checked against "
    "the database's main track-window record; if the disc was mastered with "
    "an audio offset, the candidates read adjacent TOC-gap sectors and are "
    "checked again over a wider with-gap window. A cross-check match reports "
    "VERIFIED. If the internet is unreachable the lookup degrades to the "
    "local rip-twice check and reports OFFLINE.",
)
bullet(
    doc,
    "Outcomes: ",
    "VERIFIED, NO ENTRY (disc absent from the AccurateRip database; "
    "an unverified rip is still written), OFFLINE (database unreachable; "
    "degraded to rip-twice), or MISMATCH (reads differ or the database "
    "checksum disagrees; the track is NOT written).",
)
doc.add_paragraph(
    "Drives with an audio offset master discs with adjacent-track data at "
    "each track boundary. AccurateRip's with-gap window accounts for this, "
    "but for such drives the local rip-twice mode is the most reliable "
    "choice because both reads describe the same physical frames."
)
doc.add_paragraph(
    "Separately from press offsets, a drive has a read offset: its raw "
    "audio stream can run a fixed number of frames early or late relative "
    "to the sector it was asked for. Without correction, every AccurateRip "
    "window shifts off the database vectors and a perfectly good rip "
    "reports MISMATCH on every track. CD Ripper measures this offset once "
    "per drive (see Command Line, --detect-read-offset) and stores it in "
    "the shared suite configuration; every rip then re-aligns the "
    "extracted audio and both verification reads automatically. For the "
    "final track of a disc, the shifted tail would need a read past the "
    "lead-out, which drives refuse, so the extracted audio is short by the "
    "offset (a few hundredths of a second) — the same compromise EAC makes "
    "for non-overreading drives, and the AccurateRip checksum is unaffected."
)
doc.add_paragraph(
    "The result of each track's verification is appended to its status, "
    "and the Results summary aggregates the outcomes (for example "
    "'5 VERIFIED, 1 NO ENTRY, 1 failed'), so a successful track without a "
    "verification label was ripped with verification disabled."
)
doc.add_paragraph(
    "With AccurateRip mode the app logs the computed disc ID "
    "(ID1 / ID2 / CDDB) and "
    "whether database records were found, absent, or unavailable for the disc."
)

doc.add_heading("Results", level=1)
doc.add_paragraph(
    "Results shows succeeded and failed counts, "
    "track titles, output filenames, and status. "
    "Open Output Folder opens the destination, "
    "Eject Disc requests an operating-system eject, "
    "and double-clicking a successful row reveals its file in Explorer."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is cd-rip. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app cd-rip D: --output-dir <output> --format "
    "flac\n"
    "python -m audio_suite_shared.cli_app cd-rip D: --output-dir "
    "<output> --format mp3 --tracks 1,2,3\n"
    "python -m audio_suite_shared.cli_app cd-rip D: --output-dir "
    "<output> --format flac --cue\n"
    "python -m audio_suite_shared.cli_app cd-rip D: --output-dir "
    "<output> --dry --format opus\n"
    "python -m audio_suite_shared.cli_app cd-rip D: --output-dir "
    "<output> --format flac --verify accuraterip\n"
    "python -m audio_suite_shared.cli_app cd-rip D: --output-dir "
    "<output> --format flac --verify off\n"
    "python -m audio_suite_shared.cli_app cd-rip D: --detect-read-offset"
)
doc.add_paragraph(
    "CLI options include --format, --tracks, --quality, --cue/--no-cue, --dry/--live, "
    "--verify off|twice|accuraterip (default twice), "
    "and --json/--text. Dry Run prints the selected settings (including the "
    "verification level) and does not "
    "read or write tracks. "
    "Track lists are comma-separated; invalid or empty track selections are rejected. "
    "--output-dir is required for ripping (and for a dry run) but not for "
    "--detect-read-offset."
)
doc.add_paragraph(
    "--detect-read-offset measures the drive's read offset against the "
    "AccurateRip database using the CD currently in the drive: it reads "
    "each track once with a margin, sweeps every candidate frame shift "
    "against the database checksums, requires the winning candidate to "
    "agree across the graded tracks, re-confirms it with the production "
    "checksum pipeline on every track, and persists the result for the "
    "drive in the shared suite configuration. It needs a multi-track disc "
    "present in the AccurateRip database and a working network connection; "
    "on failure nothing is persisted and the command exits non-zero."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No drive detected: ",
    "Confirm Windows sees the optical drive, refresh "
    "the drive list, and insert an audio CD.",
)
bullet(
    doc,
    "Disc cannot be read: ",
    "Check administrator rights, disc condition, "
    "drive access, FFmpeg/ffprobe availability, and the CD ripper log.",
)
bullet(
    doc,
    "No MusicBrainz match: ",
    "Use the fallback track data and enter album, artist, year, genre, "
    "title, and artist values manually.",
)
bullet(
    doc,
    "C2 diagnostic is inconclusive: ",
    "An all-zero C2 field is not proof of either reliable or unreliable C2 support; "
    "test another track or known-problem disc.",
)
bullet(
    doc,
    "Encoding failed: ",
    "Check FFmpeg availability, selected quality, output permissions, disk space, "
    "and the per-track status message.",
)
bullet(
    doc,
    "Track was not written: ",
    "A failed extraction, cancellation, missing TOC entry, failed encode, "
    "or a verification MISMATCH "
    "prevents a successful output file.",
)
bullet(
    doc,
    "Verification reports MISMATCH: ",
    "The re-read differed from the first extraction or the AccurateRip "
    "checksum disagreed. Try cleaning the disc, re-reading it, or switching "
    "to Verify Rips = Twice for drives with an audio offset.",
)
bullet(
    doc,
    "Every track reports AccurateRip MISMATCH on a known-good disc: ",
    "The drive probably has an uncorrected read offset. Run "
    "cd-rip <drive> --detect-read-offset once with a disc that is in the "
    "AccurateRip database; the measured offset is stored per drive and "
    "applied to every future rip automatically.",
)
bullet(
    doc,
    "AccurateRip shows NO ENTRY: ",
    "The disc is not (yet) in the AccurateRip database. "
    "The rip is treated as a single extraction and written anyway; "
    "the label distinguishes this from a verified rip.",
)
bullet(
    doc,
    "AccurateRip shows OFFLINE: ",
    "The database could not be reached; the mode degraded to the local "
    "rip-twice re-read. Verify the network connection.",
)

doc.add_heading("Safety and Recovery", level=1)
doc.add_paragraph(
    "Verify metadata and output settings before ripping a full disc. "
    "The raw extraction uses "
    "low-level drive access and requires administrator rights. "
    "Keep the original CD available "
    "until all tracks and metadata have been reviewed. "
    "A failed, cancelled, or verification-mismatched track does not "
    "produce a completed output "
    "file, and temporary WAV files are cleaned up after each task."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
