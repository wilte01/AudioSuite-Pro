"""Generate a fresh Library Statistics guide from the current implementation."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = (
    "D:\\Users\\wilte\\Desktop\\AudioSuite-Pro\\user_guides\\individual_programs\\"
    "library_stats\\Library_Stats_User_Guide.docx"
)


def bullet(doc: Document, title: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(title).bold = True
    paragraph.add_run(text)


doc = Document()
title = doc.add_heading("Library Statistics Dashboard - User Guide", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Based on the current Library Statistics implementation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Library Statistics Dashboard scans a music library and "
    "aggregates metadata that is "
    "already present in audio tags. It is read-only: it does "
    "not analyze audio, detect "
    "new values, or write tags. The dashboard shows summary "
    "counts, charts, missing-data "
    "lists, and tag completeness results."
)

doc.add_heading("Quick Start", level=1)
for step in (
    "Open Library_Stats.pyw.",
    "Select a music folder or use the restored last folder.",
    "Leave Cache enabled for normal repeated scans, or click "
    "Refresh for a full re-scan.",
    "Click Scan Library.",
    "Review Summary, Charts, Missing Data, and Completeness tabs.",
    "Use Export, Save CSV, or Save JSON when you need a report outside the dashboard.",
):
    doc.add_paragraph(step, style="List Number")

doc.add_heading("Scanning", level=1)
bullet(
    doc,
    "Recursive walk: ",
    "The scanner walks the selected directory tree and includes "
    "supported audio extensions.",
)
bullet(
    doc,
    "Read-only tags: ",
    "It reads BPM, key, genre, mood, tempo marking, album "
    "artist, artist, album, and release "
    "year values.",
)
bullet(
    doc,
    "Formats: ",
    "The format breakdown is based on supported file extensions"
    ", including MP3, FLAC, OGG, "
    "M4A, Opus, WAV, and AIFF variants.",
)
bullet(
    doc,
    "Unreadable files: ",
    "Files that cannot be read are skipped from tag aggregation "
    "rather than stopping the "
    "complete scan.",
)
bullet(
    doc,
    "Background scan: ",
    "The GUI scans on a worker thread and updates the dashboard when complete.",
)

doc.add_heading("Cache and Refresh", level=1)
doc.add_paragraph(
    "The optional cache is a SQLite database at "
    "~/.audiosuite/library_cache.db. It stores "
    "raw tag values, file metadata, modification times, and "
    "paths. With Cache enabled, "
    "unchanged files whose modification time matches the cache "
    "can be reused without "
    "rereading their tags. The report is still aggregated from "
    "the complete current cache "
    "scope."
)
bullet(
    doc,
    "Cache: ",
    "Use for faster repeated scans. The cache is updated for "
    "changed files and stale rows "
    "are removed.",
)
bullet(
    doc,
    "Refresh: ",
    "Clears cached entries for the selected directory and "
    "rebuilds them from the files "
    "currently present.",
)
bullet(
    doc,
    "No cache: ",
    "When Cache is unchecked, the scan reads the library "
    "directly without using the SQLite "
    "cache.",
)

doc.add_heading("Summary Tab", level=1)
doc.add_paragraph(
    "Summary shows total files, distinct albums, album artists"
    ", track artists, and scan "
    "duration. It also shows format counts, a major/minor key "
    "split, and the top artists "
    "by file count."
)

doc.add_heading("Charts Tab", level=1)
doc.add_paragraph(
    "When matplotlib is installed, Charts displays a top-genre "
    "pie chart, BPM distribution "
    "histogram, top-key bar chart, and release-year distribution"
    ". If matplotlib is "
    "unavailable, the tab reports the installation requirement "
    "instead."
)

doc.add_heading("Missing Data Tab", level=1)
doc.add_paragraph(
    "Missing Data reports counts and percentages for files "
    "without BPM, key, genre, mood, "
    "or tempo marking. It uses the total scanned file count as "
    "the denominator and displays "
    "a simple visual bar for each missing-data percentage."
)

doc.add_heading("Completeness Tab", level=1)
doc.add_paragraph(
    "Completeness separately checks whether each supported field"
    " is present: BPM, Key, "
    "Genre, Mood, Tempo Marking, Album Art, and ReplayGain. It "
    "shows total, present, "
    "missing, and percentage-complete values for each field, "
    "plus an overall percentage "
    "across all fields."
)
bullet(doc, "Green: ", "More than 80 percent complete.")
bullet(doc, "Amber: ", "50 through 80 percent complete.")
bullet(doc, "Red: ", "Below 50 percent complete.")
bullet(
    doc,
    "Missing files: ",
    "The tab lists the specific files missing each field; "
    "unreadable files count as missing "
    "for completeness purposes.",
)
bullet(
    doc,
    "Completeness exports: ",
    "Save CSV writes the per-tag table; Save JSON includes the "
    "full report and missing-file "
    "lists.",
)

doc.add_heading("Statistics and Binning", level=1)
bullet(
    doc,
    "BPM: ",
    "Grouped into <60, 60-79, 80-99, 100-119, 120-139, "
    "140-159, 160-179, and 180+ bins.",
)
bullet(
    doc,
    "Keys: ",
    "Aggregated by exact key text and classified into "
    "major/minor where the notation is "
    "recognized.",
)
bullet(
    doc,
    "Moods: ",
    "Semicolon-separated mood values are counted as individual mood labels.",
)
bullet(
    doc,
    "Years: ",
    "Parsed years are grouped into decade labels; missing or "
    "unparseable years are Unknown.",
)
bullet(doc, "Albums: ", "Distinct artist/album pairs are counted as albums.")

doc.add_heading("Exports", level=1)
doc.add_paragraph(
    "After a scan, Export lets the GUI save the metadata report "
    "as CSV or JSON. The CSV is "
    "intended for spreadsheet use; JSON contains the serialized "
    "StatsReport, including "
    "distributions, missing-data counts, top lists, and scan "
    "details. Completeness has "
    "separate CSV and JSON export buttons."
)

doc.add_heading("Command Line", level=1)
doc.add_paragraph("The registered command is stats. Examples:")
doc.add_paragraph(
    "python -m audio_suite_shared.cli_app stats <folder>\n"
    "python -m audio_suite_shared.cli_app stats <folder> "
    "--incremental --json\n"
    "python -m audio_suite_shared.cli_app stats <folder> "
    "--refresh --export-csv report.csv\n"
    "python -m audio_suite_shared.cli_app stats <folder> "
    "--completeness --detail bpm\n"
    "python -m audio_suite_shared.cli_app stats <folder> "
    "--export-json report.json"
)
doc.add_paragraph(
    "Options include --incremental, --refresh, --export-csv, "
    "--export-json, --completeness, "
    "--detail "
    "bpm|key|genre|mood|tempo_marking|album_art|replaygain, "
    "and --json. The CLI "
    "prints aggregate counts, format breakdown, missing-data "
    "percentages, and optional "
    "completeness."
)

doc.add_heading("Troubleshooting", level=1)
bullet(
    doc,
    "No files found: ",
    "Select the correct folder and confirm it contains "
    "supported audio files in the "
    "directory tree.",
)
bullet(
    doc,
    "Charts are unavailable: ",
    "Install matplotlib in the active environment; Summary, "
    "Missing Data, and Completeness "
    "do not depend on charts.",
)
bullet(
    doc,
    "Counts look stale: ",
    "Click Refresh to clear the selected directory's cache and "
    "rebuild it from current files.",
)
bullet(
    doc,
    "A tag appears missing: ",
    "The dashboard only reports existing readable tags; run the "
    "relevant analyzer or tagger "
    "first if the field has not been created.",
)
bullet(
    doc,
    "Export is disabled: ",
    "Complete a scan before exporting. Empty reports cannot be exported from the GUI.",
)

doc.add_heading("Read-Only Safety", level=1)
doc.add_paragraph(
    "Library Statistics only reads audio metadata and writes "
    "reports or cache data. It does "
    "not modify audio files or tags. The cache is separate from "
    "the library and can be "
    "refreshed or cleared for the selected directory."
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")
