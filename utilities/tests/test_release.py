"""Tests for ``utilities/release.py`` — the scripted, idempotent release cut.

The contract under test is not "the steps ran" but "the artifacts are right and
a second run changes nothing": the 1.3.0 cut had to be done by hand, and the
failure modes it exposed (a forgotten ``uv.lock``, a blank ``[Unreleased]``, a
guide re-saved by an external editor) are each pinned here.

The end-to-end test builds a throwaway repository, runs ``main()`` twice and
compares the files byte-for-byte, so idempotency is measured rather than
inferred from the ``[SKIP]`` lines.

Changed in 1.3.0: new suite for the scripted release cut.
Changed in 2.1.0: the fixture ships the shared Markdown renderer instead of a
per-tool `generate_guide.py` (PLAN_docs_consolidation P2.4).
Changed in 1.4.1: the cut refreshes the editable install's metadata before the
gate, so `doctor`'s healthy-install tests see the bumped version.
Changed in 2.1.6: the gate's compile step covers the tracked file list instead
of ``.`` - a host that has built the installer carries a vendored interpreter
whose ``Lib/test`` fixtures are deliberately invalid, and the 2.1.6 cut stopped
on it. ``TestTheGateCompilesTrackedFiles`` pins both directions, because
"ignores broken files" is only half of a gate.

Changed in 2.1.7: release fixtures place guides under ``documentation/user_guides/``.
Changed in 2.2.0: ``TestTheGateStreamsItsProgress`` pins the gate's visibility
when its stdout is redirected - children unbuffered, and each step's marker on
disk before that step finishes (the 2.1.7 cut's log looked stalled at 14%).

Changed in 2.2.1: a cut promotes both change-history files, and the technical
one is optional - a release with no developer-facing change adds no heading
there instead of failing the cut (functional backlog U2).

"""

from __future__ import annotations

import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from utilities import release  # noqa: E402

PYPROJECT = """\
# Changed in 1.0.0: first version of this fixture.
# Keep this comment block: the release note is inserted after it.

[build-system]
requires = ["setuptools>=77"]

[project]
name = "audiosuite-pro"
version = "1.3.0"
requires-python = ">=3.11"
"""

CHANGELOG = """\
# Changelog

## [Unreleased]

### Added

- **A feature worth releasing.** It does the thing.

### Fixed

- **A defect closed.** It no longer misbehaves.

## [1.3.0] - 2026-10-18

### Added

- An older release entry.
"""

TECH_CHANGELOG = """\
# Technical changelog

## [Unreleased]

### Changed

- **A tooling change worth recording.** It changes how the gate runs.

## [1.3.0] - 2026-10-18

### Fixed

- An older technical entry.
"""

#: A technical changelog with nothing to promote: a release with no
#: developer-facing change must be legal.
EMPTY_TECH_CHANGELOG = """\
# Technical changelog

## [Unreleased]

Nothing yet.

## [1.3.0] - 2026-10-18

### Fixed

- An older technical entry.
"""

STUB_RENDERER = '''\
"""Stub shared guide renderer: renders the version from pyproject.toml."""

import sys
import tomllib
from pathlib import Path

from docx import Document

root = Path(__file__).resolve().parents[1]
version = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))[
    "project"
]["version"]
document = Document()
document.add_heading("Demo Tool User Guide", level=1)
document.add_paragraph(f"Version {version} \\u2014 November 2026")
document.save(
    str(
        root
        / "documentation"
        / "user_guides"
        / "individual_programs"
        / "demo"
        / "Demo_User_Guide.docx"
    )
)
sys.exit(0)
'''


def _write_guide(path: Path, label: str) -> None:
    """Create a minimal ``.docx`` carrying a version subtitle."""
    from docx import Document

    document = Document()
    document.add_heading("Demo Tool User Guide", level=1)
    document.add_paragraph(label)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _replace_entry(path: Path, name: str, data: bytes) -> None:
    """Rewrite one entry of a zip container, leaving the rest byte-identical."""
    with zipfile.ZipFile(path) as archive:
        entries = {entry: archive.read(entry) for entry in archive.namelist()}
    entries[name] = data
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for entry, payload in entries.items():
            archive.writestr(entry, payload)


class TestVersionBump:
    """``[project] version`` is the only version constant; one edit moves it."""

    def test_bumps_version_and_adds_release_note(self) -> None:
        text, old = release.bump_project_version(PYPROJECT, "1.4.0")

        assert old == "1.3.0"
        assert 'version = "1.4.0"' in text
        assert 'version = "1.3.0"' not in text
        assert "bumped [project] version 1.3.0 -> 1.4.0" in text
        assert "Changed in 1.4.0:" in text

    def test_second_call_is_a_noop(self) -> None:
        once, _old = release.bump_project_version(PYPROJECT, "1.4.0")
        twice, old = release.bump_project_version(once, "1.4.0")

        assert old is None
        assert twice == once
        assert twice.count("bumped [project] version") == 1

    def test_missing_version_is_an_error(self) -> None:
        with pytest.raises(release.ReleaseError, match="no \\[project\\] version"):
            release.bump_project_version('[project]\nname = "x"\n', "1.4.0")

    def test_line_endings_survive_the_bump(self) -> None:
        crlf = PYPROJECT.replace("\n", "\r\n")

        text, _old = release.bump_project_version(crlf, "1.4.0")

        assert 'version = "1.4.0"\r\n' in text
        assert "\n" not in text.replace("\r\n", "")


class TestChangelogRotation:
    """Promotion keeps the released body and reopens the section."""

    def test_promotes_entries_and_reopens_unreleased(self) -> None:
        text, changed = release.promote_changelog(CHANGELOG, "1.4.0", "2026-11-02")

        assert changed
        unreleased = text.index("## [Unreleased]")
        placeholder = text.index("Nothing yet.", unreleased)
        released = text.index("## [1.4.0] - 2026-11-02", placeholder)
        assert unreleased < placeholder < released
        assert "**A feature worth releasing.**" in text[released:]
        assert "## [1.3.0] - 2026-10-18" in text[released:]

    def test_second_call_is_a_noop(self) -> None:
        once, _changed = release.promote_changelog(CHANGELOG, "1.4.0", "2026-11-02")
        twice, changed = release.promote_changelog(once, "1.4.0", "2026-11-02")

        assert changed is False
        assert twice == once

    def test_empty_unreleased_refuses_to_release(self) -> None:
        empty = (
            "# Changelog\n\n## [Unreleased]\n\nNothing yet.\n\n"
            "## [1.3.0] - 2026-10-18\n\n- An older release.\n"
        )

        with pytest.raises(release.ReleaseError, match="no entries"):
            release.promote_changelog(empty, "1.4.0", "2026-11-02")

    def test_an_optional_file_with_nothing_to_promote_is_left_alone(self) -> None:
        """A release with no developer-facing change must not fail the cut."""
        empty = (
            "# Technical changelog\n\n## [Unreleased]\n\nNothing yet.\n\n"
            "## [1.3.0] - 2026-10-18\n\n- An older entry.\n"
        )

        text, changed = release.promote_changelog(
            empty, "1.4.0", "2026-11-02", require_entries=False
        )

        assert changed is False
        assert text == empty

    def test_an_optional_file_still_promotes_what_it_has(self) -> None:
        text, changed = release.promote_changelog(
            TECH_CHANGELOG, "1.4.0", "2026-11-02", require_entries=False
        )

        assert changed
        assert "## [1.4.0] - 2026-11-02" in text
        released = text[text.index("[1.4.0]") :]
        reopened = text[text.index("[Unreleased]") : text.index("[1.4.0]")]
        assert "**A tooling change worth recording.**" in released
        assert "Nothing yet." in reopened

    def test_a_required_file_is_unchanged_by_the_optional_rule(self) -> None:
        """The default still refuses an empty user-facing section."""
        empty = (
            "# Changelog\n\n## [Unreleased]\n\nNothing yet.\n\n"
            "## [1.3.0] - 2026-10-18\n\n- An older release.\n"
        )

        with pytest.raises(release.ReleaseError, match="no entries"):
            release.promote_changelog(
                empty, "1.4.0", "2026-11-02", require_entries=True
            )

    def test_unbounded_unreleased_section_is_an_error(self) -> None:
        text = "# Changelog\n\n## [Unreleased]\n\n- Something.\n"

        with pytest.raises(release.ReleaseError, match="not followed"):
            release.promote_changelog(text, "1.4.0", "2026-11-02")

    def test_already_released_section_reopens_the_placeholder(self) -> None:
        text = (
            "# Changelog\n\n## [Unreleased]\n\n## [1.4.0] - 2026-11-02\n\n"
            "- Something.\n"
        )

        new_text, changed = release.promote_changelog(text, "1.4.0", "2026-11-02")

        assert changed
        assert "## [Unreleased]\n\nNothing yet.\n\n## [1.4.0]" in new_text


class TestInstallRefresh:
    """A bump must refresh the installed metadata before the gate reads it."""

    def test_uses_uv_for_a_dependency_free_editable_reinstall(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """uv (not pip) and no dependency work, from the release root.

        pip cannot do this in a uv-managed venv: there is no ``setuptools``,
        so ``--no-build-isolation`` fails, and an isolated build would fetch a
        backend the lock never pinned.
        """
        calls: list[list[str]] = []

        def fake_run(args: list[str], **_kwargs: object) -> object:
            calls.append(args)
            return subprocess.CompletedProcess(args, 0, "", "")

        monkeypatch.setattr(release.shutil, "which", lambda name: "/usr/bin/uv")
        monkeypatch.setattr(release.subprocess, "run", fake_run)
        release.refresh_install_metadata(tmp_path, "python")

        assert len(calls) == 1
        args = calls[0]
        assert args[:4] == ["/usr/bin/uv", "pip", "install", "--no-deps"]
        assert args[4:6] == ["--python", "python"]
        assert args[6] == "-e"
        assert args[7] == str(tmp_path)

    def test_a_failed_refresh_is_a_release_error(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A silent failure would let the gate read the old version."""

        def fake_run(args: list[str], **_kwargs: object) -> object:
            return subprocess.CompletedProcess(args, 1, "boom", "")

        monkeypatch.setattr(release.shutil, "which", lambda name: "/usr/bin/uv")
        monkeypatch.setattr(release.subprocess, "run", fake_run)
        with pytest.raises(release.ReleaseError, match="metadata"):
            release.refresh_install_metadata(tmp_path, "python")

    def test_missing_uv_is_a_release_error(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Without uv the environment cannot be refreshed, so say so."""
        monkeypatch.setattr(release.shutil, "which", lambda name: None)
        with pytest.raises(release.ReleaseError, match="uv is not on PATH"):
            release.refresh_install_metadata(tmp_path, "python")


class TestBanditVerdict:
    """The gate reads the severity block, because exit 1 is normal here."""

    def _output(self, high: int, medium: int, low: int = 41) -> str:
        return (
            "Run metrics:\n"
            "\tTotal issues (by severity):\n"
            f"\t\tLow: {low}\n"
            f"\t\tMedium: {medium}\n"
            f"\t\tHigh: {high}\n"
            "\tTotal issues (by confidence):\n"
            "\t\tHigh: 18\n"
        )

    def test_zero_high_and_medium_passes_even_with_exit_one(self) -> None:
        ok, summary = release.bandit_verdict(1, self._output(0, 0))

        assert ok
        assert "Medium=0" in summary

    def test_a_medium_finding_fails(self) -> None:
        ok, _summary = release.bandit_verdict(1, self._output(0, 2))

        assert not ok

    def test_tool_error_fails_even_with_a_clean_block(self) -> None:
        ok, _summary = release.bandit_verdict(2, self._output(0, 0))

        assert not ok

    def test_unparseable_output_fails_loudly(self) -> None:
        ok, summary = release.bandit_verdict(0, "Traceback: bandit broke")

        assert not ok
        assert "could not parse" in summary


class TestGuideChange:
    """A version-only regeneration may touch exactly one zip entry."""

    def test_version_rewrite_is_accepted(self, tmp_path: Path) -> None:
        guide = tmp_path / "Demo_User_Guide.docx"
        _write_guide(guide, "Version 1.3.0 \u2014 October 2026")
        before = release.entry_data(guide)

        _write_guide(guide, "Version 1.4.0 \u2014 November 2026")

        assert release.guide_change_problems(guide, before, "1.4.0") == []

    def test_unchanged_regeneration_is_rejected(self, tmp_path: Path) -> None:
        guide = tmp_path / "Demo_User_Guide.docx"
        _write_guide(guide, "Version 1.3.0 \u2014 October 2026")
        before = release.entry_data(guide)

        problems = release.guide_change_problems(guide, before, "1.4.0")

        assert problems, "a regeneration that changed nothing must fail the run"
        assert "hardcoded" in problems[0]

    def test_external_tool_drift_is_rejected(self, tmp_path: Path) -> None:
        guide = tmp_path / "Demo_User_Guide.docx"
        _write_guide(guide, "Version 1.3.0 \u2014 October 2026")
        before = release.entry_data(guide)

        _write_guide(guide, "Version 1.4.0 \u2014 November 2026")
        _replace_entry(guide, "docProps/app.xml", b"<Properties>re-saved</Properties>")

        problems = release.guide_change_problems(guide, before, "1.4.0")

        assert any("external-tool drift" in problem for problem in problems)

    def test_wrong_subtitle_is_rejected(self, tmp_path: Path) -> None:
        guide = tmp_path / "Demo_User_Guide.docx"
        _write_guide(guide, "Version 1.3.0 \u2014 October 2026")
        before = release.entry_data(guide)

        _write_guide(guide, "Version 9.9.9 \u2014 October 2026")

        problems = release.guide_change_problems(guide, before, "1.4.0")

        assert any("does not start with" in problem for problem in problems)


class TestGuideDiscovery:
    """Only guides that actually display a version are regenerated."""

    def test_finds_only_labelled_guides(self, tmp_path: Path) -> None:
        labelled = (
            tmp_path
            / "documentation"
            / "user_guides"
            / "individual_programs"
            / "demo"
            / "Demo_User_Guide.docx"
        )
        unlabelled = labelled.parent.parent / "other" / "Other_User_Guide.docx"
        _write_guide(labelled, "Version 1.3.0 \u2014 October 2026")
        _write_guide(unlabelled, "No version here")

        found = release.version_labelled_guides(tmp_path)

        assert found == [labelled]

    def test_no_labelled_guide_fails_loudly(self, tmp_path: Path) -> None:
        unlabelled = (
            tmp_path
            / "documentation"
            / "user_guides"
            / "individual_programs"
            / "other"
            / "Other_User_Guide.docx"
        )
        _write_guide(unlabelled, "No version here")

        with pytest.raises(release.ReleaseError, match="vacuous"):
            release.version_labelled_guides(tmp_path)


@pytest.fixture()
def release_repo(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A throwaway repository whose lock step never shells out to uv."""
    (tmp_path / "pyproject.toml").write_text(PYPROJECT, encoding="utf-8")
    (tmp_path / "CHANGELOG.md").write_text(CHANGELOG, encoding="utf-8")
    (tmp_path / "TECHNICAL_CHANGELOG.md").write_text(TECH_CHANGELOG, encoding="utf-8")
    demo = tmp_path / "documentation" / "user_guides" / "individual_programs" / "demo"
    demo.mkdir(parents=True)
    (demo / "Demo_User_Guide.md").write_text("# demo\n", encoding="utf-8")
    utilities = tmp_path / "utilities"
    utilities.mkdir()
    (utilities / "build_user_guides.py").write_text(STUB_RENDERER, encoding="utf-8")
    _write_guide(demo / "Demo_User_Guide.docx", "Version 1.3.0 \u2014 October 2026")

    def fake_lock_is_current(root: Path, offline: bool) -> bool:
        lock = root / "uv.lock"
        if not lock.is_file():
            return False
        declared = release.declared_version(
            (root / "pyproject.toml").read_text(encoding="utf-8")
        )
        return f'version = "{declared}"' in lock.read_text(encoding="utf-8")

    def fake_update_lock(root: Path, offline: bool) -> None:
        declared = release.declared_version(
            (root / "pyproject.toml").read_text(encoding="utf-8")
        )
        (root / "uv.lock").write_text(
            f'[[package]]\nname = "audiosuite-pro"\nversion = "{declared}"\n',
            encoding="utf-8",
        )

    monkeypatch.setattr(release, "lock_is_current", fake_lock_is_current)
    monkeypatch.setattr(release, "update_lock", fake_update_lock)
    return tmp_path


class TestEndToEnd:
    """One command cuts the release; a second run changes no byte."""

    def test_first_run_updates_every_artifact(
        self, release_repo: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        code = release.main(
            [
                "1.4.0",
                "--root",
                str(release_repo),
                "--date",
                "2026-11-02",
                "--skip-gate",
            ]
        )

        out = capsys.readouterr().out
        assert code == 0, out
        assert 'version = "1.4.0"' in (release_repo / "pyproject.toml").read_text(
            encoding="utf-8"
        )
        assert "## [1.4.0] - 2026-11-02" in (release_repo / "CHANGELOG.md").read_text(
            encoding="utf-8"
        )
        assert "## [1.4.0] - 2026-11-02" in (
            release_repo / "TECHNICAL_CHANGELOG.md"
        ).read_text(encoding="utf-8")
        guide = (
            release_repo
            / "documentation"
            / "user_guides"
            / "individual_programs"
            / "demo"
            / "Demo_User_Guide.docx"
        )
        assert release.guide_version_label(guide) == (
            "Version 1.4.0 \u2014 November 2026"
        )
        assert '"1.4.0"' in (release_repo / "uv.lock").read_text(encoding="utf-8")

    def test_second_run_is_byte_identical(
        self, release_repo: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        args = [
            "1.4.0",
            "--root",
            str(release_repo),
            "--date",
            "2026-11-02",
            "--skip-gate",
        ]
        assert release.main(args) == 0
        capsys.readouterr()

        artifacts = [
            release_repo / "pyproject.toml",
            release_repo / "CHANGELOG.md",
            release_repo / "uv.lock",
            release_repo
            / "documentation"
            / "user_guides"
            / "individual_programs"
            / "demo"
            / "Demo_User_Guide.docx",
        ]
        first = {path: path.read_bytes() for path in artifacts}

        assert release.main(args) == 0
        out = capsys.readouterr().out

        assert {path: path.read_bytes() for path in artifacts} == first
        assert out.count("[SKIP]") == 5, out
        assert "Demo_User_Guide.docx already labels 1.4.0" in out

    def test_dry_run_writes_nothing(
        self, release_repo: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        artifacts = [
            release_repo / "pyproject.toml",
            release_repo / "CHANGELOG.md",
            release_repo
            / "documentation"
            / "user_guides"
            / "individual_programs"
            / "demo"
            / "Demo_User_Guide.docx",
        ]
        before = {path: path.read_bytes() for path in artifacts}

        code = release.main(["1.4.0", "--root", str(release_repo), "--dry-run"])

        out = capsys.readouterr().out
        assert code == 0, out
        assert "would bump pyproject.toml 1.3.0 -> 1.4.0" in out
        assert "would promote CHANGELOG.md" in out
        assert "would regenerate demo/Demo_User_Guide.docx" in out
        assert {path: path.read_bytes() for path in artifacts} == before
        assert not (release_repo / "uv.lock").exists()

    def test_a_release_with_no_technical_work_still_cuts(
        self, release_repo: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        """The technical file is optional; the user-facing one is not."""
        technical = release_repo / "TECHNICAL_CHANGELOG.md"
        technical.write_text(EMPTY_TECH_CHANGELOG, encoding="utf-8")

        code = release.main(
            [
                "1.4.0",
                "--root",
                str(release_repo),
                "--date",
                "2026-11-02",
                "--skip-gate",
            ]
        )

        out = capsys.readouterr().out
        assert code == 0, out
        assert "TECHNICAL_CHANGELOG.md holds no entries for [1.4.0]" in out
        assert technical.read_text(encoding="utf-8") == EMPTY_TECH_CHANGELOG
        assert "## [1.4.0] - 2026-11-02" in (release_repo / "CHANGELOG.md").read_text(
            encoding="utf-8"
        )

    def test_a_user_facing_release_with_no_entries_is_still_refused(
        self, release_repo: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        """The optional rule must not make the user changelog optional."""
        (release_repo / "CHANGELOG.md").write_text(
            "# Changelog\n\n## [Unreleased]\n\nNothing yet.\n\n"
            "## [1.3.0] - 2026-10-18\n\n- An older release.\n",
            encoding="utf-8",
        )

        code = release.main(["1.4.0", "--root", str(release_repo), "--skip-gate"])

        assert code == 1
        assert "no entries" in capsys.readouterr().err

    def test_malformed_version_is_refused_before_any_write(
        self, release_repo: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        before = (release_repo / "pyproject.toml").read_bytes()

        code = release.main(["1.4", "--root", str(release_repo), "--skip-gate"])

        assert code == 1
        assert "not an X.Y.Z version" in capsys.readouterr().err
        assert (release_repo / "pyproject.toml").read_bytes() == before


class TestTheGateCompilesTrackedFiles:
    """The compile step must see first-party source, not whatever is on disk.

    ``compileall .`` walks ignored build output, and a host that has built the
    installer carries ``packaging/build/app`` - a vendored interpreter whose
    ``Lib/test`` fixtures are deliberately invalid - so the 2.1.6 cut failed a
    gate that had nothing to do with the tree. Both directions are pinned: an
    untracked broken file is invisible, and a tracked one still fails, because a
    step that ignores everything would be worse than the defect.
    """

    @staticmethod
    def _repo(tmp_path: Path) -> Path:
        """A throwaway git repository holding one valid tracked module."""

        def run(*args: str) -> None:
            subprocess.run(
                ["git", *args], cwd=tmp_path, capture_output=True, check=True
            )

        run("init")
        (tmp_path / "good.py").write_text("VALUE = 1\n", encoding="utf-8")
        run("add", "good.py")
        return tmp_path

    def test_a_broken_untracked_artefact_cannot_fail_the_gate(
        self, tmp_path: Path
    ) -> None:
        repo = self._repo(tmp_path)
        staged = repo / "packaging" / "build" / "Lib" / "test"
        staged.mkdir(parents=True)
        # Non-UTF-8 byte and no coding declaration: the shape that failed the cut.
        (staged / "badsyntax_pep3120.py").write_bytes(b"print('b\xf6se')\n")

        assert release.compile_tracked(repo, sys.executable) == 0

    def test_a_broken_tracked_file_still_fails_it(self, tmp_path: Path) -> None:
        repo = self._repo(tmp_path)
        (repo / "bad.py").write_bytes(b"print('b\xf6se')\n")
        subprocess.run(
            ["git", "add", "bad.py"], cwd=repo, capture_output=True, check=True
        )

        assert release.compile_tracked(repo, sys.executable) != 0


class TestTheGateStreamsItsProgress:
    """A backgrounded cut must show which step is running while it runs.

    A cut is launched with stdout redirected to a log, and Python block-buffers
    a redirected stdout: each child then writes in 8 KiB bursts, so the 2.1.7
    run arrived in a few chunks and read as stuck at 14% for most of 21
    minutes. Both halves are pinned here. The children must run unbuffered, and
    the parent's own marker must be on disk **before** the step it announces
    runs - a buffered ``[START] pytest`` would otherwise surface after pytest's
    output and misreport the order the gate ran in.
    """

    #: Bandit's severity block, so the security step is judged, not skipped.
    _BANDIT = (
        "\tTotal issues (by severity):\n"
        "\t\tLow: 41\n"
        "\t\tMedium: 0\n"
        "\t\tHigh: 0\n"
        "\tTotal issues (by confidence):\n"
        "\t\tHigh: 18\n"
    )

    def test_every_gate_child_runs_unbuffered(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """ruff, mypy, bandit and pytest all receive ``PYTHONUNBUFFERED``."""
        calls: list[tuple[list[str], dict[str, object]]] = []

        def fake_run(args: list[str], **kwargs: object) -> object:
            calls.append((args, kwargs))
            stdout = self._BANDIT if "bandit" in args else ""
            return subprocess.CompletedProcess(args, 0, stdout, "")

        monkeypatch.setattr(release, "compile_tracked", lambda root, python: 0)
        monkeypatch.setattr(release.subprocess, "run", fake_run)

        assert release.run_gate(tmp_path, "python") == 0

        # Non-vacuity: the gate's own steps, not just the compile helper.
        assert len(calls) >= 5, [args for args, _ in calls]
        assert any("pytest" in args for args, _ in calls), [a for a, _ in calls]
        for args, kwargs in calls:
            env = kwargs.get("env")
            assert isinstance(env, dict), args
            assert env.get("PYTHONUNBUFFERED") == "1", args

    def test_a_step_marker_is_on_disk_before_that_step_runs(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """``[START] pytest`` is readable while pytest is still running.

        Asserted against a real, still-open text file rather than ``capsys``:
        ``capsys`` captures regardless of flushing, so it cannot tell a
        flushed line from a buffered one and would pass either way.
        """
        log = tmp_path / "gate.log"
        seen: list[str] = []

        def fake_run(args: list[str], **_kwargs: object) -> object:
            stdout = self._BANDIT if "bandit" in args else ""
            if "pytest" in args:
                # A second handle sees only bytes already flushed to disk.
                seen.append(log.read_text(encoding="utf-8"))
            return subprocess.CompletedProcess(args, 0, stdout, "")

        monkeypatch.setattr(release, "compile_tracked", lambda root, python: 0)
        monkeypatch.setattr(release.subprocess, "run", fake_run)

        with log.open("w", encoding="utf-8", newline="") as handle:
            monkeypatch.setattr(sys, "stdout", handle)
            assert release.run_gate(tmp_path, "python") == 0

        assert seen, "the gate must reach its test step"
        assert "[START] pytest" in seen[0]
        assert "[OK] mypy" in seen[0], seen[0]
        assert "[OK] pytest" not in seen[0], "pytest's result cannot precede its run"
        assert "[DONE] gate passed" in log.read_text(encoding="utf-8")
